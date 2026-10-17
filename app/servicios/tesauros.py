"""
Tesauros: importarlos y normalizar con ellos las materias de la descripción.

- Importar: un archivo SKOS (Turtle, RDF/XML, N-Triples o JSON-LD, como el
  Tesauro de la UNESCO), o una tabla Excel o CSV (para un tesauro que no
  publica SKOS, como el ambiental de Minambiente). Columnas de la tabla:
  término preferente, términos no preferentes («usado por», separados por
  «;»), nota de alcance, término general (superior), URI y código.
- Normalizar: un término escrito se busca entre las etiquetas de los
  conceptos, sin tildes ni mayúsculas. Si es la preferente, se usa; si es
  una no preferente («archivos históricos» → «Archivo histórico»), se usa la
  preferente; si no aparece, se ofrecen los parecidos (trigramas) y, si no
  hay ninguno, queda como término local, marcado como tal.
- Materia: la entidad «materia» del vocabulario del fondo que corresponde a
  un concepto. Hay una sola por concepto y fondo: dos documentos con el
  mismo tema llegan al mismo nodo del grafo.
"""

import csv
import hashlib
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass, field

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models.descripcion import EntidadVocabulario
from app.models.tesauro import ConceptoTesauro, EtiquetaConcepto, Tesauro
from app.servicios.auditoria import registrar

UMBRAL_PARECIDO = 0.45
MAX_BYTES = 80 * 1024 * 1024
MAX_CONCEPTOS = 60000


class ErrorTesauro(Exception):
    def __init__(self, mensaje: str, codigo: int = 422):
        super().__init__(mensaje)
        self.codigo = codigo


def normalizar(texto: str | None) -> str:
    """Igual que el vocabulario del fondo: sin tildes, en minúsculas, sin
    espacios de más (y sin los signos de los extremos)."""
    sin_tildes = "".join(c for c in unicodedata.normalize("NFKD", texto or "") if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sin_tildes).strip(" .;:,-").lower()


@dataclass
class ConceptoLeido:
    clave: str
    preferente: str
    uri: str | None = None
    no_preferentes: list[str] = field(default_factory=list)
    ocultas: list[str] = field(default_factory=list)
    nota: str | None = None
    notacion: str | None = None
    superior: str | None = None  # clave del concepto superior


# --- Lectura -----------------------------------------------------------------------------------------


def _formato_rdf(archivo: str) -> list[str]:
    ext = archivo.lower().rsplit(".", 1)[-1] if "." in archivo else ""
    preferido = {"ttl": "turtle", "rdf": "xml", "xml": "xml", "owl": "xml", "nt": "nt", "jsonld": "json-ld",
                 "json": "json-ld", "n3": "n3"}.get(ext)
    return ([preferido] if preferido else []) + [f for f in ("turtle", "xml", "nt", "json-ld") if f != preferido]


def leer_skos(contenido: bytes, archivo: str, idioma: str = "es") -> tuple[list[ConceptoLeido], str | None]:
    from rdflib import OWL, RDF, Graph, Literal, Namespace, URIRef

    skos = Namespace("http://www.w3.org/2004/02/skos/core#")
    g, ultimo = None, None
    for formato in _formato_rdf(archivo):
        try:
            g = Graph()
            g.parse(data=contenido, format=formato)
            break
        except Exception as exc:  # noqa: BLE001 — se prueba el siguiente formato
            g, ultimo = None, exc
    if g is None:
        raise ErrorTesauro(f"No se pudo leer el archivo como SKOS ({ultimo}).")

    def textos(nodo, prop) -> list[str]:
        """Las del idioma pedido; si no hay, las sin idioma; la etiqueta
        preferente, además, en cualquier idioma antes que ninguna."""
        valores = [o for o in g.objects(nodo, prop) if isinstance(o, Literal)]
        del_idioma = [str(v) for v in valores if (v.language or "").lower().startswith(idioma)]
        sin_idioma = [str(v) for v in valores if not v.language]
        otros = [str(v) for v in valores][:1] if prop == skos.prefLabel else []
        return del_idioma or sin_idioma or otros

    esquema = next((str(s) for s in g.subjects(RDF.type, skos.ConceptScheme)), None)
    salida = []
    for c in set(g.subjects(RDF.type, skos.Concept)):
        if not isinstance(c, URIRef) or (c, OWL.deprecated, Literal(True)) in g:
            continue
        pref = textos(c, skos.prefLabel)
        if not pref:
            continue
        superior = next((str(o) for o in g.objects(c, skos.broader) if isinstance(o, URIRef)), None)
        notas = textos(c, skos.scopeNote) or textos(c, skos.definition)
        notacion = next((str(o) for o in g.objects(c, skos.notation)), None)
        salida.append(ConceptoLeido(clave=str(c), uri=str(c), preferente=pref[0], no_preferentes=textos(c, skos.altLabel),
                                    ocultas=textos(c, skos.hiddenLabel), nota=notas[0] if notas else None,
                                    notacion=notacion, superior=superior))
    if not salida:
        raise ErrorTesauro("El archivo no tiene conceptos SKOS (skos:Concept con skos:prefLabel).")
    return salida, esquema


COLUMNAS = {
    "preferente": ("termino preferente", "termino", "descriptor", "preferente", "preflabel", "etiqueta preferente",
                   "termino aceptado"),
    "no_preferentes": ("terminos no preferentes", "no preferentes", "no preferente", "usado por", "up", "altlabel",
                       "sinonimos", "variantes", "terminos equivalentes"),
    "nota": ("nota de alcance", "nota", "na", "scopenote", "definicion", "nota de aplicacion"),
    "superior": ("termino general", "tg", "superior", "broader", "termino superior", "termino generico"),
    "uri": ("uri", "url", "enlace"),
    "notacion": ("codigo", "notacion", "notation"),
}


def _filas_tabla(contenido: bytes, archivo: str) -> list[list[str]]:
    if archivo.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ErrorTesauro(f"No se pudo abrir el Excel ({exc}).") from exc
        hoja = libro.worksheets[0]
        return [["" if v is None else str(v).strip() for v in fila] for fila in hoja.iter_rows(values_only=True)]
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:5000], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    return [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), dialecto)]


def leer_tabla(contenido: bytes, archivo: str) -> list[ConceptoLeido]:
    filas = [f for f in _filas_tabla(contenido, archivo) if any(f)]
    if not filas:
        raise ErrorTesauro("La tabla está vacía.")
    cabecera = [normalizar(c) for c in filas[0]]
    indice = {}
    for campo, nombres in COLUMNAS.items():
        for i, c in enumerate(cabecera):
            if c in nombres:
                indice[campo] = i
                break
    if "preferente" not in indice:
        raise ErrorTesauro("Falta la columna «Término preferente» en la primera fila de la tabla.")
    salida, vistos = [], set()

    def celda(fila, campo):
        i = indice.get(campo)
        return fila[i].strip() if i is not None and i < len(fila) and fila[i] else ""

    for fila in filas[1:]:
        pref = " ".join(celda(fila, "preferente").split())
        if not pref or normalizar(pref) in vistos:
            continue
        vistos.add(normalizar(pref))
        variantes = [v.strip() for v in re.split(r"[;|\n]", celda(fila, "no_preferentes")) if v.strip()]
        superior = " ".join(celda(fila, "superior").split())
        salida.append(ConceptoLeido(clave=normalizar(pref), preferente=pref, uri=celda(fila, "uri") or None,
                                    no_preferentes=variantes, nota=celda(fila, "nota") or None,
                                    notacion=celda(fila, "notacion") or None,
                                    superior=normalizar(superior) if superior else None))
    if not salida:
        raise ErrorTesauro("La tabla no tiene términos preferentes.")
    return salida


def leer(contenido: bytes, archivo: str, idioma: str = "es") -> tuple[list[ConceptoLeido], str | None, str]:
    nombre = archivo.lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        return leer_tabla(contenido, archivo), None, "excel"
    if nombre.endswith((".csv", ".tsv", ".txt")):
        return leer_tabla(contenido, archivo), None, "csv"
    conceptos, esquema = leer_skos(contenido, archivo, idioma)
    return conceptos, esquema, "skos"


# --- Importación -------------------------------------------------------------------------------------


def importar(db: Session, *, nombre: str, contenido: bytes, archivo: str, fuente: str | None, licencia: str | None,
             idioma: str, usuario_id: uuid.UUID | None) -> Tesauro:
    if not contenido:
        raise ErrorTesauro("El archivo está vacío.")
    if len(contenido) > MAX_BYTES:
        raise ErrorTesauro(f"El archivo pasa de {MAX_BYTES // (1024 * 1024)} MB.", 413)
    huella = hashlib.sha256(contenido).hexdigest()
    if db.scalar(select(Tesauro.id).where(Tesauro.huella == huella, Tesauro.estado == "vigente")):
        raise ErrorTesauro("Ese mismo archivo ya está importado.", 409)
    conceptos, esquema, formato = leer(contenido, archivo, idioma)
    if len(conceptos) > MAX_CONCEPTOS:
        raise ErrorTesauro(f"El tesauro tiene {len(conceptos)} conceptos; el máximo es {MAX_CONCEPTOS}.")
    t = Tesauro(id=uuid.uuid4(), nombre=" ".join(nombre.split())[:200], fuente=(fuente or "").strip()[:500] or None,
                licencia=(licencia or "").strip()[:200] or None, idioma=idioma, uri_esquema=esquema, formato=formato,
                archivo=archivo[:300], huella=huella, conceptos=len(conceptos), importado_por_id=usuario_id)
    db.add(t)
    db.flush()
    ids: dict[str, uuid.UUID] = {}
    etiquetas = []
    for c in conceptos:
        cid = uuid.uuid4()
        ids[c.clave] = cid
        db.add(ConceptoTesauro(id=cid, tesauro_id=t.id, clave=c.clave[:500], uri=(c.uri or None) and c.uri[:500],
                               preferente=c.preferente[:300], notacion=(c.notacion or None) and c.notacion[:60],
                               nota_alcance=c.nota))
        vistas = set()
        for texto, tipo in [(c.preferente, "preferente")] + [(x, "no_preferente") for x in c.no_preferentes] \
                + [(x, "oculta") for x in c.ocultas]:
            n = normalizar(texto)[:300]
            if n and n not in vistas:
                vistas.add(n)
                etiquetas.append(EtiquetaConcepto(concepto_id=cid, etiqueta=texto[:300], normalizada=n, tipo=tipo))
    db.flush()
    db.add_all(etiquetas)
    for c in conceptos:  # término general: después de crear todos
        if c.superior and c.superior in ids and ids[c.superior] != ids[c.clave]:
            db.get(ConceptoTesauro, ids[c.clave]).superior_id = ids[c.superior]
    db.flush()
    registrar(db, modulo="vocabularios", accion="tesauro_importado", usuario_id=usuario_id, entidad_tipo="tesauro",
              entidad_id=t.id, detalle=t.nombre,
              nuevo={"conceptos": t.conceptos, "etiquetas": len(etiquetas), "formato": formato, "huella": huella,
                     "archivo": t.archivo, "licencia": t.licencia})
    return t


def anular(db: Session, t: Tesauro, motivo: str, usuario_id: uuid.UUID) -> Tesauro:
    """Un tesauro anulado deja de ofrecerse para normalizar. Las materias ya
    normalizadas con él conservan su concepto: la descripción no cambia."""
    if t.estado == "anulado":
        raise ErrorTesauro("El tesauro ya está anulado.", 409)
    if not (motivo or "").strip():
        raise ErrorTesauro("Indique el motivo de la anulación.")
    t.estado = "anulado"
    registrar(db, modulo="vocabularios", accion="tesauro_anulado", usuario_id=usuario_id, entidad_tipo="tesauro",
              entidad_id=t.id, detalle=motivo.strip()[:500], anterior={"estado": "vigente"}, nuevo={"estado": "anulado"})
    return t


# --- Normalización -----------------------------------------------------------------------------------


def _vigentes():
    return select(ConceptoTesauro.id).join(Tesauro, Tesauro.id == ConceptoTesauro.tesauro_id).where(
        Tesauro.estado == "vigente", ConceptoTesauro.estado == "vigente")


def normalizar_termino(db: Session, texto: str, limite: int = 6, tesauro_id: uuid.UUID | None = None) -> dict:
    """Qué concepto corresponde a un término escrito.

    Devuelve {"termino", "exacto": concepto o None, "via": preferente |
    no_preferente | None, "parecidos": [...]}. Un término que coincide con
    una etiqueta (sin tildes ni mayúsculas) es exacto: si la etiqueta es no
    preferente, la normalización lleva a la preferente."""
    buscado = normalizar(texto)
    salida = {"termino": texto, "exacto": None, "via": None, "parecidos": []}
    if not buscado:
        return salida
    vigentes = _vigentes()
    if tesauro_id:
        vigentes = vigentes.where(ConceptoTesauro.tesauro_id == tesauro_id)
    exactas = db.execute(select(EtiquetaConcepto).where(EtiquetaConcepto.normalizada == buscado,
                                                        EtiquetaConcepto.concepto_id.in_(vigentes))
                         .order_by(EtiquetaConcepto.tipo.desc())).scalars().all()  # preferente antes que no_preferente
    if exactas:
        e = sorted(exactas, key=lambda x: {"preferente": 0, "no_preferente": 1, "oculta": 2}[x.tipo])[0]
        salida["exacto"] = concepto_out(db, db.get(ConceptoTesauro, e.concepto_id)) | {"etiqueta_coincidente": e.etiqueta}
        salida["via"] = "preferente" if e.tipo == "preferente" else "no_preferente"
    similitud = func.similarity(EtiquetaConcepto.normalizada, buscado)
    filas = db.execute(select(EtiquetaConcepto, similitud.label("s"))
                       .where(EtiquetaConcepto.concepto_id.in_(vigentes),
                              or_(similitud >= UMBRAL_PARECIDO,
                                  EtiquetaConcepto.normalizada.like(f"{buscado.replace('%', '')}%")))
                       .order_by(similitud.desc()).limit(limite * 4)).all()
    vistos = {salida["exacto"]["id"]} if salida["exacto"] else set()
    for e, s in filas:
        if str(e.concepto_id) in vistos:
            continue
        vistos.add(str(e.concepto_id))
        salida["parecidos"].append(concepto_out(db, db.get(ConceptoTesauro, e.concepto_id))
                                   | {"etiqueta_coincidente": e.etiqueta, "similitud": round(float(s), 2)})
        if len(salida["parecidos"]) >= limite:
            break
    return salida


def concepto_out(db: Session, c: ConceptoTesauro) -> dict:
    t = db.get(Tesauro, c.tesauro_id)
    ruta, actual, vistos = [], c, set()
    while actual.superior_id and actual.superior_id not in vistos and len(ruta) < 12:
        vistos.add(actual.superior_id)
        actual = db.get(ConceptoTesauro, actual.superior_id)
        ruta.insert(0, actual.preferente)
    no_pref = db.scalars(select(EtiquetaConcepto.etiqueta).where(EtiquetaConcepto.concepto_id == c.id,
                                                                 EtiquetaConcepto.tipo == "no_preferente")).all()
    return {"id": str(c.id), "preferente": c.preferente, "uri": c.uri, "notacion": c.notacion,
            "nota_alcance": c.nota_alcance, "no_preferentes": list(no_pref), "ruta": ruta,
            "tesauro": {"id": str(t.id), "nombre": t.nombre, "licencia": t.licencia}}


def tesauro_out(t: Tesauro) -> dict:
    return {"id": str(t.id), "nombre": t.nombre, "fuente": t.fuente, "licencia": t.licencia, "idioma": t.idioma,
            "formato": t.formato, "archivo": t.archivo, "uri_esquema": t.uri_esquema, "conceptos": t.conceptos,
            "estado": t.estado, "importado_en": t.importado_en}


# --- Materias del fondo ------------------------------------------------------------------------------


def materia_de_concepto(db: Session, *, fondo_id: uuid.UUID, concepto_id: uuid.UUID,
                        usuario_id: uuid.UUID) -> EntidadVocabulario:
    """La materia del fondo para ese concepto: la que ya existe o una nueva
    con el término preferente. Una sola por concepto y fondo."""
    c = db.get(ConceptoTesauro, concepto_id)
    if c is None or c.estado != "vigente":
        raise ErrorTesauro("Ese concepto no existe en los tesauros.", 404)
    existente = db.scalar(select(EntidadVocabulario).where(EntidadVocabulario.fondo_id == fondo_id,
                                                           EntidadVocabulario.clase == "materia",
                                                           EntidadVocabulario.concepto_tesauro_id == c.id,
                                                           EntidadVocabulario.estado == "activa"))
    if existente is not None:
        return existente
    from app.servicios import vocabulario

    m = vocabulario.crear(db, fondo_id=fondo_id, clase="materia", nombre=c.preferente, subtipo=None, origen="persona",
                          confianza=None, motor=None, usuario_id=usuario_id)
    m.concepto_tesauro_id = c.id
    m.historia = c.nota_alcance
    db.flush()
    return m


def resolver(db: Session, fondo_id: uuid.UUID, texto: str, usuario_id: uuid.UUID) -> tuple[EntidadVocabulario | None, dict]:
    """Normaliza un término escrito al describir: si coincide exactamente con
    un concepto (preferente o no preferente), devuelve la materia de ese
    concepto. Si no, None (lo decide la persona: un parecido o término local)."""
    n = normalizar_termino(db, texto)
    if n["exacto"]:
        return materia_de_concepto(db, fondo_id=fondo_id, concepto_id=uuid.UUID(n["exacto"]["id"]),
                                   usuario_id=usuario_id), n
    return None, n
