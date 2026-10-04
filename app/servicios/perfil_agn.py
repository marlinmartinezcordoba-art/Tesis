"""
Perfil RiC-Col: la adaptación colombiana del AGN aplicada a RICORA.

Fuente: Archivo General de la Nación, «Esquema de Metadatos para las
Entidades Públicas Colombianas», versión 1.4 (31 de octubre de 2025),
adaptación de RiC-CM 1.0. El propio documento se declara instrumento de
referencia no obligatorio y propuesta preliminar (pp. 6-7 y 79).

Tres capas separadas (especificación integral, sección 7):
- «ric»: el modelo internacional (RiC-CM 1.0 y RiC-O 1.1), que manda en
  los códigos y en las propiedades del RDF;
- «agn»: lo que pide la adaptación colombiana (obligatoriedad, reglas,
  vocabularios, correspondencia con el FUID);
- «ricora»: decisiones propias del sistema, que nunca se presentan como
  obligación del AGN.

Cada fila guarda el código tal como lo cita el AGN y la correspondencia
correcta comprobada contra el OWL oficial (una prueba lo exige). Cuando
difieren, la fila lo dice: el esquema del AGN arrastra códigos del borrador
de 2016 de RiC-CM y propiedades que no existen en RiC-O 1.1 (ERRATAS).
"""

from dataclasses import dataclass

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.descripcion import IdentificadorEntidad, Relacion
from app.models.instanciacion import Instanciacion
from app.models.recurso_documental import RecursoDocumental

FUENTES = {
    "agn": {"nombre": "Esquema de Metadatos para las Entidades Públicas Colombianas (AGN)", "version": "1.4",
            "fecha": "2025-10-31", "caracter": "Instrumento técnico de referencia, no obligatorio salvo acuerdo del AGN",
            "url": "https://www.archivogeneral.gov.co/sites/default/files/2026-01/ESQUEMA_DE_METADATOS.pdf"},
    "ric_cm": {"nombre": "Records in Contexts – Conceptual Model (ICA-EGAD)", "version": "1.0", "fecha": "2023-11-30",
               "url": "https://www.ica.org/app/uploads/2023/12/RiC-CM-1.0.pdf"},
    "ric_o": {"nombre": "Records in Contexts – Ontology (ICA-EGAD)", "version": "1.1", "fecha": "2025-05-22",
              "url": "https://www.ica.org/standards/RiC/RiC-O_1-1.html"},
}

CAPAS = {"ric": "Modelo internacional (RiC-CM 1.0 / RiC-O 1.1)", "agn": "Adaptación colombiana (AGN v1.4)",
         "ricora": "Decisión propia de RICORA"}
ESTADOS = {"cubierto": "Cubierto", "parcial": "Parcial", "ausente": "Ausente", "no_adoptado": "No adoptado (decisión)"}
NIVELES = {"basico": "Básico (modelo mínimo viable)", "intermedio": "Intermedio", "avanzado": "Avanzado"}


@dataclass(frozen=True)
class Fila:
    id: str
    origen: str  # tabla o anexo del esquema del AGN
    entidad: str  # entidad según el AGN, con el nombre en español
    elemento: str
    codigo_agn: str  # tal como lo cita el AGN
    obligatoriedad: str  # Obligatorio, Recomendado, Opcional o —
    regla_agn: str
    ric_cm: str  # correspondencia correcta en RiC-CM 1.0
    ric_o: tuple[str, ...]  # términos de RiC-O 1.1 (nombre local); vacío: sin propiedad
    campo: str  # de dónde sale en RICORA
    capa: str
    estado: str
    nivel: str
    nota: str = ""
    codigo_agn_correcto: bool = True


F = Fila
CATALOGO: tuple[Fila, ...] = (
    # --- Tabla 3. Record (RiC-E04) - Documento de archivo -------------------------------------------------
    F("AGN-REC-01", "Tabla 3", "Documento", "Identificador único", "ric:identifier", "Obligatorio",
      "Formato CO-{ENTIDAD}-{SERIE}-{NÚMERO}, máximo 50 caracteres",
      "RiC-A22 Identifier", ("identifier",), "codigo_referencia (código del CCD, compuesto desde el nivel superior) "
      "y URI persistente urn:uuid", "agn", "parcial", "basico",
      "El código del CCD sí es obligatorio (Acuerdo 001 de 2024); el formato CO-… del AGN es un ejemplo de un "
      "instrumento no obligatorio y no se impone. La medición de calidad avisa si pasa de 50 caracteres."),
    F("AGN-REC-02", "Tabla 3", "Documento", "Título formal", "ric:name", "Obligatorio",
      "Máximo 255 caracteres, sin abreviaciones", "RiC-A28 Name (rico:title)", ("title",), "titulo",
      "agn", "cubierto", "basico", "RICORA admite hasta 300 caracteres (ISAD-G no limita); la medición avisa si "
      "pasa de 255. rico:name existe; rico:title es su especialización para documentos."),
    F("AGN-REC-03", "Tabla 3", "Documento", "Descripción del contenido", "ric:description", "Recomendado",
      "Entre 50 y 500 caracteres", "RiC-A38 Scope and Content", ("scopeAndContent",), "alcance_contenido",
      "agn", "cubierto", "intermedio", "«ric:description» no existe en RiC-O 1.1.", False),
    F("AGN-REC-04", "Tabla 3", "Documento", "Tipo de documento", "ric:hasDocumentaryFormType", "Obligatorio",
      "Vocabulario de la TRD más el vocabulario COAR", "RiC-A17 Documentary Form Type",
      ("hasDocumentaryFormType", "DocumentaryFormType"), "forma_documental_id (vocabulario del fondo, SKOS)",
      "agn", "parcial", "basico",
      "El AGN lo hace corresponder con ISAD(G) 3.1.4 «Nivel de descripción», que es otra cosa: el nivel va en "
      "rico:hasRecordSetType o en la clase. COAR no se adopta (ver decisiones)."),
    F("AGN-REC-05", "Tabla 3", "Documento", "Dimensiones físicas o lógicas", "ric:hasExtent", "Obligatorio",
      "Especificar la unidad de medida (folios, MB…)", "RiC-A35 Record Resource Extent",
      ("hasExtent", "RecordResourceExtent", "recordResourceExtent"), "folios (unidad: folios) y tamaño de cada "
      "instanciación (bytes, páginas)", "agn", "cubierto", "basico"),
    F("AGN-REC-06", "Tabla 3", "Documento", "Soporte del documento", "ric:hasCarrierType", "Obligatorio",
      "Tipo MIME para digitales, descripción normalizada para físicos", "RiC-A05 Carrier Type (de la Instantiation)",
      ("hasCarrierType", "CarrierType"), "soporte (FUID) y, en cada instanciación, soporte del original o tipo "
      "MIME y formato PRONOM", "agn", "cubierto", "basico",
      "En RiC el soporte es de la instanciación, no del documento: RICORA lo exporta en la Instantiation."),
    F("AGN-REC-07", "Tabla 3", "Documento", "Idioma del documento", "ric:language", "Obligatorio",
      "ISO 639-2, «spa» por defecto", "RiC-A25 Language", ("hasOrHadLanguage",), "idiomas (ISO 639-3)",
      "agn", "cubierto", "intermedio",
      "«spa» es el mismo código en ISO 639-2/T y en 639-3; RICORA usa 639-3 porque distingue lenguas indígenas "
      "colombianas que 639-2 agrupa. No se pone «spa» por defecto: el idioma lo confirma una persona.", False),
    # --- Tabla 4 y anexo 3. Instantiation (RiC-E06) ------------------------------------------------------
    F("AGN-INS-01", "Tabla 4", "Instanciación", "Soporte", "ric:carrierType · RiC-P13", "Obligatorio",
      "Vocabulario controlado: papel, digital, microfilm, audiovisual", "RiC-A05 Carrier Type",
      ("hasCarrierType",), "instanciaciones.soporte (vocabulario cerrado) o formato del archivo", "agn",
      "cubierto", "basico", "", False),
    F("AGN-INS-02", "Tabla 4", "Instanciación", "Localización física (signatura topográfica)",
      "ric:physicalLocation · RiC-P35", "Obligatorio", "Incluir signatura topográfica, depósito, estantería",
      "RiC-A27 Physical Location (en RiC-O es atributo del Lugar, no de la instanciación)", (),
      "ubicacion_fisica + deposito, estante, entrepano; caja y carpeta en el FUID", "agn", "cubierto",
      "intermedio", "Se guarda y sale en la ficha, el inventario y la ficha ISAD(G); no va al RDF porque RiC-O 1.1 "
      "no tiene propiedad de dato para la ubicación de una instanciación.", False),
    F("AGN-INS-03", "Tabla 4", "Instanciación", "Extensión", "ric:extent · RiC-P32", "Obligatorio",
      "Unidades normalizadas (MB, GB, cajas, legajos)", "RiC-A23 Instantiation Extent",
      ("instantiationExtent", "hasExtent"), "tamano_bytes, paginas (con su unidad en rico:Extent)", "agn",
      "cubierto", "basico", "", False),
    F("AGN-INS-04", "Tabla 4", "Instanciación", "Características técnicas", "ric:technicalCharacteristics · RiC-P28",
      "Recomendado", "Formato, resolución, compresión, hash SHA-256; PUID de PRONOM", "RiC-A41 Technical "
      "Characteristics", ("technicalCharacteristics",), "formato PRONOM (Siegfried), MIME, versión, páginas",
      "agn", "cubierto", "avanzado", "El hash va en la nota de autenticidad (RiC-A03), no aquí.", False),
    F("AGN-INS-05", "Tabla 4", "Instanciación", "Evidencias de autenticidad", "ric:hasAuthentication · RiC-P25",
      "Opcional", "Firmas digitales (X.509), sellos, marcas de agua", "RiC-A03 Authenticity Note",
      ("authenticityNote",), "huella SHA-256, verificaciones de fijeza y validación de formato (veraPDF, JHOVE)",
      "agn", "parcial", "avanzado", "No se verifican firmas digitales X.509 (Ley 527 de 1999).", False),
    F("AGN-INS-06", "Tabla 4", "Instanciación", "Estado de conservación", "ric:conditionOfRecord · RiC-A17",
      "Recomendado", "Bueno, Regular, Malo, Restaurado; documentar daños", "RiC-A31 Physical Characteristics",
      ("physicalCharacteristicsNote",), "estado_conservacion (vocabulario del AGN) + caracteristicas_fisicas",
      "agn", "cubierto", "avanzado", "RiC-A17 es Documentary Form Type, no el estado de conservación.", False),
    F("AGN-INS-07", "Tabla 4", "Instanciación", "Técnica de producción", "ric:hasProductionTechnique · RiC-P29",
      "Opcional", "Manuscrito, mecanografiado, impreso, nacido digital, digitalizado", "RiC-A33 Production Technique",
      ("productionTechnique",), "solo «digitalizado»: la derivación desde el original físico (RiC-R014)", "agn",
      "parcial", "avanzado", "Manuscrito, mecanografiado o impreso no se registran.", False),
    F("AGN-INS-08", "Tabla 4", "Instanciación", "Identificador", "ric:hasOrHadIdentifier · RiC-P1", "Obligatorio",
      "URI persistente, código de barras, signatura", "RiC-A22 Identifier", ("identifier",),
      "urn:uuid y URI dereferenciable pública", "agn", "cubierto", "basico", "", False),
    F("AGN-INS-09", "Tabla 4", "Instanciación", "Vínculo con el documento", "ric:isOrWasCarrierOf · ric:instantiates "
      "· RiC-R66", "Obligatorio", "Vincular con el Record Resource", "RiC-R025 has or had instantiation",
      ("hasOrHadInstantiation", "isOrWasInstantiationOf"), "relación has_or_had_instantiation", "agn", "cubierto",
      "basico", "RiC-R066 es «is or was enforced by».", False),
    F("AGN-INS-10", "Tabla 4", "Instanciación", "Condiciones de acceso (el AGN la llama «accesibilidad»)",
      "ric:hasAccessibility · RiC-P38", "Obligatorio", "Público, Restringido, Confidencial; con fundamento legal",
      "RiC-A08 Conditions of Access", ("conditionsOfAccess",), "declaración de derechos Ley 1712 (pública, "
      "clasificada, reservada con plazo) heredable", "agn", "cubierto", "basico",
      "RICORA usa las categorías de la Ley 1712 (pública, clasificada, reservada); «restringido» y «confidencial» "
      "no son categorías legales.", False),
    # --- Tabla 5. Agent (RiC-E07) -----------------------------------------------------------------------
    F("AGN-AGE-01", "Tabla 5", "Agente", "Nombre autorizado", "ric:name", "Obligatorio",
      "Forma completa; consultar autoridades nacionales", "RiC-A28 Name (rico:AgentName)",
      ("name", "hasOrHadAgentName", "AgentName"), "nombre y formas del nombre con vigencia (ISAAR-CPF)", "agn",
      "cubierto", "basico"),
    F("AGN-AGE-02", "Tabla 5", "Agente", "Tipo de agente", "ric:agentType", "Obligatorio",
      "Person, Corporate Body, Family, Group", "Clases RiC-E08 a E13 (no es un atributo)",
      ("Person", "CorporateBody", "Family", "Group", "Position", "Mechanism"), "subtipo del agente → clase RiC-O",
      "agn", "cubierto", "basico", "«ric:agentType» no existe en RiC-O 1.1; el tipo es la clase.", False),
    F("AGN-AGE-03", "Tabla 5", "Agente", "Identificador del agente", "ric:identifier", "Recomendado",
      "ORCID para personas, ROR para instituciones", "RiC-A22 Identifier", ("identifier",),
      "identificadores con esquema: interno, VIAF, Wikidata, ISNI, LCNAF, ORCID, ROR (owl:sameAs)", "agn",
      "cubierto", "avanzado", "ORCID y ROR se validan con su dígito de control."),
    F("AGN-AGE-04", "Tabla 5", "Agente", "Historia institucional o biográfica", "ric:description", "Recomendado",
      "Resumen de actividades relevantes", "RiC-A21 History", ("history",), "historia (ficha de autoridad)",
      "agn", "cubierto", "intermedio", "", False),
    F("AGN-AGE-05", "Tabla 5", "Agente", "Lugares", "ric:hasPlace", "Opcional", "Coordenadas cuando sea relevante",
      "RiC-R075 has or had location (Lugar, E22)", ("hasOrHadLocation", "Place"), "relación del agente con un lugar "
      "del vocabulario, con coordenadas", "agn", "cubierto", "avanzado", "", False),
    # --- Tablas 6 y 7. Actividad y función ----------------------------------------------------------------
    F("AGN-ACT-01", "Tabla 6", "Actividad", "Denominación, descripción y fechas", "ric:name · ric:description · "
      "ric:hasBeginningDate · ric:hasEndDate", "Obligatorio", "Terminología del mapa de procesos; ISO 8601",
      "RiC-E15 Activity; RiC-A28; RiC-R069 / RiC-R071", ("Activity", "name", "hasBeginningDate", "hasEndDate"),
      "actividad del vocabulario con periodo EDTF, ejercida por un agente y regulada por un mandato", "agn",
      "cubierto", "intermedio", "«ric:description» no existe; las fechas en RiC son entidades Fecha (E18).", False),
    F("AGN-FUN-01", "Tabla 7", "Función", "Función con su marco normativo", "Function (RiC-E08) · ric:hasLegalStatus",
      "Obligatorio", "Terminología del manual de funciones; citar las normas", "RiC-CM 1.0 no tiene la entidad "
      "Función: se representa con Activity (E15) y Activity Type (A02); el marco normativo con Mandate (E17, "
      "RiC-R067)", ("ActivityType", "hasActivityType", "Mandate", "authorizes"),
      "tipo de actividad (función de la TRD) con jerarquía SKOS y mandato que la autoriza", "agn", "cubierto",
      "intermedio", "E08 es Persona. hasLegalStatus (A26) es el estatuto jurídico de un agente, no el marco de "
      "una función.", False),
    # --- Tablas 8 y 9. Persona y entidad corporativa -------------------------------------------------------
    F("AGN-PER-01", "Tabla 8", "Persona", "Nombres, fechas de existencia y ocupación", "Person (RiC-E08) · "
      "ric:hasAppellation · ric:hasBirthDate · ric:hasOccupation", "Obligatorio",
      "Apellidos, Nombres; fecha de nacimiento solo si es relevante (Ley 1581)", "RiC-E08 Person; RiC-A28; "
      "RiC-R070; RiC-A30 Occupation Type", ("Person", "hasOrHadAgentName", "hasBirthDate",
                                           "hasOrHadOccupationOfType"),
      "persona con formas del nombre, existencia EDTF y cargos (RiC-E12)", "agn", "parcial", "intermedio",
      "La ocupación se expresa como cargo ocupado (Position), no como tipo de ocupación.", False),
    F("AGN-COR-01", "Tabla 9", "Entidad corporativa", "Denominación, fechas, mandato y estatuto",
      "Corporate Body (RiC-E11) · ric:hasOtherName · ric:hasMandate", "Obligatorio",
      "Denominación oficial según acto de creación; competencias legales", "RiC-E11 Corporate Body; RiC-A28; "
      "RiC-A26 Legal Status; RiC-R067 authorizes (desde el Mandato)",
      ("CorporateBody", "hasOrHadAgentName", "hasOrHadLegalStatus", "authorizes"),
      "entidad corporativa con formas del nombre, estatuto jurídico, hitos y mandatos", "agn", "cubierto",
      "intermedio", "«ric:hasOtherName» y «ric:hasMandate» no existen en RiC-O 1.1.", False),
    # --- Tabla 10. Relaciones entre documentos -------------------------------------------------------------
    F("AGN-REL-01", "Tabla 10", "Relación", "Jerárquica (forma parte / contiene)", "ric:isPartOf · ric:includes",
      "Obligatorio", "Para documentos que forman parte de expedientes", "RiC-R024 includes or included",
      ("includesOrIncluded", "isOrWasIncludedIn"), "inclusión orgánica y adicionales", "agn", "cubierto", "basico",
      "", False),
    F("AGN-REL-02", "Tabla 10", "Relación", "Asociativa", "ric:isAssociatedWith", "Opcional",
      "Documentos temáticamente relacionados", "RiC-R001 is related to", ("isRelatedTo",),
      "relacionadas y secuencia (RiC-R008)", "agn", "cubierto", "intermedio", "", False),
    F("AGN-REL-03", "Tabla 10", "Relación", "Derivación (copias, extractos)", "ric:derivedFrom", "Opcional",
      "Para copias, extractos, transcripciones", "RiC-R012 has copy (entre documentos); RiC-R014 derivación "
      "(entre instanciaciones)", ("hasCopy", "hasOrHadDerivedInstantiation"),
      "derivación y migración entre instanciaciones; la copia entre documentos descritos no", "agn", "parcial",
      "avanzado", "", False),
    # --- Anexo 2. FUID ↔ RiC-Col ---------------------------------------------------------------------------
    F("AGN-FUID-01", "Anexo 2", "FUID", "Código de serie o subserie", "RiC-P001: Identifier (no P29)", "—",
      "Persistencia en el SGDEA", "RiC-A22 Identifier", ("identifier",), "codigo_referencia (CCD)", "agn",
      "cubierto", "basico", "", False),
    F("AGN-FUID-02", "Anexo 2", "FUID", "Nombre de la serie o subserie", "RiC-E02: Record Set + RiC-P005: Name",
      "—", "", "RiC-E03 Record Set; RiC-A28", ("RecordSet", "title"), "titulo de la agrupación", "agn",
      "cubierto", "basico", "Record Set es E03; E02 es Record Resource.", False),
    F("AGN-FUID-03", "Anexo 2", "FUID", "Dependencia productora", "RiC-E04: Agent + RiC-R01: isCreatedBy", "—",
      "Incluir NIT o código institucional", "RiC-E07 Agent; RiC-R027 has creator", ("hasCreator",),
      "relación has_creator al agente (con sus identificadores)", "agn", "cubierto", "basico",
      "E04 es Record y R001 es «is related to».", False),
    F("AGN-FUID-04", "Anexo 2", "FUID", "Fechas extremas", "RiC-P004: Date / RiC-P005: Date Range", "—",
      "ISO 8601", "RiC-E18 Date con RiC-A29 Normalized Value y RiC-A42 Date Type", ("Date", "normalizedDateValue"),
      "fechas_extremas_edtf (EDTF, compatible con ISO 8601-2) y nodos Fecha", "agn", "cubierto", "basico",
      "En RiC-CM 1.0 la fecha es una entidad, no un atributo.", False),
    F("AGN-FUID-05", "Anexo 2", "FUID", "Soporte", "RiC-P013: Carrier Type", "—", "Papel, Digital, Mixto",
      "RiC-A05 Carrier Type", ("hasCarrierType",), "soporte", "agn", "cubierto", "basico", "", False),
    F("AGN-FUID-06", "Anexo 2", "FUID", "Acceso y restricción", "RiC-P018 / RiC-P019", "—",
      "Vincular con la Ley 1712 y la Ley 1581", "RiC-A08 Conditions of Access; RiC-A09 Conditions of Use",
      ("conditionsOfAccess", "conditionsOfUse"), "clasificación Ley 1712 + condiciones de uso + datos personales",
      "agn", "cubierto", "basico", "", False),
    F("AGN-FUID-07", "Anexo 2", "FUID", "Frecuencia de uso", "RiC-P031: Record State (adaptación local)", "—", "",
      "Sin equivalente en RiC-CM: Record State (A39) es el estado de producción (original, copia, borrador)", (),
      "frecuencia_consulta (columna del FUID)", "agn", "cubierto", "intermedio",
      "Se usa en el inventario; no se exporta a RDF para no desvirtuar rico:RecordState.", False),
    F("AGN-FUID-08", "Anexo 2", "FUID", "Valor documental y disposición final", "RiC-R07: Has Appraisal Decision + "
      "RiC-P025: Appraisal Type", "—", "Compatible con la TVD", "RiC-E16 Rule (regla de retención) con RiC-R063 "
      "regulates or regulated", ("Rule", "regulatesOrRegulated"), "regla de retención de la TRD heredada + "
      "valoracion", "agn", "cubierto", "intermedio", "hasAppraisalDecision no existe en RiC-O 1.1.", False),
    F("AGN-FUID-09", "Anexo 2", "FUID", "Soporte de preservación", "RiC-E06: Instantiation + RiC-P035: Storage "
      "Location", "—", "", "RiC-E06 Instantiation; la ubicación del almacenamiento es metadato de preservación "
      "(PREMIS storage), no de RiC", ("Instantiation",), "almacén y segunda copia, en el AIP con PREMIS", "agn",
      "cubierto", "avanzado", "", False),
    F("AGN-FUID-10", "Anexo 2", "FUID", "Volumen documental", "RiC-P002: Extent", "—", "Indicar la unidad",
      "RiC-A35 Record Resource Extent", ("recordResourceExtent", "RecordResourceExtent"), "folios, caja, carpeta",
      "agn", "cubierto", "basico", "", False),
    F("AGN-FUID-11", "Anexo 2", "FUID", "Procedimiento o proceso asociado", "RiC-E12: Function + RiC-R12: Is Function "
      "Of", "—", "Vincular con el mapa de procesos", "RiC-E15 Activity / RiC-A02 Activity Type; RiC-R033 documents",
      ("documents", "ActivityType"), "actividad documentada y tipo de actividad (función)", "agn", "cubierto",
      "intermedio", "E12 es Puesto (Position).", False),
    F("AGN-FUID-12", "Anexo 2", "FUID", "Serie antecedente o derivada", "RiC-R23: Precedes / RiC-R24: Follows", "—",
      "", "RiC-R008 precedes or preceded", ("precedesOrPreceded", "followsOrFollowed"), "secuencia",
      "agn", "cubierto", "intermedio", "R023 es vínculo genético y R024 es inclusión.", False),
    F("AGN-FUID-13", "Anexo 2", "FUID", "Productor antecedente o sucesor", "RiC-R05 / RiC-R06", "—",
      "Fusiones o cambios institucionales", "RiC-R016 relación temporal entre agentes (has successor)",
      ("hasSuccessor", "AgentTemporalRelation"), "relación has_successor con fecha", "agn", "cubierto", "intermedio",
      "R005 es subdivisión de grupo y R006 es subevento.", False),
    F("AGN-FUID-14", "Anexo 2", "FUID", "Palabras clave o temas", "RiC-P003: Subject", "—",
      "Vocabularios del AGN o tesauros sectoriales", "RiC-R019 has or had subject", ("hasOrHadSubject",),
      "agentes, lugares y actividades mencionados (vocabulario controlado)", "agn", "cubierto", "intermedio",
      "", False),
    F("AGN-FUID-15", "Anexo 2", "FUID", "Observaciones", "RiC-P006: Description", "—", "",
      "RiC-A43 General Description", ("generalDescription",), "nota", "agn", "cubierto", "intermedio", "", False),
    F("AGN-FUID-16", "Anexo 2", "FUID", "Ubicación física o digital", "RiC-P035: Storage Location", "—",
      "URI o ruta del SGDEA/OAIS", "Sin propiedad de dato en RiC-O 1.1 para una instanciación", (),
      "caja, carpeta, signatura topográfica; ruta interna del almacén (no se publica)", "agn", "parcial",
      "intermedio", "La ruta del almacén no se publica por seguridad.", False),
    F("AGN-FUID-17", "Anexo 2", "FUID", "Responsable de custodia", "RiC-E04: Agent + RiC-R09: Is Custodian Of", "—",
      "Coincidir con el responsable del plan de preservación", "RiC-R039 has or had holder",
      ("hasOrHadHolder", "RecordResourceHoldingRelation"), "custodio del documento y de la instanciación, con fechas",
      "agn", "cubierto", "intermedio", "R009 es precedencia en el tiempo.", False),
    # --- Lineamientos transversales del AGN ----------------------------------------------------------------
    F("AGN-TRA-01", "pp. 43 y 50", "Transversal", "Datos personales y sensibles", "—", "Obligatorio",
      "Ley 1581 de 2012: minimización y control de acceso", "Sin equivalente en RiC", (),
      "datos_personales (no contiene, personales, sensibles, de menores)", "agn", "cubierto", "intermedio",
      "No se exporta a RDF. La medición avisa si un documento con datos sensibles o de menores queda público."),
    F("AGN-TRA-02", "pp. 18-19", "Transversal", "Accesibilidad", "—", "Recomendado",
      "Ley 1680 de 2013; WCAG 2.1; versiones accesibles y descripción alternativa", "Sin equivalente en RiC", (),
      "nota_accesibilidad (sale en la ficha pública)", "agn", "cubierto", "avanzado"),
    F("AGN-TRA-03", "p. 41", "Transversal", "Metadatos de preservación", "PREMIS", "Recomendado",
      "Eventos, fijeza, formato, entorno", "Fuera de RiC: PREMIS 3.0", (), "PREMIS 3.0 en el AIP, validado contra "
      "el XSD oficial", "agn", "cubierto", "avanzado"),
    F("AGN-TRA-04", "p. 73", "Transversal", "Validación de metadatos", "XSD o SHACL", "Recomendado",
      "Controles de calidad automáticos", "RiC-O 1.1 con formas SHACL", (), "validación SHACL de cada exportación "
      "y comprobación contra el OWL oficial", "agn", "cubierto", "avanzado"),
    F("AGN-TRA-05", "pp. 11-12", "Transversal", "Niveles de madurez e indicadores de calidad", "—", "Recomendado",
      "Modelo mínimo viable; KPI de calidad de metadatos", "—", (), "medición de calidad por documento y por fondo "
      "según este perfil", "agn", "cubierto", "intermedio"),
    F("AGN-TRA-06", "p. 44", "Transversal", "Perfiles de aplicación por tipología", "—", "Recomendado",
      "Plantillas para expedientes contractuales, judiciales, pensionales", "—", (),
      "no existen plantillas por tipología", "agn", "ausente", "avanzado"),
    F("AGN-TRA-07", "p. 38", "Transversal", "Firma digital", "—", "Opcional",
      "Ley 527 de 1999; Decreto 2364 de 2012", "RiC-A03 Authenticity Note", ("authenticityNote",),
      "no se valida la firma; solo fijeza y formato", "agn", "ausente", "avanzado"),
    F("AGN-TRA-08", "Tabla 3", "Transversal", "Vocabulario COAR de tipos de recurso", "—", "Obligatorio",
      "Junto al vocabulario de la TRD", "—", (), "no se adopta", "agn", "no_adoptado", "avanzado",
      "COAR describe tipos de recursos de repositorios académicos (artículo, tesis), no tipos documentales de "
      "archivo; el tipo documental sale de la TRD del fondo."),
    # --- Decisiones propias de RICORA -----------------------------------------------------------------------
    F("RIC-EXT-01", "—", "Relación", "Certeza, fuente y estado de cada relación", "—", "—",
      "", "RiC-O 1.1: rico:Relation con relationCertainty, relationSource, relationState",
      ("relationCertainty", "relationSource", "relationState"), "nodo de relación con certeza del motor y fuente",
      "ricora", "cubierto", "avanzado", "Es RiC-O oficial; la decisión propia es exponer la certeza de la IA."),
    F("RIC-EXT-02", "—", "Mecanismo", "Programas y motor de IA como agentes con versión", "—", "—", "",
      "RiC-E13 Mechanism; RiC-A41", ("Mechanism", "technicalCharacteristics"),
      "vocabulario de mecanismos de solo lectura (OCR, PRONOM, motor)", "ricora", "cubierto", "avanzado"),
    F("RIC-EXT-03", "—", "Documento", "Origen y confianza de cada dato (motor o persona)", "—", "—", "",
      "Sin equivalente en RiC", (), "origen_* y confianza_* en la descripción; la IA propone y la persona decide",
      "ricora", "cubierto", "intermedio"),
)

# Erratas del esquema del AGN v1.4 (comprobadas contra el OWL de RiC-O 1.1 y
# RiC-CM 1.0). Se usan la corrección, nunca el código del AGN.
ERRATAS = (
    {"pagina": "45 y 80", "dice": "RiC-CM v1.0 «publicado por el ICA en 2021»",
     "correcto": "RiC-CM 1.0 se publicó el 30 de noviembre de 2023"},
    {"pagina": "91", "dice": "«RiC-CM está en versión 0.2 (octubre 2021)»",
     "correcto": "Contradice la portada (adaptación de RiC-CM 1.0); la 0.2 es un borrador"},
    {"pagina": "33-34, 45-49, 74",
     "dice": "Agente como E03, E04 y E07; Documento como E02, E03 y E04; Evento como E04, E10 y E15; Norma como E07 "
             "y E13; Lugar como E02 y E07; Fecha como E01",
     "correcto": "E04 Record, E03 Record Set, E07 Agent, E14 Event, E15 Activity, E16 Rule, E18 Date, E22 Place"},
    {"pagina": "48-49, 53", "dice": "«Function (RiC-E08)» y «(RiC-E12)»",
     "correcto": "RiC-CM 1.0 no tiene entidad Función (Activity E15 + Activity Type A02); E08 es Person, E12 es "
                 "Position"},
    {"pagina": "20-24, 36-37, 83-89", "dice": "Atributos RiC-P01 a RiC-P38",
     "correcto": "Los códigos P son del borrador de 2016; en RiC-CM 1.0 los atributos son RiC-A01 a RiC-A45"},
    {"pagina": "20, 24, 36-37, 52, 83-87",
     "dice": "R002 «is created by», R026 «has transfer date», R032-R035 creador/destinatario/sujeto/custodio, "
             "R126, y códigos de dos cifras R01, R07, R09, R12",
     "correcto": "Son 86 relaciones (R001-R086). Creador R027, destinatario R032, sujeto R019, custodio R039; "
                 "R002 es todo-parte, R026 procedencia orgánica; R126 no existe"},
    {"pagina": "45-50, 89",
     "dice": "Propiedades ric:description, ric:language, ric:carrierType, ric:physicalLocation, ric:instantiates, "
             "ric:isPartOf, ric:hasAccessibility, ric:conditionOfRecord, ric:agentType, ric:storageLocation…",
     "correcto": "No existen en RiC-O 1.1 y el prefijo es rico:. Ver la columna RiC-O de este perfil"},
    {"pagina": "83-86", "dice": "Record State (RiC-P031) para «Vigente», «frecuencia de uso» o «Versión 2.0»",
     "correcto": "Record State (RiC-A39) es el estado de producción: original, copia, borrador"},
    {"pagina": "35 y 41", "dice": "ISO 23081 «adaptada como NTC 4095:2013»",
     "correcto": "La NTC 4095 es la norma general de descripción archivística (adaptación de ISAD(G))"},
    {"pagina": "45", "dice": "Tipo de documento ↔ ISAD(G) 3.1.4 «Nivel de descripción»",
     "correcto": "El nivel de descripción no es el tipo documental"},
    {"pagina": "33", "dice": "La tabla 2 omite Instantiation «por su complejidad»",
     "correcto": "La misma tabla la incluye y el anexo 3 la desarrolla"},
    {"pagina": "56-70", "dice": "Ejemplos XML con espacio de nombres example.org",
     "correcto": "No son RiC ni XML bien formado; RICORA exporta RiC-O 1.1 (Turtle y JSON-LD) validado con SHACL"},
    {"pagina": "86-87", "dice": "«tu adaptación es válida», «Perfecto», «Corregido (no P4)»",
     "correcto": "Comentarios de revisión que quedaron en el texto publicado"},
)


def catalogo() -> dict:
    filas = [f.__dict__ | {"ric_o": list(f.ric_o)} for f in CATALOGO]
    resumen = {e: sum(1 for f in CATALOGO if f.estado == e) for e in ESTADOS}
    return {"fuentes": FUENTES, "capas": CAPAS, "estados": ESTADOS, "niveles": NIVELES, "filas": filas,
            "resumen": resumen, "erratas": list(ERRATAS),
            "codigos_agn_incorrectos": sum(1 for f in CATALOGO if not f.codigo_agn_correcto)}


# --- Calidad de metadatos según el perfil (KPI del AGN, fase 3) --------------------------------------

LIMITE_TITULO = 255
LIMITE_IDENTIFICADOR = 50
NIVELES_RECORD = ("unidad_documental", "parte_documental")


def _instancias(db: Session, recurso: RecursoDocumental) -> list[Instanciacion]:
    return list(db.scalars(select(Instanciacion).join(Relacion, Relacion.destino_id == Instanciacion.id).where(
        Relacion.origen_id == recurso.id, Relacion.codigo_ric == "has_or_had_instantiation",
        Relacion.estado == "vigente")).all())


def _acceso(db: Session, recurso: RecursoDocumental) -> str | None:
    from app.servicios import derechos

    d = derechos._vigente(db, recurso.id) or derechos.declaracion_de_recurso(db, recurso)
    return d.acceso if d is not None else None


def calidad(db: Session, recurso: RecursoDocumental) -> dict:
    """Cada criterio del perfil con su nivel, si aplica a este nivel de
    descripción y si se cumple. Avisos: reglas del AGN que no son error."""
    from app.servicios import isadg, retencion

    valores = {e["elemento"]: e["valor"] for e in isadg.ficha(db, recurso)}
    instancias = _instancias(db, recurso)
    fisicos = [i for i in instancias if i.estado == "registro_fisico"]
    digitales = [i for i in instancias if i.estado != "registro_fisico"]
    es_record = recurso.nivel in NIVELES_RECORD
    acceso = _acceso(db, recurso)
    productores = db.scalars(select(Relacion.destino_id).where(
        Relacion.origen_id == recurso.id, Relacion.codigo_ric == "has_creator", Relacion.estado == "vigente")).all()
    if not productores:
        # El productor se declara una vez en el nivel donde está (ISAD-G,
        # descripción multinivel) y los niveles inferiores lo heredan.
        superior, vistos = recurso.incluido_en_id, set()
        while superior and superior not in vistos and not productores:
            vistos.add(superior)
            productores = db.scalars(select(Relacion.destino_id).where(
                Relacion.origen_id == superior, Relacion.codigo_ric == "has_creator",
                Relacion.estado == "vigente")).all()
            sup = db.get(RecursoDocumental, superior)
            superior = sup.incluido_en_id if sup else None
    externos = bool(productores) and db.scalar(select(IdentificadorEntidad.id).where(
        IdentificadorEntidad.entidad_id.in_(productores), IdentificadorEntidad.estado == "vigente",
        IdentificadorEntidad.esquema != "interno").limit(1)) is not None

    criterios = [
        ("identificador", "AGN-REC-01", "basico", True, bool(recurso.codigo_referencia)),
        ("titulo", "AGN-REC-02", "basico", True, bool(recurso.titulo)),
        ("fechas", "AGN-FUID-04", "basico", True, bool(valores.get("3.1.3"))),
        ("productor", "AGN-FUID-03", "basico", True, bool(productores)),
        ("tipo_documental", "AGN-REC-04", "basico", es_record, recurso.forma_documental_id is not None),
        ("extension", "AGN-REC-05", "basico", True,
         recurso.folios is not None or any(i.tamano_bytes for i in digitales)),
        ("soporte", "AGN-REC-06", "basico", True,
         bool(recurso.soporte) or any(i.soporte for i in fisicos) or any(i.formato_mime for i in digitales)),
        ("acceso", "AGN-INS-10", "basico", True, acceso is not None or bool(recurso.condiciones_acceso)),
        ("instanciacion", "AGN-INS-09", "basico", es_record, bool(instancias)),
        ("idioma", "AGN-REC-07", "intermedio", True, bool(recurso.idiomas)),
        ("alcance", "AGN-REC-03", "intermedio", True, bool(recurso.alcance_contenido)),
        ("datos_personales", "AGN-TRA-01", "intermedio", True, recurso.datos_personales is not None),
        ("disposicion", "AGN-FUID-08", "intermedio", recurso.nivel in ("serie", "subserie", "expediente"),
         retencion.retencion_de(db, recurso) is not None),
        ("ubicacion_fisica", "AGN-INS-02", "intermedio", bool(fisicos),
         all(i.ubicacion_fisica or i.deposito or i.estante for i in fisicos)),
        ("estado_conservacion", "AGN-INS-06", "avanzado", bool(fisicos), all(i.estado_conservacion for i in fisicos)),
        ("formato_identificado", "AGN-INS-04", "avanzado", bool(digitales), all(i.formato_puid for i in digitales)),
        ("integridad_verificada", "AGN-INS-05", "avanzado", bool(digitales),
         all(i.estado_integridad == "integra" for i in digitales)),
        ("identificador_externo_productor", "AGN-AGE-03", "avanzado", bool(productores), externos),
        ("accesibilidad", "AGN-TRA-02", "avanzado", True, bool(recurso.nota_accesibilidad)),
    ]
    nombres = {f.id: f.elemento for f in CATALOGO}
    salida = [{"clave": c, "fila": fila, "elemento": nombres[fila], "nivel": nivel, "aplica": aplica,
               "cumple": cumple if aplica else None} for c, fila, nivel, aplica, cumple in criterios]
    avisos = []
    if recurso.titulo and len(recurso.titulo) > LIMITE_TITULO:
        avisos.append(f"El título tiene {len(recurso.titulo)} caracteres; el AGN recomienda hasta {LIMITE_TITULO}.")
    if recurso.codigo_referencia and len(recurso.codigo_referencia) > LIMITE_IDENTIFICADOR:
        avisos.append(f"El identificador pasa de {LIMITE_IDENTIFICADOR} caracteres (AGN, tabla 3).")
    if recurso.alcance_contenido and len(recurso.alcance_contenido) < 50:
        avisos.append("El alcance y contenido tiene menos de 50 caracteres (AGN recomienda entre 50 y 500).")
    incoherencias = []
    if recurso.datos_personales in ("sensibles", "menores") and acceso in (None, "publico"):
        incoherencias.append(
            "Contiene datos " + ("sensibles" if recurso.datos_personales == "sensibles" else "de niños, niñas o "
                                 "adolescentes") + " y su acceso es público: clasifíquelo (Ley 1712, art. 18, literal "
                                                   "a) o anonimice la versión de consulta (Ley 1581, arts. 5 a 7).")
    por_nivel = {}
    for n in NIVELES:
        aplican = [c for c in salida if c["nivel"] == n and c["aplica"]]
        por_nivel[n] = {"aplican": len(aplican), "cumplen": sum(1 for c in aplican if c["cumple"]),
                        "porcentaje": round(100 * sum(1 for c in aplican if c["cumple"]) / len(aplican))
                        if aplican else None}
    aplican = [c for c in salida if c["aplica"]]
    return {"id": str(recurso.id), "titulo": recurso.titulo, "nivel": recurso.nivel, "criterios": salida,
            "por_nivel": por_nivel, "avisos": avisos, "incoherencias": incoherencias,
            "porcentaje": round(100 * sum(1 for c in aplican if c["cumple"]) / len(aplican)) if aplican else None,
            "nivel_alcanzado": _nivel_alcanzado(por_nivel)}


def _nivel_alcanzado(por_nivel: dict) -> str | None:
    """El nivel más alto con todos sus criterios, y los de los niveles
    anteriores, cumplidos (el modelo mínimo viable es el básico)."""
    alcanzado = None
    for n in NIVELES:
        if por_nivel[n]["porcentaje"] not in (None, 100):
            break
        alcanzado = n
    return alcanzado


LIMITE_FONDO = 2000


def calidad_fondo(db: Session, fondo: RecursoDocumental) -> dict:
    """Indicadores del fondo: porcentaje de cumplimiento de cada criterio
    sobre las descripciones a las que aplica, y las que tienen incoherencias."""
    recursos = list(db.scalars(select(RecursoDocumental).where(
        (RecursoDocumental.fondo_id == fondo.id) | (RecursoDocumental.id == fondo.id),
        (RecursoDocumental.publicado_en.is_not(None)) | (RecursoDocumental.nivel != "unidad_documental"))
        .order_by(RecursoDocumental.creado_en).limit(LIMITE_FONDO)).all())
    resultados = [calidad(db, r) for r in recursos]
    criterios = {}
    for res in resultados:
        for c in res["criterios"]:
            if not c["aplica"]:
                continue
            k = criterios.setdefault(c["clave"], {"clave": c["clave"], "fila": c["fila"], "elemento": c["elemento"],
                                                  "nivel": c["nivel"], "aplican": 0, "cumplen": 0})
            k["aplican"] += 1
            k["cumplen"] += 1 if c["cumple"] else 0
    for k in criterios.values():
        k["porcentaje"] = round(100 * k["cumplen"] / k["aplican"])
    alcanzado = {n: sum(1 for r in resultados if r["nivel_alcanzado"] == n) for n in NIVELES}
    return {"fondo": {"id": str(fondo.id), "titulo": fondo.titulo}, "descripciones": len(resultados),
            "truncado": len(recursos) == LIMITE_FONDO, "criterios": list(criterios.values()),
            "nivel_alcanzado": alcanzado | {"ninguno": sum(1 for r in resultados if r["nivel_alcanzado"] is None)},
            "incoherencias": [{"id": r["id"], "titulo": r["titulo"], "detalle": r["incoherencias"]}
                              for r in resultados if r["incoherencias"]],
            "pendientes": sorted(({"id": r["id"], "titulo": r["titulo"], "nivel": r["nivel"],
                                   "porcentaje": r["porcentaje"]} for r in resultados if r["porcentaje"] is not None
                                  and r["porcentaje"] < 100), key=lambda x: x["porcentaje"])[:20]}


COLUMNAS_XLSX = (("id", "ID", 13), ("capa", "Capa", 9), ("origen", "Origen AGN", 11), ("entidad", "Entidad", 16),
                 ("elemento", "Elemento", 30), ("codigo_agn", "Código citado por el AGN", 30),
                 ("codigo_agn_correcto", "¿Código del AGN correcto?", 12), ("obligatoriedad", "Obligatoriedad AGN", 13),
                 ("regla_agn", "Regla AGN", 34), ("ric_cm", "RiC-CM 1.0 (correcto)", 34), ("ric_o", "RiC-O 1.1", 30),
                 ("campo", "Campo en RICORA", 40), ("estado", "Estado", 12), ("nivel", "Nivel de madurez", 12),
                 ("nota", "Nota", 50))


def catalogo_xlsx() -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Perfil RiC-Col"
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F3A5F")
    ws.append([t for _, t, _ in COLUMNAS_XLSX])
    for f in CATALOGO:
        fila = []
        for clave, _, _ in COLUMNAS_XLSX:
            v = getattr(f, clave)
            fila.append({"capa": CAPAS.get(v), "estado": ESTADOS.get(v), "nivel": NIVELES.get(v)}.get(clave, v)
                        if clave in ("capa", "estado", "nivel") else
                        ", ".join(f"rico:{t}" for t in v) if clave == "ric_o" else
                        ("Sí" if v else "No") if clave == "codigo_agn_correcto" else v)
        ws.append(fila)
    erratas = wb.create_sheet("Erratas del esquema AGN")
    erratas.append(["Página", "El esquema dice", "Lo correcto (RiC-CM 1.0 / RiC-O 1.1)"])
    for e in ERRATAS:
        erratas.append([e["pagina"], e["dice"], e["correcto"]])
    fuentes = wb.create_sheet("Fuentes")
    fuentes.append(["Fuente", "Versión", "Fecha", "Dirección"])
    for d in FUENTES.values():
        fuentes.append([d["nombre"], d["version"], d["fecha"], d["url"]])
    for hoja, anchos in ((ws, [a for _, _, a in COLUMNAS_XLSX]), (erratas, [14, 70, 70]), (fuentes, [60, 10, 12, 70])):
        for celda in hoja[1]:
            celda.font, celda.fill = encabezado, relleno
        for i, ancho in enumerate(anchos):
            hoja.column_dimensions[chr(65 + i)].width = ancho
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(wrap_text=True, vertical="top")
        hoja.freeze_panes = "A2"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
