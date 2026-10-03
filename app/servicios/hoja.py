"""
Hoja de cálculo común para el patrón «historial reciente con exportación
completa»: una sola forma de armar el archivo (encabezado en negrita,
fechas en la hora de Bogotá sin zona, columnas con ancho legible), la usan
todas las vistas que crecen sin límite.
"""

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from app.core.config import settings


def _celda(valor):
    if isinstance(valor, datetime):
        return valor.astimezone(ZoneInfo(settings.zona_horaria)).replace(tzinfo=None) if valor.tzinfo else valor
    if isinstance(valor, (list, tuple, set)):
        return ", ".join(str(x) for x in valor)
    if isinstance(valor, dict):
        return "; ".join(f"{k}: {v}" for k, v in valor.items())
    return valor


def libro(hojas: list[tuple[str, list[str], list[list]]]) -> bytes:
    """`hojas`: (título, encabezado, filas). Devuelve el .xlsx en bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for titulo, encabezado, filas in hojas:
        hoja = wb.create_sheet(titulo[:31])
        hoja.append(encabezado)
        for c in hoja[1]:
            c.font = Font(bold=True)
        for fila in filas:
            hoja.append([_celda(v) for v in fila])
        hoja.freeze_panes = "A2"
        for i, nombre in enumerate(encabezado, start=1):
            ancho = max([len(str(nombre))] + [len(str(_celda(f[i - 1]) or "")) for f in filas[:200]])
            hoja.column_dimensions[get_column_letter(i)].width = min(60, max(10, ancho + 2))
            if any(isinstance(_celda(f[i - 1]), datetime) for f in filas[:5]):
                for celda in hoja[get_column_letter(i)][1:]:
                    celda.number_format = "yyyy-mm-dd hh:mm"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
