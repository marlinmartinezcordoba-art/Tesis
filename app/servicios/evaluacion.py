"""
Evaluación ciega del motor frente a archivistas (objetivo 3 de la tesis).

Diseño (decisiones en documentacion/modulo-evaluacion.md):
- Patrón de referencia: la descripción **ciega** de cada archivista, hecha
  sin haber visto nunca la propuesta del motor. El servidor lo exige: quien
  vio la propuesta de un documento (en la condición asistida, al calificarla
  o en los resultados), o lo trabajó en Descripción, ya no puede
  describirlo a ciegas.
- Solo entran documentos sin describir: si ya lo estuvieran, su vocabulario
  le daría al motor la respuesta por el contexto.
- La descripción se hace en RiC (versión 2): entidades con su clase de
  RiC-O y relaciones con su propiedad de RiC-O (servicios/evaluacion_ric).
  La propuesta del motor se convierte a lo mismo, con la regla de la
  publicación.
- Métricas de entidades: precisión, exhaustividad y F1 por tipo, en dos
  versiones: estricta (mismo valor normalizado) y flexible (similitud de
  texto desde el umbral de la evaluación). Emparejamiento uno a uno.
- Métricas de relaciones: las mismas, por propiedad de RiC-O; una relación
  acierta si tiene la misma propiedad y sus dos extremos coinciden. Además,
  la confusión de propiedades (mismos extremos, otra relación) y la
  conformidad de cada descripción con RiC-O 1.1.
- Acuerdo entre archivistas: F1 entre dos anotaciones ciegas del mismo
  documento (en extracción de entidades no hay «negativos», así que el
  kappa no está definido; se usa F1, como proponen Hripcsak y Rothschild,
  2005). Para la rúbrica 1–5, kappa de Cohen ponderado cuadrático.
- Tiempos: mediana de minutos por documento, ciega frente a asistida.
"""

import io
import re
import statistics
import unicodedata
from collections import Counter
import uuid
from difflib import SequenceMatcher
from itertools import combinations

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.base import ahora
from app.models.descripcion import Relacion, TrabajoDescripcion, TrabajoInstanciacion
from app.models.evaluacion import (CRITERIOS, Anotacion, Calificacion, Evaluacion, EvaluacionDocumento,
                                   Exposicion)
from app.models.instanciacion import Instanciacion
from app.servicios.auditoria import registrar

# Los tipos de entidad que se comparan (los seis reutilizables y la fecha).
TIPOS = ("agente", "lugar", "fecha", "actividad", "tipo_actividad", "mandato", "forma_documental")


class ErrorEvaluacion(Exception):
    def __init__(self, mensaje: str, codigo: int = 422):
        super().__init__(mensaje)
        self.codigo = codigo


# --- Comparación (funciones puras) -----------------------------------------------------------------


def normalizar(texto: str | None) -> str:
    sin_tildes = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^\w\s]", " ", sin_tildes.lower()).split())


def clave_de(e: dict) -> str:
    """Con qué se agrupa una entidad para compararla: su tipo. El papel de
    un agente (productor, destinatario…) no es de la entidad sino de su
    relación con el documento, y se mide con las relaciones."""
    return e.get("tipo")


def texto_de(e: dict) -> str:
    """Lo que se compara: la fecha, por su EDTF si lo tiene; lo demás, por su valor."""
    if e.get("tipo") == "fecha" and e.get("edtf"):
        return e["edtf"].strip()
    return normalizar(e.get("valor"))


def emparejar(sistema: list[dict], referencia: list[dict], umbral: float) -> int:
    """Pares uno a uno (cada entidad cuenta una sola vez), del más parecido
    al menos, dentro del mismo tipo. Devuelve cuántos pares superan el
    umbral (1.0 = solo iguales tras normalizar)."""
    pares = 0
    for grupo in {clave_de(e) for e in sistema} & {clave_de(e) for e in referencia}:
        s = [texto_de(e) for e in sistema if clave_de(e) == grupo]
        r = [texto_de(e) for e in referencia if clave_de(e) == grupo]
        candidatos = sorted(((SequenceMatcher(None, a, b).ratio() if a != b else 1.0, i, j)
                             for i, a in enumerate(s) for j, b in enumerate(r)), reverse=True)
        usados_s, usados_r = set(), set()
        for parecido, i, j in candidatos:
            if parecido < umbral:
                break
            if i not in usados_s and j not in usados_r:
                usados_s.add(i)
                usados_r.add(j)
                pares += 1
    return pares


def conteo(sistema: list[dict], referencia: list[dict], umbral: float) -> dict:
    vp = emparejar(sistema, referencia, umbral)
    return {"vp": vp, "fp": len(sistema) - vp, "fn": len(referencia) - vp}


def prf(c: dict) -> dict:
    p = c["vp"] / (c["vp"] + c["fp"]) if c["vp"] + c["fp"] else None
    r = c["vp"] / (c["vp"] + c["fn"]) if c["vp"] + c["fn"] else None
    f = 2 * p * r / (p + r) if p is not None and r is not None and p + r else (0.0 if p is not None and r is not None else None)
    return {**c, "precision": _r(p), "exhaustividad": _r(r), "f1": _r(f)}


def _r(x: float | None) -> float | None:
    return round(x, 4) if x is not None else None


def _sumar(a: dict, b: dict) -> dict:
    return {k: a.get(k, 0) + b.get(k, 0) for k in ("vp", "fp", "fn")}


def kappa_ponderado(pares: list[tuple[int, int]], categorias=(1, 2, 3, 4, 5)) -> float | None:
    """Kappa de Cohen con pesos cuadráticos entre dos calificadores, para
    una escala ordinal. None si no hay datos o no hay variación esperada."""
    if not pares:
        return None
    k = len(categorias)
    indice = {c: i for i, c in enumerate(categorias)}
    n = len(pares)
    observada = [[0.0] * k for _ in range(k)]
    for a, b in pares:
        observada[indice[a]][indice[b]] += 1 / n
    fila = [sum(observada[i]) for i in range(k)]
    columna = [sum(observada[i][j] for i in range(k)) for j in range(k)]
    peso = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    desacuerdo_obs = sum(peso[i][j] * observada[i][j] for i in range(k) for j in range(k))
    desacuerdo_esp = sum(peso[i][j] * fila[i] * columna[j] for i in range(k) for j in range(k))
    return _r(1 - desacuerdo_obs / desacuerdo_esp) if desacuerdo_esp else None


def entidades_de_propuesta(propuesta: dict | None) -> list[dict]:
    """Los nodos de la propuesta del motor: una entidad con dos roles es un
    solo nodo (los roles son relaciones)."""
    from app.servicios import evaluacion_ric

    return evaluacion_ric.desde_roles(propuesta)["entidades"]


def entidades_de_anotacion(datos: dict | None) -> list[dict]:
    from app.servicios import evaluacion_ric

    return [e for e in evaluacion_ric.a_ric(datos)["entidades"]
            if e.get("tipo") in TIPOS and (e.get("valor") or "").strip()]


# --- Flujo ---------------------------------------------------------------------------------------------


def _evaluacion(db: Session, evaluacion_id: uuid.UUID) -> Evaluacion:
    ev = db.get(Evaluacion, evaluacion_id)
    if ev is None:
        raise ErrorEvaluacion("La evaluación no existe.", 404)
    return ev


def crear(db: Session, *, fondo_id: uuid.UUID, nombre: str, protocolo: str | None, umbral: float,
          usuario_id: uuid.UUID) -> Evaluacion:
    if not 0.5 <= umbral <= 1:
        raise ErrorEvaluacion("El umbral de similitud va de 0,5 a 1 (1 = solo valores iguales).")
    ev = Evaluacion(id=uuid.uuid4(), fondo_id=fondo_id, nombre=" ".join(nombre.split())[:200],
                    protocolo=(protocolo or "").strip() or None, umbral_similitud=umbral, creada_por_id=usuario_id)
    db.add(ev)
    db.flush()
    registrar(db, modulo="evaluacion", accion="evaluacion_creada", usuario_id=usuario_id, entidad_tipo="evaluacion",
              entidad_id=ev.id, detalle=ev.nombre, nuevo={"umbral": umbral, "fondo_id": str(fondo_id)})
    return ev


def descrito(db: Session, inst_id: uuid.UUID) -> bool:
    return db.scalar(select(Relacion.id).where(Relacion.destino_id == inst_id, Relacion.estado == "vigente",
                                               Relacion.codigo_ric == "has_or_had_instantiation").limit(1)) is not None


def agregar_documentos(db: Session, ev: Evaluacion, ids: list[uuid.UUID], usuario_id: uuid.UUID) -> int:
    if ev.estado != "preparacion":
        raise ErrorEvaluacion("Los documentos se agregan mientras la evaluación está en preparación.", 409)
    ya = set(db.scalars(select(EvaluacionDocumento.instanciacion_id).where(EvaluacionDocumento.evaluacion_id == ev.id)))
    nuevos = 0
    for inst_id in dict.fromkeys(ids):
        inst = db.get(Instanciacion, inst_id)
        if inst is None or inst.fondo_id != ev.fondo_id:
            raise ErrorEvaluacion("Uno de los documentos no es de este fondo.")
        if inst.estado != "listo_para_descripcion" or not (inst.texto_extraido or "").strip():
            raise ErrorEvaluacion(f"«{inst.nombre_original}» no está listo para describir o no tiene texto.")
        if descrito(db, inst.id):
            raise ErrorEvaluacion(f"«{inst.nombre_original}» ya está descrito: su vocabulario le daría la respuesta "
                                  "al motor y contaminaría la evaluación.")
        from app.servicios import derechos

        if derechos.instanciacion_restringida(db, inst):
            raise ErrorEvaluacion(f"«{inst.nombre_original}» tiene información clasificada o reservada: no entra a la "
                                  "evaluación, porque su texto se enviaría al motor y lo verían las evaluadoras.")
        if inst.id not in ya:
            db.add(EvaluacionDocumento(id=uuid.uuid4(), evaluacion_id=ev.id, instanciacion_id=inst.id))
            nuevos += 1
    db.flush()
    registrar(db, modulo="evaluacion", accion="documentos_agregados", usuario_id=usuario_id, entidad_tipo="evaluacion",
              entidad_id=ev.id, nuevo={"agregados": nuevos})
    return nuevos


def generar_propuestas(db: Session, ev: Evaluacion, usuario_id: uuid.UUID) -> dict:
    """El motor propone cada documento como lo haría en Descripción (con el
    vocabulario del fondo como contexto). Nadie ve el resultado aquí: queda
    guardado y se compara al final."""
    from app.servicios import motor, propuestas_ia, vocabulario

    if ev.estado != "preparacion":
        raise ErrorEvaluacion("Las propuestas se generan antes de iniciar la evaluación.", 409)
    hechas, avisos = 0, []
    for d in db.scalars(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id,
                                                          EvaluacionDocumento.propuesta.is_(None))).all():
        inst = db.get(Instanciacion, d.instanciacion_id)
        contexto = vocabulario.contexto_para_motor(db, ev.fondo_id, inst.texto_extraido or "")
        p = motor.proponer([motor.Documento(id=inst.id, nombre=inst.nombre_original, texto=inst.texto_extraido)],
                           "unidad_documental", contexto)
        # La propuesta queda como registro inalterable (RF-AI-002), también si falló.
        propuestas_ia.ubicar_fragmentos(db, p)
        registro = propuestas_ia.registrar(db, p, origen="evaluacion", fondo_id=ev.fondo_id, nivel="unidad_documental",
                                           instanciaciones=[inst.id], solicitada_por_id=usuario_id, evaluacion_id=ev.id,
                                           estado="evaluada")
        if not p.disponible:
            avisos.append(f"{inst.nombre_original}: {p.aviso or 'el motor no respondió'}")
            continue
        d.propuesta, d.motor, d.version_prompt, d.generada_en = p.a_dict(), p.motor, p.version_prompt, ahora()
        d.propuesta_id = registro.id if registro else None
        hechas += 1
    db.flush()
    registrar(db, modulo="evaluacion", accion="propuestas_generadas", usuario_id=usuario_id, entidad_tipo="evaluacion",
              entidad_id=ev.id, nuevo={"generadas": hechas, "sin_propuesta": len(avisos)})
    return {"generadas": hechas, "avisos": avisos}


def fijar_criterio(db: Session, ev: Evaluacion, datos: dict, usuario_id: uuid.UUID) -> dict:
    """El criterio de comparabilidad se fija antes de iniciar: después de ver
    los datos, cambiarlo invalidaría la respuesta a la pregunta."""
    from app.servicios import evaluacion_analisis as an

    if ev.estado != "preparacion":
        raise ErrorEvaluacion("El criterio de comparabilidad se fija antes de iniciar y no cambia después.", 409)
    try:
        nuevo = an.validar_criterio(datos)
    except an.ErrorAnalisis as exc:
        raise ErrorEvaluacion(str(exc), exc.codigo) from exc
    anterior, ev.criterio = an.criterio_de(ev.criterio), nuevo
    registrar(db, modulo="evaluacion", accion="criterio_fijado", usuario_id=usuario_id, entidad_tipo="evaluacion",
              entidad_id=ev.id, anterior=anterior, nuevo=nuevo)
    db.flush()
    return nuevo


def cambiar_estado(db: Session, ev: Evaluacion, estado: str, usuario_id: uuid.UUID) -> None:
    documentos = db.scalars(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id)).all()
    if estado == "en_curso":
        if ev.estado != "preparacion":
            raise ErrorEvaluacion("Solo se inicia una evaluación en preparación.", 409)
        if not documentos or any(d.propuesta is None for d in documentos):
            raise ErrorEvaluacion("Antes de iniciar, cada documento necesita la propuesta del motor.", 409)
        from app.servicios import evaluacion_analisis as an

        ev.criterio = an.criterio_de(ev.criterio)  # queda congelado tal como estaba al iniciar
        ev.iniciada_en = ahora()
    elif estado == "cerrada":
        if ev.estado != "en_curso":
            raise ErrorEvaluacion("Solo se cierra una evaluación en curso.", 409)
        ev.cerrada_en = ahora()
    else:
        raise ErrorEvaluacion("Estado desconocido.")
    anterior, ev.estado = ev.estado, estado
    registrar(db, modulo="evaluacion", accion="evaluacion_iniciada" if estado == "en_curso" else "evaluacion_cerrada",
              usuario_id=usuario_id, entidad_tipo="evaluacion", entidad_id=ev.id,
              anterior={"estado": anterior}, nuevo={"estado": estado, "documentos": len(documentos),
                                                    **({"criterio": ev.criterio} if estado == "en_curso" else {})})


def _documento(db: Session, ev: Evaluacion, inst_id: uuid.UUID) -> EvaluacionDocumento:
    d = db.scalar(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id,
                                                   EvaluacionDocumento.instanciacion_id == inst_id))
    if d is None:
        raise ErrorEvaluacion("Ese documento no está en la evaluación.", 404)
    return d


def contaminado(db: Session, ev: Evaluacion, inst_id: uuid.UUID, usuario_id: uuid.UUID) -> str | None:
    """Por qué esta persona ya no puede describir el documento a ciegas, o None."""
    if db.scalar(select(Exposicion.id).where(Exposicion.instanciacion_id == inst_id,
                                             Exposicion.usuario_id == usuario_id).limit(1)):
        return "ya vio la propuesta del motor para este documento"
    if db.scalar(select(TrabajoInstanciacion.trabajo_id).join(TrabajoDescripcion, TrabajoDescripcion.id ==
                                                              TrabajoInstanciacion.trabajo_id)
                 .where(TrabajoInstanciacion.instanciacion_id == inst_id, TrabajoDescripcion.usuario_id == usuario_id)
                 .limit(1)):
        return "abrió este documento en Descripción, donde el motor le mostró su propuesta"
    return None


def _exponer(db: Session, ev: Evaluacion, inst_id: uuid.UUID, usuario_id: uuid.UUID, motivo: str) -> None:
    db.add(Exposicion(evaluacion_id=ev.id, instanciacion_id=inst_id, usuario_id=usuario_id, motivo=motivo))
    registrar(db, modulo="evaluacion", accion="propuesta_vista", usuario_id=usuario_id, entidad_tipo="instanciacion",
              entidad_id=inst_id, nuevo={"evaluacion_id": str(ev.id), "motivo": motivo})


def iniciar_anotacion(db: Session, ev: Evaluacion, inst_id: uuid.UUID, condicion: str, usuario_id: uuid.UUID) -> Anotacion:
    if ev.estado != "en_curso":
        raise ErrorEvaluacion("La evaluación no está en curso.", 409)
    d = _documento(db, ev, inst_id)
    existente = db.scalar(select(Anotacion).where(Anotacion.evaluacion_id == ev.id, Anotacion.instanciacion_id == inst_id,
                                                  Anotacion.evaluador_id == usuario_id, Anotacion.condicion == condicion))
    if existente is not None and existente.estado != "anulada":
        return existente
    if condicion == "ciega":
        motivo = contaminado(db, ev, inst_id, usuario_id)
        if motivo:
            raise ErrorEvaluacion(f"No puede describirlo a ciegas: {motivo}.", 409)
    else:
        abierta = db.scalar(select(Anotacion.id).where(Anotacion.evaluacion_id == ev.id,
                                                       Anotacion.instanciacion_id == inst_id,
                                                       Anotacion.evaluador_id == usuario_id,
                                                       Anotacion.condicion == "ciega", Anotacion.estado == "en_curso"))
        if abierta:
            raise ErrorEvaluacion("Primero envíe su descripción a ciegas: ver la propuesta la invalidaría.", 409)
        _exponer(db, ev, inst_id, usuario_id, "asistida")
    if existente is not None:  # anulada: se reabre con una nueva fecha de inicio
        existente.estado, existente.iniciada_en, existente.enviada_en = "en_curso", ahora(), None
        existente.datos = _inicial(d) if condicion == "asistida" else None  # la asistida vuelve con la propuesta
        a = existente
    else:
        a = Anotacion(id=uuid.uuid4(), evaluacion_id=ev.id, instanciacion_id=inst_id, evaluador_id=usuario_id,
                      condicion=condicion, datos=_inicial(d) if condicion == "asistida" else None)
        db.add(a)
    db.flush()
    registrar(db, modulo="evaluacion", accion="anotacion_iniciada", usuario_id=usuario_id, entidad_tipo="instanciacion",
              entidad_id=inst_id, nuevo={"evaluacion_id": str(ev.id), "condicion": condicion})
    return a


def _inicial(d: EvaluacionDocumento) -> dict:
    """La condición asistida empieza con la propuesta del motor puesta, ya
    en RiC: entidades y relaciones."""
    from app.servicios import evaluacion_ric

    return evaluacion_ric.desde_roles(d.propuesta)


def _limpiar(datos: dict) -> dict:
    from app.servicios import evaluacion_ric

    try:
        return evaluacion_ric.limpiar(datos)
    except evaluacion_ric.ErrorRic as exc:
        raise ErrorEvaluacion(str(exc)) from exc


def guardar(db: Session, a: Anotacion, datos: dict, usuario_id: uuid.UUID, enviar: bool) -> Anotacion:
    if a.evaluador_id != usuario_id:
        raise ErrorEvaluacion("Esa anotación es de otra persona.", 403)
    if a.estado != "en_curso":
        raise ErrorEvaluacion("La anotación ya se envió.", 409)
    ev = db.get(Evaluacion, a.evaluacion_id)
    if ev.estado != "en_curso":
        raise ErrorEvaluacion("La evaluación no está en curso.", 409)
    from app.servicios import evaluacion_ric

    a.datos = _limpiar(datos)
    if enviar:
        if not a.datos["entidades"] and not a.datos["titulo"]:
            raise ErrorEvaluacion("La descripción está vacía.")
        # La referencia tiene que ser RiC válida: si no, se compararía el
        # motor con algo que RiC-O no admite.
        revision = evaluacion_ric.conformidad(a.datos)
        if not revision["conforme"]:
            raise ErrorEvaluacion("La descripción todavía no es RiC válida: " + " · ".join(revision["problemas"][:6]))
        a.estado, a.enviada_en = "enviada", ahora()
        registrar(db, modulo="evaluacion", accion="anotacion_enviada", usuario_id=usuario_id,
                  entidad_tipo="instanciacion", entidad_id=a.instanciacion_id,
                  nuevo={"evaluacion_id": str(ev.id), "condicion": a.condicion, "entidades": len(a.datos["entidades"]),
                         "relaciones": len(a.datos["relaciones"]),
                         "segundos": round((a.enviada_en - a.iniciada_en).total_seconds())})
    db.flush()
    return a


def ver_propuesta(db: Session, ev: Evaluacion, inst_id: uuid.UUID, usuario_id: uuid.UUID) -> dict:
    """Para calificar con la rúbrica: verla deja constancia (exposición)."""
    if ev.estado != "en_curso":
        raise ErrorEvaluacion("La evaluación no está en curso.", 409)
    d = _documento(db, ev, inst_id)
    if db.scalar(select(Anotacion.id).where(Anotacion.evaluacion_id == ev.id, Anotacion.instanciacion_id == inst_id,
                                            Anotacion.evaluador_id == usuario_id, Anotacion.condicion == "ciega",
                                            Anotacion.estado == "en_curso")):
        raise ErrorEvaluacion("Primero envíe su descripción a ciegas: ver la propuesta la invalidaría.", 409)
    _exponer(db, ev, inst_id, usuario_id, "calificacion")
    return _inicial(d)


def calificar(db: Session, ev: Evaluacion, inst_id: uuid.UUID, puntajes: dict[str, int], usuario_id: uuid.UUID) -> None:
    if ev.estado != "en_curso":
        raise ErrorEvaluacion("La evaluación no está en curso.", 409)
    _documento(db, ev, inst_id)
    if not db.scalar(select(Exposicion.id).where(Exposicion.evaluacion_id == ev.id, Exposicion.instanciacion_id == inst_id,
                                                 Exposicion.usuario_id == usuario_id).limit(1)):
        raise ErrorEvaluacion("Para calificar la propuesta hay que verla primero.", 409)
    if set(puntajes) != set(CRITERIOS) or any(not isinstance(v, int) or not 1 <= v <= 5 for v in puntajes.values()):
        raise ErrorEvaluacion("Califique los tres criterios de 1 a 5.")
    for criterio, puntaje in puntajes.items():
        c = db.scalar(select(Calificacion).where(Calificacion.evaluacion_id == ev.id, Calificacion.instanciacion_id == inst_id,
                                                 Calificacion.evaluador_id == usuario_id, Calificacion.criterio == criterio))
        if c is None:
            db.add(Calificacion(id=uuid.uuid4(), evaluacion_id=ev.id, instanciacion_id=inst_id, evaluador_id=usuario_id,
                                criterio=criterio, puntaje=puntaje))
        else:
            c.puntaje = puntaje  # el valor anterior queda en la auditoría
    registrar(db, modulo="evaluacion", accion="calificacion_registrada", usuario_id=usuario_id,
              entidad_tipo="instanciacion", entidad_id=inst_id, nuevo={"evaluacion_id": str(ev.id), **puntajes})
    db.flush()


def tareas(db: Session, ev: Evaluacion, usuario_id: uuid.UUID) -> list[dict]:
    salida = []
    for d in db.scalars(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id)).all():
        inst = db.get(Instanciacion, d.instanciacion_id)
        mias = {a.condicion: a for a in db.scalars(select(Anotacion).where(
            Anotacion.evaluacion_id == ev.id, Anotacion.instanciacion_id == inst.id, Anotacion.evaluador_id == usuario_id))}
        calificado = db.scalar(select(Calificacion.id).where(Calificacion.evaluacion_id == ev.id,
                                                             Calificacion.instanciacion_id == inst.id,
                                                             Calificacion.evaluador_id == usuario_id).limit(1)) is not None
        salida.append({"instanciacion_id": str(inst.id), "nombre": inst.nombre_original,
                       "ciega": _estado_anotacion(mias.get("ciega")), "asistida": _estado_anotacion(mias.get("asistida")),
                       "puede_ciega": mias.get("ciega") is not None or contaminado(db, ev, inst.id, usuario_id) is None,
                       "calificado": calificado})
    return salida


def _estado_anotacion(a: Anotacion | None) -> dict | None:
    return {"id": str(a.id), "estado": a.estado} if a else None


def anotacion_out(db: Session, a: Anotacion) -> dict:
    from app.servicios import evaluacion_ric

    inst = db.get(Instanciacion, a.instanciacion_id)
    datos = evaluacion_ric.a_ric(a.datos) if a.datos else {"entidades": [], "relaciones": []}
    return {"id": str(a.id), "condicion": a.condicion, "estado": a.estado, "datos": datos,
            "conformidad": evaluacion_ric.conformidad(a.datos) if a.datos else None,
            "iniciada_en": a.iniciada_en, "enviada_en": a.enviada_en,
            "documento": {"id": str(inst.id), "nombre": inst.nombre_original, "texto": inst.texto_extraido}}


# --- Resultados ----------------------------------------------------------------------------------------


def resultados(db: Session, ev: Evaluacion, usuario_id: uuid.UUID) -> dict:
    documentos = db.scalars(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id)).all()
    for d in documentos:  # quien mira los resultados ve las propuestas: ya no describe a ciegas
        if not db.scalar(select(Exposicion.id).where(Exposicion.evaluacion_id == ev.id,
                                                     Exposicion.instanciacion_id == d.instanciacion_id,
                                                     Exposicion.usuario_id == usuario_id, Exposicion.motivo == "resultados")):
            db.add(Exposicion(evaluacion_id=ev.id, instanciacion_id=d.instanciacion_id, usuario_id=usuario_id,
                              motivo="resultados"))
    anotaciones = db.scalars(select(Anotacion).where(Anotacion.evaluacion_id == ev.id, Anotacion.estado == "enviada")).all()
    por_doc: dict[uuid.UUID, list[Anotacion]] = {}
    for a in anotaciones:
        por_doc.setdefault(a.instanciacion_id, []).append(a)

    from app.servicios import evaluacion_ric as ric

    estricto = {t: {"vp": 0, "fp": 0, "fn": 0} for t in TIPOS}
    flexible = {t: {"vp": 0, "fp": 0, "fn": 0} for t in TIPOS}
    rel_estricto: dict[str, dict] = {}
    rel_flexible: dict[str, dict] = {}
    confusiones: dict[tuple[str, str], int] = {}
    acuerdo = {"vp": 0, "fp": 0, "fn": 0}
    acuerdo_rel = {"vp": 0, "fp": 0, "fn": 0}
    conformidad = {"motor": [], "ciega": [], "asistida": []}
    por_documento = []
    from app.servicios import evaluacion_analisis as an

    acuerdo_tipo = {t: {"vp": 0, "fp": 0, "fn": 0} for t in TIPOS}
    cuentas_doc, cuentas_doc_rel = [], []
    supervision_tipo: dict[str, dict] = {}
    lista_desacuerdos = []

    def acumular(destino: dict, cuentas: dict) -> None:
        for k, c in cuentas.items():
            destino[k] = _sumar(destino.get(k, {}), c)

    for d in documentos:
        sistema = entidades_de_propuesta(d.propuesta)
        motor_ric = ric.desde_roles(d.propuesta)
        sistema_rel = ric.tripletas(motor_ric)
        rev_motor = ric.conformidad(motor_ric)
        conformidad["motor"].append(rev_motor)
        anotadas = por_doc.get(d.instanciacion_id, [])
        for a in anotadas:
            conformidad[a.condicion].append(ric.conformidad(a.datos))
        ciegas = [a for a in anotadas if a.condicion == "ciega"]
        f1_doc, f1_rel_doc = [], []
        doc_ent, doc_rel = {"vp": 0, "fp": 0, "fn": 0}, {"vp": 0, "fp": 0, "fn": 0}
        for a in anotadas:
            if a.condicion == "asistida":
                for t, c in an.supervision(sistema, entidades_de_anotacion(a.datos), ev.umbral_similitud).items():
                    supervision_tipo[t] = {k: supervision_tipo.get(t, {}).get(k, 0) + v for k, v in c.items()}
        for a in ciegas:  # cada descripción ciega es una referencia
            referencia = entidades_de_anotacion(a.datos)
            for t in TIPOS:
                s_t = [e for e in sistema if e["tipo"] == t]
                r_t = [e for e in referencia if e["tipo"] == t]
                estricto[t] = _sumar(estricto[t], conteo(s_t, r_t, 1.0))
                flexible[t] = _sumar(flexible[t], conteo(s_t, r_t, ev.umbral_similitud))
            c_doc = conteo(sistema, referencia, ev.umbral_similitud)
            doc_ent = _sumar(doc_ent, c_doc)
            f1_doc.append(prf(c_doc)["f1"])
            referencia_rel = ric.tripletas(a.datos)
            lista_desacuerdos += an.desacuerdos(str(d.instanciacion_id), str(a.id), sistema, referencia, sistema_rel,
                                                referencia_rel, ev.umbral_similitud)
            acumular(rel_estricto, ric.conteo_por_propiedad(sistema_rel, referencia_rel, 1.0)[0])
            cuentas, confundidas = ric.conteo_por_propiedad(sistema_rel, referencia_rel, ev.umbral_similitud)
            acumular(rel_flexible, cuentas)
            for par in confundidas:
                confusiones[par] = confusiones.get(par, 0) + 1
            total = {"vp": 0, "fp": 0, "fn": 0}
            for c in cuentas.values():
                total = _sumar(total, c)
            f1_rel_doc.append(prf(total)["f1"])
            doc_rel = _sumar(doc_rel, total)
        cuentas_doc.append(doc_ent)
        cuentas_doc_rel.append(doc_rel)
        for a, b in combinations(ciegas, 2):  # acuerdo entre archivistas
            ea, eb = entidades_de_anotacion(a.datos), entidades_de_anotacion(b.datos)
            acuerdo = _sumar(acuerdo, conteo(ea, eb, ev.umbral_similitud))
            for t in TIPOS:  # el techo de cada tipo, para el criterio de comparabilidad
                acuerdo_tipo[t] = _sumar(acuerdo_tipo[t], conteo([e for e in ea if e["tipo"] == t],
                                                                 [e for e in eb if e["tipo"] == t], ev.umbral_similitud))
            vp, _ = ric.emparejar(ric.tripletas(a.datos), ric.tripletas(b.datos), ev.umbral_similitud)
            acuerdo_rel = _sumar(acuerdo_rel, {"vp": vp, "fp": len(ric.tripletas(a.datos)) - vp,
                                               "fn": len(ric.tripletas(b.datos)) - vp})
        inst = db.get(Instanciacion, d.instanciacion_id)

        def media(valores):
            v = [x for x in valores if x is not None]
            return _r(statistics.mean(v)) if v else None
        por_documento.append({"instanciacion_id": str(inst.id), "nombre": inst.nombre_original,
                              "entidades_motor": len(sistema), "relaciones_motor": len(sistema_rel),
                              "referencias_ciegas": len(ciegas), "f1_flexible": media(f1_doc),
                              "f1_relaciones_flexible": media(f1_rel_doc), "motor_conforme": rev_motor["conforme"],
                              "descrito_durante_la_evaluacion": descrito(db, inst.id)})

    def tabla(cuentas: dict) -> dict:
        total = {"vp": 0, "fp": 0, "fn": 0}
        for c in cuentas.values():
            total = _sumar(total, c)
        filas = {t: prf(c) for t, c in cuentas.items() if c["vp"] + c["fp"] + c["fn"]}
        f1s = [f["f1"] for f in filas.values() if f["f1"] is not None]
        return {"por_tipo": filas, "micro": prf(total), "macro_f1": _r(statistics.mean(f1s)) if f1s else None}

    minutos = {c: [(a.enviada_en - a.iniciada_en).total_seconds() / 60 for a in anotaciones if a.condicion == c]
               for c in ("ciega", "asistida")}
    tiempos = {c: {"n": len(v), "mediana_minutos": _r(statistics.median(v)) if v else None,
                   "media_minutos": _r(statistics.mean(v)) if v else None} for c, v in minutos.items()}
    # Las evaluadoras no se identifican: «Evaluadora 1, 2…» en el orden en que empezaron.
    orden = sorted({a.evaluador_id for a in anotaciones},
                   key=lambda u: min(x.iniciada_en for x in anotaciones if x.evaluador_id == u))
    alias = {u: f"Evaluadora {i}" for i, u in enumerate(orden, 1)}
    nombres_doc = {d.instanciacion_id: db.get(Instanciacion, d.instanciacion_id).nombre_original for d in documentos}
    pareados: dict[tuple, dict] = {}
    for a in anotaciones:
        pareados.setdefault((a.evaluador_id, a.instanciacion_id), {})[a.condicion] = \
            _r((a.enviada_en - a.iniciada_en).total_seconds() / 60)
    tiempos_documento = [{"evaluadora": alias[u], "documento": nombres_doc.get(i), "ciega": v.get("ciega"),
                          "asistida": v.get("asistida"),
                          "diferencia": _r(v["ciega"] - v["asistida"]) if "ciega" in v and "asistida" in v else None}
                         for (u, i), v in sorted(pareados.items(), key=lambda x: (alias[x[0][0]], str(nombres_doc.get(x[0][1]))))]
    prueba_tiempos = an.wilcoxon([(t["ciega"], t["asistida"]) for t in tiempos_documento if t["diferencia"] is not None])

    def resumen_conformidad(revisiones: list[dict]) -> dict:
        relaciones = sum(x["relaciones"] for x in revisiones)
        validas = sum(x["relaciones_validas"] for x in revisiones)
        frecuencia: dict[str, int] = {}
        for x in revisiones:
            for problema in x["problemas"]:
                frecuencia[problema] = frecuencia.get(problema, 0) + 1
        return {"descripciones": len(revisiones), "conformes": sum(1 for x in revisiones if x["conforme"]),
                "relaciones": relaciones, "relaciones_validas": validas,
                "tasa_relaciones_validas": _r(validas / relaciones) if relaciones else None,
                "problemas": [{"problema": k, "veces": v}
                              for k, v in sorted(frecuencia.items(), key=lambda x: (-x[1], x[0]))[:20]]}

    calificaciones = db.scalars(select(Calificacion).where(Calificacion.evaluacion_id == ev.id)).all()
    rubrica = {}
    for criterio in CRITERIOS:
        del_criterio = [c for c in calificaciones if c.criterio == criterio]
        por_item: dict[uuid.UUID, dict[uuid.UUID, int]] = {}
        for c in del_criterio:
            por_item.setdefault(c.instanciacion_id, {})[c.evaluador_id] = c.puntaje
        pares = [(p[a], p[b]) for p in por_item.values() for a, b in combinations(sorted(p, key=str), 2)]
        rubrica[criterio] = {"n": len(del_criterio),
                             "media": _r(statistics.mean(c.puntaje for c in del_criterio)) if del_criterio else None,
                             "pares": len(pares), "kappa_ponderado": kappa_ponderado(pares)}
    registrar(db, modulo="evaluacion", accion="resultados_consultados", usuario_id=usuario_id,
              entidad_tipo="evaluacion", entidad_id=ev.id, nuevo={"anotaciones": len(anotaciones)})
    db.flush()
    motores = sorted({(d.motor, d.version_prompt) for d in documentos if d.motor}, key=str)
    salida = {"evaluacion": evaluacion_out(db, ev), "umbral": ev.umbral_similitud,
            "motores": [{"motor": m, "version_prompt": v} for m, v in motores],
            "estricto": tabla(estricto), "flexible": tabla(flexible),
            "relaciones": {"estricto": tabla(rel_estricto), "flexible": tabla(rel_flexible),
                           "acuerdo_entre_archivistas": prf(acuerdo_rel) if sum(acuerdo_rel.values()) else None,
                           "confusiones": [{"motor": m, "archivista": a, "veces": n} for (m, a), n in
                                           sorted(confusiones.items(), key=lambda x: (-x[1], x[0]))],
                           "propiedades": {c: ric.nombre_rico(c) for c in ric.PROPIEDADES}},
            "conformidad": {k: resumen_conformidad(v) for k, v in conformidad.items()},
            "acuerdo_entre_archivistas": prf(acuerdo) if sum(acuerdo.values()) else None,
            "tiempos": tiempos, "tiempos_por_documento": tiempos_documento, "prueba_tiempos": prueba_tiempos, "rubrica": rubrica, "documentos": por_documento,
            "acuerdo_por_tipo": {t: prf(c) for t, c in acuerdo_tipo.items() if sum(c.values())},
            "intervalos": {"entidades_flexible": an.intervalo_f(cuentas_doc), "relaciones_flexible": an.intervalo_f(cuentas_doc_rel)},
            "supervision": an.resumen_supervision(supervision_tipo),
            "errores": _resumen_errores(db, ev, lista_desacuerdos),
            "anotaciones": {"ciegas": sum(1 for a in anotaciones if a.condicion == "ciega"),
                            "asistidas": sum(1 for a in anotaciones if a.condicion == "asistida"),
                            "evaluadores": len({a.evaluador_id for a in anotaciones})}}
    salida["comparabilidad"] = an.veredicto(salida, ev.criterio)
    return salida


def _resumen_errores(db: Session, ev: Evaluacion, lista: list[dict]) -> dict:
    from app.models.evaluacion import ClasificacionError
    from app.servicios import evaluacion_analisis as an

    filas = db.scalars(select(ClasificacionError).where(ClasificacionError.evaluacion_id == ev.id)).all()
    resumen = an.resumen_clasificaciones([x["clave"] for x in lista], filas)
    resumen["por_clase"] = dict(Counter(x["clase"] for x in lista))
    resumen["lista"] = [x | {"tipo_error": resumen["finales"].get(x["clave"])} for x in lista]
    del resumen["finales"]
    return resumen


def desacuerdos(db: Session, ev: Evaluacion, usuario_id: uuid.UUID) -> dict:
    """Los desacuerdos para clasificarlos. Cada investigadora ve su propia
    clasificación; la de la otra, solo después de haber clasificado ella
    (independencia), y el consenso cuando ya existe."""
    from app.models.evaluacion import ClasificacionError
    from app.servicios import evaluacion_analisis as an

    datos = resultados(db, ev, usuario_id)
    filas = db.scalars(select(ClasificacionError).where(ClasificacionError.evaluacion_id == ev.id)).all()
    por_clave: dict[str, list] = {}
    for f in filas:
        por_clave.setdefault(f.clave, []).append(f)
    nombres_doc = {x["instanciacion_id"]: x["nombre"] for x in datos["documentos"]}
    items = []
    for x in datos["errores"]["lista"]:
        propias = por_clave.get(x["clave"], [])
        mia = next((f.tipo_error for f in propias if f.ronda == "individual" and f.clasificadora_id == usuario_id), None)
        otras = [f.tipo_error for f in propias if f.ronda == "individual" and f.clasificadora_id != usuario_id]
        consenso = next((f.tipo_error for f in propias if f.ronda == "consenso"), None)
        items.append(x | {"documento": nombres_doc.get(x["documento_id"]), "mia": mia,
                          "otra": (otras[0] if otras else None) if mia else None, "otra_clasifico": bool(otras),
                          "consenso": consenso,
                          "requiere_consenso": consenso is None and mia is not None and bool(otras) and otras[0] != mia})
    resumen = {k: v for k, v in datos["errores"].items() if k != "lista"}
    return {"estado": ev.estado, "tipos_error": an.TIPOS_ERROR, "clases": an.CLASES_DESACUERDO, "resumen": resumen,
            "desacuerdos": items}


def clasificar_desacuerdo(db: Session, ev: Evaluacion, clave: str, tipo_error: str, ronda: str, nota: str | None,
                          usuario_id: uuid.UUID) -> None:
    from app.servicios import evaluacion_analisis as an

    claves = {x["clave"] for x in resultados(db, ev, usuario_id)["errores"]["lista"]}
    try:
        an.clasificar(db, ev, clave, tipo_error, ronda, nota, claves, usuario_id)
    except an.ErrorAnalisis as exc:
        raise ErrorEvaluacion(str(exc), exc.codigo) from exc


def evaluacion_out(db: Session, ev: Evaluacion) -> dict:
    from app.servicios.evaluacion_analisis import criterio_de

    documentos = db.scalars(select(EvaluacionDocumento).where(EvaluacionDocumento.evaluacion_id == ev.id)).all()
    return {"id": str(ev.id), "nombre": ev.nombre, "protocolo": ev.protocolo, "estado": ev.estado,
            "umbral_similitud": ev.umbral_similitud, "creada_en": ev.creada_en, "iniciada_en": ev.iniciada_en,
            "cerrada_en": ev.cerrada_en, "documentos": len(documentos),
            "criterio": criterio_de(ev.criterio),
            "con_propuesta": sum(1 for d in documentos if d.propuesta is not None)}


def hoja_de_calculo(datos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Entidades"
    hoja.append(["Medida", "Tipo", "VP", "FP", "FN", "Precisión", "Exhaustividad", "F1"])
    for medida in ("estricto", "flexible"):
        for tipo, f in datos[medida]["por_tipo"].items():
            hoja.append([medida, tipo, f["vp"], f["fp"], f["fn"], f["precision"], f["exhaustividad"], f["f1"]])
        m = datos[medida]["micro"]
        hoja.append([medida, "total (micro)", m["vp"], m["fp"], m["fn"], m["precision"], m["exhaustividad"], m["f1"]])
        hoja.append([medida, "F1 macro", None, None, None, None, None, datos[medida]["macro_f1"]])
    rel = libro.create_sheet("Relaciones RiC")
    rel.append(["Medida", "Propiedad RiC-O", "VP", "FP", "FN", "Precisión", "Exhaustividad", "F1"])
    nombres = datos["relaciones"]["propiedades"]
    for medida in ("estricto", "flexible"):
        tabla_r = datos["relaciones"][medida]
        for codigo, f in tabla_r["por_tipo"].items():
            rel.append([medida, nombres.get(codigo, codigo), f["vp"], f["fp"], f["fn"], f["precision"],
                        f["exhaustividad"], f["f1"]])
        m = tabla_r["micro"]
        rel.append([medida, "total (micro)", m["vp"], m["fp"], m["fn"], m["precision"], m["exhaustividad"], m["f1"]])
        rel.append([medida, "F1 macro", None, None, None, None, None, tabla_r["macro_f1"]])
    rel.append([])
    rel.append(["Confusión de propiedades", "Motor", "Archivista", "Veces"])
    for c in datos["relaciones"]["confusiones"]:
        rel.append([None, nombres.get(c["motor"], c["motor"]), nombres.get(c["archivista"], c["archivista"]), c["veces"]])
    conf = libro.create_sheet("Conformidad RiC-O")
    conf.append(["Quién describe", "Descripciones", "Conformes", "Relaciones", "Relaciones válidas", "Tasa"])
    for quien, c in datos["conformidad"].items():
        conf.append([quien, c["descripciones"], c["conformes"], c["relaciones"], c["relaciones_validas"],
                     c["tasa_relaciones_validas"]])
    conf.append([])
    conf.append(["Quién describe", "Problema", "Veces"])
    for quien, c in datos["conformidad"].items():
        for x in c["problemas"]:
            conf.append([quien, x["problema"], x["veces"]])
    otra = libro.create_sheet("Acuerdo, tiempos y rúbrica")
    a = datos["acuerdo_entre_archivistas"] or {}
    otra.append(["Acuerdo entre archivistas (F1)", a.get("f1")])
    otra.append(["Acuerdo entre archivistas en relaciones (F1)", (datos["relaciones"]["acuerdo_entre_archivistas"] or {}).get("f1")])
    for c, t in datos["tiempos"].items():
        otra.append([f"Tiempo {c} (mediana, minutos)", t["mediana_minutos"], "n", t["n"]])
    for criterio, r in datos["rubrica"].items():
        otra.append([f"Rúbrica · {criterio} (media)", r["media"], "kappa ponderado", r["kappa_ponderado"],
                     "pares", r["pares"]])
    for nombre, i in (("F1 entidades (flexible)", datos["intervalos"]["entidades_flexible"]),
                      ("F1 relaciones (flexible)", datos["intervalos"]["relaciones_flexible"])):
        otra.append([f"IC 95 % {nombre}", i["inferior"], i["superior"], "documentos", i["documentos"],
                     "remuestreos", i["remuestreos"]])
    w = datos["prueba_tiempos"]
    otra.append(["Wilcoxon, tiempos ciega frente a asistida", "n", w["n"], "z", w["z"], "p", w["p"], w["nota"] or ""])
    comp = libro.create_sheet("Comparabilidad")
    c = datos["comparabilidad"]
    comp.append(["Criterio fijado antes de iniciar", "margen F", c["criterio"]["margen_f"], "mínimo de casos",
                 c["criterio"]["minimo_casos"], "rúbrica mínima", c["criterio"]["rubrica_minima"],
                 "conformidad asistida", c["criterio"]["conformidad_asistida"]])
    comp.append(["Tipo", "Casos de referencia", "F del motor (flexible)", "Acuerdo entre archivistas (techo)",
                 "Diferencia", "Resultado", "Motivos"])
    for t in c["tipos"]:
        comp.append([t["tipo"], t["casos"], t["f_motor"], t["techo"], t["diferencia"], t["estado"], "; ".join(t["motivos"])])
    sup = libro.create_sheet("Supervisión (asistida)")
    sup.append(["Tipo", "Aceptadas", "Corregidas", "Rechazadas", "Agregadas por la archivista"])
    for t, x in datos["supervision"]["por_tipo"].items():
        sup.append([t, x["aceptada"], x["corregida"], x["rechazada"], x["agregada"]])
    tot = datos["supervision"]["total"]
    sup.append(["Total", tot.get("aceptada", 0), tot.get("corregida", 0), tot.get("rechazada", 0), tot.get("agregada", 0)])
    tie = libro.create_sheet("Tiempos por documento")
    tie.append(["Evaluadora", "Documento", "Minutos a ciegas", "Minutos asistida", "Diferencia"])
    for t in datos["tiempos_por_documento"]:
        tie.append([t["evaluadora"], t["documento"], t["ciega"], t["asistida"], t["diferencia"]])
    des = libro.create_sheet("Desacuerdos")
    des.append(["Clave", "Documento", "Clase", "Tipo", "Propiedad (motor)", "Propiedad (archivista)", "Motor",
                "Archivista", "Tipo de error (final)"])
    nombres_doc = {d["instanciacion_id"]: d["nombre"] for d in datos["documentos"]}
    from app.servicios.evaluacion_analisis import TIPOS_ERROR

    for x in datos["errores"]["lista"]:
        des.append([x["clave"], nombres_doc.get(x["documento_id"]), x["clase_nombre"], x["tipo"],
                    nombres.get(x["propiedad"], x["propiedad"]) if x["propiedad"] else None,
                    nombres.get(x["propiedad_archivista"], x["propiedad_archivista"]) if x["propiedad_archivista"] else None,
                    x["motor"], x["archivista"], TIPOS_ERROR.get(x["tipo_error"]) if x["tipo_error"] else None])
    des.append([])
    e = datos["errores"]
    des.append(["Kappa entre clasificadoras", e["kappa"], "pares", e["pares"], "pendientes de consenso",
                e["pendientes_consenso"], "sin clasificar", e["sin_clasificar"]])
    docs = libro.create_sheet("Documentos")
    docs.append(["Documento", "Entidades del motor", "Relaciones del motor", "Referencias ciegas", "F1 flexible",
                 "F1 relaciones (flexible)", "Motor conforme con RiC-O", "Descrito durante la evaluación"])
    for d in datos["documentos"]:
        docs.append([d["nombre"], d["entidades_motor"], d["relaciones_motor"], d["referencias_ciegas"], d["f1_flexible"],
                     d["f1_relaciones_flexible"], "sí" if d["motor_conforme"] else "no",
                     "sí" if d["descrito_durante_la_evaluacion"] else "no"])
    for h in libro.worksheets:
        for celda in h[1]:
            celda.font = Font(bold=True)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
