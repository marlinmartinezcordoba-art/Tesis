"""
Módulo 4 · Generación de instrumentos de descripción.

Solo lee y agrega lo que ya existe (descripción y vocabularios): catálogo
navegable, inventario documental sobre el FUID del AGN, guía del fondo e
índice de términos. Nunca escribe sobre el grafo archivístico.

Regla de procedencia: todo lo que sale de aquí (pantalla o archivo) se arma
con listas cerradas de campos permitidos, y `sin_campos_internos()` lo
comprueba antes de responder. Origen, confianza, motor, estado de revisión
y fragmentos citados no salen jamás.
"""

import io
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.models.descripcion import EntidadVocabulario, Fecha, Relacion
from app.models.instanciacion import Instanciacion
from app.models.recurso_documental import NIVEL_DESCRIPCION, RecursoDocumental
from app.servicios import alertas, consulta, motor, vocabulario

NIVEL_PLURAL = {"seccion": "secciones", "serie": "series", "subserie": "subseries", "expediente": "expedientes",
                "unidad_documental": "unidades documentales", "parte_documental": "partes documentales"}
NIVEL_NOMBRE = {"fondo": "Fondo", "seccion": "Sección", "serie": "Serie", "subserie": "Subserie",
                "expediente": "Expediente", "unidad_documental": "Unidad documental",
                "parte_documental": "Parte documental"}

# Claves que no pueden aparecer en ninguna respuesta ni archivo de este módulo.
CAMPOS_INTERNOS = consulta.CAMPOS_PROHIBIDOS | {"estado_revision", "origen_dato", "fragmento_inicio",
                                                 "fragmento_instanciacion_id", "confirmada_por_id"}


class ErrorInstrumento(Exception):
    def __init__(self, mensaje: str, codigo: int = 422):
        super().__init__(mensaje)
        self.codigo = codigo


def sin_campos_internos(datos):
    """Comprobación final antes de que algo salga del backend: si una clave
    interna se coló (por un cambio futuro en otra parte), falla en voz alta
    en lugar de filtrarla en silencio."""
    if isinstance(datos, dict):
        for clave, valor in datos.items():
            if clave in CAMPOS_INTERNOS:
                raise AssertionError(f"Campo interno «{clave}» en una salida del módulo de instrumentos.")
            sin_campos_internos(valor)
    elif isinstance(datos, list):
        for x in datos:
            sin_campos_internos(x)
    return datos


# --- Árbol del fondo -----------------------------------------------------------------------------


@dataclass
class Arbol:
    """El fondo completo en memoria: nodos visibles (el fondo y lo publicado),
    hijos por nodo y fechas extremas ya propagadas hacia arriba."""

    fondo: RecursoDocumental
    nodos: dict[uuid.UUID, RecursoDocumental] = field(default_factory=dict)
    hijos: dict[uuid.UUID, list[uuid.UUID]] = field(default_factory=lambda: defaultdict(list))
    fechas: dict[uuid.UUID, tuple[date, date]] = field(default_factory=dict)
    legibles: dict[uuid.UUID, list[str]] = field(default_factory=dict)

    def subarbol(self, raiz: uuid.UUID) -> list[uuid.UUID]:
        pila, vistos = [raiz], []
        while pila:
            n = pila.pop()
            vistos.append(n)
            pila.extend(self.hijos.get(n, []))
        return vistos

    def ancestros(self, nodo_id: uuid.UUID) -> list[RecursoDocumental]:
        cadena, actual = [], self.nodos.get(nodo_id)
        while actual is not None and actual.incluido_en_id and actual.id != self.fondo.id:
            actual = self.nodos.get(actual.incluido_en_id)
            if actual is not None:
                cadena.append(actual)
        return list(reversed(cadena))


def _orden(r: RecursoDocumental):
    return (NIVEL_DESCRIPCION.index(r.nivel), (r.codigo_referencia or "").lower(), r.titulo.lower())


def arbol(db: Session, fondo: RecursoDocumental) -> Arbol:
    a = Arbol(fondo=fondo)
    filas = db.scalars(select(RecursoDocumental).where(
        RecursoDocumental.fondo_id == fondo.id, RecursoDocumental.id != fondo.id,
        RecursoDocumental.publicado_en.is_not(None))).all()
    a.nodos[fondo.id] = fondo
    for r in filas:
        a.nodos[r.id] = r
    for r in sorted(filas, key=_orden):
        padre = r.incluido_en_id if r.incluido_en_id in a.nodos else fondo.id
        a.hijos[padre].append(r.id)
    # Fechas de creación (RiC-R080 is creation date of) normalizadas, de
    # cada descripción, propagadas a sus niveles superiores.
    # Se usan los límites del intervalo EDTF (c. 1948 cubre todo 1948); una
    # fecha sin límites conocidos (año desconocido) no cuenta.
    from app.servicios import fechas as servicio_fechas

    propias: dict[uuid.UUID, list[date]] = defaultdict(list)
    for destino, inicio, fin, normalizada, edtf, expresion in db.execute(
            select(Relacion.destino_id, Fecha.inicio, Fecha.fin, Fecha.normalizada, Fecha.edtf, Fecha.expresion)
            .join(Fecha, Fecha.id == Relacion.origen_id)
            .where(Relacion.codigo_ric == "is_creation_date_of", Relacion.estado == "vigente",
                   Relacion.destino_id.in_(list(a.nodos)))).all():
        limites = [d for d in (inicio or normalizada, fin or normalizada) if d]
        propias[destino].extend(limites)
        if limites:
            a.legibles.setdefault(destino, []).append(servicio_fechas.legible(edtf, expresion))

    def recorrer(n: uuid.UUID) -> list[date]:
        todas = list(propias.get(n, []))
        for h in a.hijos.get(n, []):
            todas.extend(recorrer(h))
        if todas:
            a.fechas[n] = (min(todas), max(todas))
        return todas

    recorrer(fondo.id)
    return a


def _fechas_texto(a: Arbol, r: RecursoDocumental) -> str | None:
    # Un documento con una sola fecha propia: tal como se precisó (c. 1948).
    if len(a.legibles.get(r.id, [])) == 1 and not a.hijos.get(r.id):
        return a.legibles[r.id][0]
    if r.id in a.fechas:
        ini, fin = a.fechas[r.id]
        if (ini.month, ini.day, fin.month, fin.day) == (1, 1, 12, 31):  # años completos: no se inventa el día
            return str(ini.year) if ini.year == fin.year else f"{ini.year} – {fin.year}"
        return _fecha_corta(ini) if ini == fin else f"{_fecha_corta(ini)} – {_fecha_corta(fin)}"
    return r.fechas_extremas


def _fecha_corta(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _nodo(a: Arbol, r: RecursoDocumental) -> dict:
    unidades = sum(1 for n in a.subarbol(r.id) if a.nodos[n].nivel == "unidad_documental")
    return {"id": str(r.id), "nivel": r.nivel, "titulo": r.titulo, "codigo_referencia": r.codigo_referencia,
            "fechas_extremas": _fechas_texto(a, r), "hijos": len(a.hijos.get(r.id, [])), "unidades_documentales": unidades}


def nivel(db: Session, fondo: RecursoDocumental, nodo_id: uuid.UUID | None = None) -> dict:
    """Un nivel del árbol para la navegación por migas de pan."""
    a = arbol(db, fondo)
    nodo_id = nodo_id or fondo.id
    if nodo_id not in a.nodos:
        raise ErrorInstrumento("Ese nivel no existe en el fondo o no tiene una descripción publicada.", 404)
    actual = a.nodos[nodo_id]
    return sin_campos_internos({
        "fondo": {"id": str(fondo.id), "titulo": fondo.titulo},
        "migas": [{"id": str(x.id), "nivel": x.nivel, "titulo": x.titulo} for x in a.ancestros(nodo_id)],
        "actual": _nodo(a, actual),
        "hijos": [_nodo(a, a.nodos[h]) for h in a.hijos.get(nodo_id, [])],
    })


# --- Ficha de consulta --------------------------------------------------------------------------


def preservacion(db: Session, instanciacion_id: str) -> dict:
    """Estado de preservación en solo lectura, tal como lo mantiene el
    módulo 5: integridad, riesgo de formato y migraciones hechas."""
    from app.servicios import preservacion as modulo

    inst = db.get(Instanciacion, uuid.UUID(instanciacion_id))
    if inst is None:
        return {"estado": "sin_evaluar"}
    r = modulo.riesgo_de(db, inst)
    if inst.estado_integridad in ("alterada", "ausente"):
        estado = "alerta_integridad"
    elif r["nivel"] != "bajo" and not r["mitigado_por"]:
        estado = "riesgo_obsolescencia"
    elif inst.estado_integridad == "sin_verificar":
        estado = "sin_verificar"
    else:
        estado = "buen_estado"
    return {"estado": estado, "integridad": inst.estado_integridad, "ultima_verificacion_en": inst.ultima_verificacion_en,
            "riesgo": r["nivel"], "formato": inst.formato_nombre, "puid": inst.formato_puid,
            "algoritmo_huella": inst.algoritmo_huella, "huella": inst.huella,
            "derivada_de": str(inst.derivada_de_id) if inst.derivada_de_id else None}


def ficha(db: Session, recurso: RecursoDocumental) -> dict:
    fondo = db.get(RecursoDocumental, recurso.fondo_id)
    a = arbol(db, fondo)
    if recurso.id not in a.nodos:
        raise ErrorInstrumento("Esa descripción no está publicada.", 404)
    base = consulta.ficha_publica(db, recurso)
    ids = [uuid.UUID(e["entidad_id"]) for e in base["entidades"]]
    forma = db.get(EntidadVocabulario, recurso.forma_documental_id) if recurso.forma_documental_id else None
    if forma:
        ids.append(forma.id)
    conexiones = vocabulario.conexiones_de(db, ids)
    entidades = []
    for e in base["entidades"]:
        propia = uuid.UUID(e["entidad_id"])
        en_vocabulario = e["tipo"] in ("agente", "lugar", "forma_documental", "actividad", "tipo_actividad", "mandato")
        entidades.append({**e, "en_vocabulario": en_vocabulario,
                          "documentos": conexiones.get(propia) if en_vocabulario else None})
    return sin_campos_internos({
        **{k: v for k, v in base.items() if k not in ("forma_documental",)},
        "entidades": entidades,
        "forma_documental": {"id": str(forma.id), "nombre": forma.nombre, "documentos": conexiones[forma.id]} if forma else None,
        "migas": [{"id": str(x.id), "nivel": x.nivel, "titulo": x.titulo} for x in a.ancestros(recurso.id)],
        "fechas_extremas": _fechas_texto(a, recurso),
        "control": {c: getattr(recurso, c) for c in ("codigo_referencia", "caja", "carpeta", "folios", "soporte")},
        "instanciaciones": [{**i, "preservacion": preservacion(db, i["id"])} for i in base["instanciaciones"]],
        "hijos": len(a.hijos.get(recurso.id, [])),
    })


# --- Inventario documental (FUID) ---------------------------------------------------------------

# Columnas del Formato Único de Inventario Documental (Acuerdo 042 de 2002
# del AGN), en el orden del formato. Las obligatorias se marcan pendientes
# si faltan; las demás pueden quedar vacías.
COLUMNAS_FUID = [
    ("numero_orden", "N.º de orden", True),
    ("codigo", "Código", True),
    ("nombre", "Nombre de la serie, subserie o asunto", True),
    ("fecha_inicial", "Fecha inicial", True),
    ("fecha_final", "Fecha final", True),
    ("caja", "Caja", True),
    ("carpeta", "Carpeta", True),
    ("folios", "N.º de folios", True),
    ("soporte", "Soporte", True),
    ("notas", "Notas", False),
]
PENDIENTE = "Pendiente"


def _nombre_fila(a: Arbol, r: RecursoDocumental) -> str:
    agrupaciones = [x.titulo for x in a.ancestros(r.id) if x.nivel in ("serie", "subserie")]
    return " / ".join(agrupaciones + [r.titulo])


def _folios(a: Arbol, r: RecursoDocumental) -> int | None:
    if r.folios is not None:
        return r.folios
    # Un expediente sin folios propios suma los de sus unidades, solo si
    # todas los tienen (una suma parcial sería un dato falso).
    unidades = [a.nodos[n] for n in a.subarbol(r.id) if n != r.id and a.nodos[n].nivel == "unidad_documental"]
    if unidades and all(u.folios is not None for u in unidades):
        return sum(u.folios for u in unidades)
    return None


def filas_inventario(db: Session, recurso: RecursoDocumental) -> dict:
    """Un renglón por unidad documental, o por expediente cuando el
    expediente se describió como un todo (sin unidades documentales
    publicadas dentro)."""
    fondo = db.get(RecursoDocumental, recurso.fondo_id or recurso.id)
    a = arbol(db, fondo)
    if recurso.id not in a.nodos:
        raise ErrorInstrumento("Ese nivel no está publicado.", 404)
    renglones = []
    for n in _en_orden(a, recurso.id):
        r = a.nodos[n]
        es_fila = r.nivel == "unidad_documental" or (
            r.nivel == "expediente" and not any(a.nodos[x].nivel == "unidad_documental" for x in a.subarbol(n) if x != n))
        if not es_fila:
            continue
        ini, fin = a.fechas.get(n, (None, None))
        valores = {
            "codigo": r.codigo_referencia, "nombre": _nombre_fila(a, r),
            "fecha_inicial": _fecha_corta(ini) if ini else None, "fecha_final": _fecha_corta(fin) if fin else None,
            "caja": r.caja, "carpeta": r.carpeta, "folios": _folios(a, r), "soporte": r.soporte, "notas": r.nota,
        }
        valores["numero_orden"] = len(renglones) + 1
        pendientes = [c for c, _, obligatoria in COLUMNAS_FUID if obligatoria and valores.get(c) in (None, "")]
        renglones.append({"id": str(r.id), "nivel": r.nivel, "valores": valores, "pendientes": pendientes})
    por_campo = defaultdict(int)
    for f in renglones:
        for c in f["pendientes"]:
            por_campo[c] += 1
    return sin_campos_internos({
        "nivel": {"id": str(recurso.id), "nivel": recurso.nivel, "titulo": recurso.titulo},
        "fondo": {"id": str(fondo.id), "titulo": fondo.titulo},
        "columnas": [{"clave": c, "nombre": nombre, "obligatoria": ob} for c, nombre, ob in COLUMNAS_FUID],
        "filas": renglones,
        "pendientes": sum(por_campo.values()),
        "pendientes_por_campo": dict(por_campo),
    })


def _en_orden(a: Arbol, raiz: uuid.UUID) -> list[uuid.UUID]:
    salida = []

    def ir(n):
        salida.append(n)
        for h in a.hijos.get(n, []):
            ir(h)

    ir(raiz)
    return salida


def alerta_pendientes(db: Session, inventario: dict) -> dict | None:
    """Panel central de alertas: una alerta pendiente por nivel del árbol,
    con el conteo de la última generación. Sin pendientes, no se crea
    ninguna, y la que hubiera queda atendida por el propio sistema."""
    nivel_id = inventario["nivel"]["id"]
    total = inventario["pendientes"]
    existente = alertas.pendiente(db, "inventario_campos_pendientes", nivel_id)
    if total == 0:
        if existente is not None:
            alertas.atender(db, existente, None, "Resuelta: la última generación del inventario ya no tiene pendientes.")
        return None
    mensaje = (f"Inventario de «{inventario['nivel']['titulo']}»: {total} campo{'s' if total != 1 else ''} "
               f"obligatorio{'s' if total != 1 else ''} del FUID sin dato.")
    detalle = {"pendientes": total, "por_campo": inventario["pendientes_por_campo"],
               "renglones": len(inventario["filas"])}
    if existente is not None:
        existente.mensaje, existente.detalle = mensaje, detalle
        alerta = existente
    else:
        alerta = alertas.crear(db, tipo="inventario_campos_pendientes", severidad="media", modulo="instrumentos",
                               entidad_tipo="recurso_documental", entidad_id=nivel_id, mensaje=mensaje,
                               fondo_id=uuid.UUID(inventario["fondo"]["id"]), detalle=detalle)
    db.flush()
    return {"id": str(alerta.id), "mensaje": mensaje}


def inventario_xlsx(inventario: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario"
    fino = Side(style="thin", color="9A9A9A")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    columnas = inventario["columnas"]
    ultima = get_column_letter(len(columnas))

    hoja.merge_cells(f"A1:{ultima}1")
    hoja["A1"] = "FORMATO ÚNICO DE INVENTARIO DOCUMENTAL"
    hoja["A1"].font = Font(bold=True, size=13)
    hoja["A1"].alignment = Alignment(horizontal="center")
    encabezado = [
        ("Fondo", inventario["fondo"]["titulo"]),
        ("Nivel inventariado", f"{NIVEL_NOMBRE[inventario['nivel']['nivel']]}: {inventario['nivel']['titulo']}"),
        ("Objeto", "Inventario documental"),
        ("Fecha de elaboración", date.today().strftime("%d/%m/%Y")),
        ("Campos pendientes", str(inventario["pendientes"])),
    ]
    for i, (k, v) in enumerate(encabezado, start=2):
        hoja.cell(row=i, column=1, value=k).font = Font(bold=True)
        hoja.merge_cells(start_row=i, start_column=2, end_row=i, end_column=len(columnas))
        hoja.cell(row=i, column=2, value=v)

    fila_titulos = len(encabezado) + 3
    relleno_titulo = PatternFill("solid", fgColor="DDE3EA")
    for j, col in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_titulos, column=j, value=col["nombre"])
        celda.font, celda.fill, celda.border = Font(bold=True), relleno_titulo, borde
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    relleno_pendiente = PatternFill("solid", fgColor="FFF1D6")
    fuente_pendiente = Font(italic=True, color="9A5B00")
    for i, fila in enumerate(inventario["filas"], start=fila_titulos + 1):
        for j, col in enumerate(columnas, start=1):
            valor = fila["valores"].get(col["clave"])
            celda = hoja.cell(row=i, column=j)
            celda.border = borde
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            if col["clave"] in fila["pendientes"]:
                celda.value, celda.fill, celda.font = PENDIENTE, relleno_pendiente, fuente_pendiente
                celda.comment = Comment("Falta dato: complételo en la descripción.", "RICORA")
            else:
                celda.value = valor

    anchos = {"numero_orden": 8, "codigo": 16, "nombre": 48, "fecha_inicial": 13, "fecha_final": 13,
              "caja": 9, "carpeta": 10, "folios": 10, "soporte": 14, "notas": 30}
    for j, col in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(j)].width = anchos.get(col["clave"], 14)
    hoja.freeze_panes = hoja.cell(row=fila_titulos + 1, column=1)
    hoja.print_title_rows = f"{fila_titulos}:{fila_titulos}"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# --- Guía del fondo -----------------------------------------------------------------------------

INSTRUCCION_GUIA = (
    "Eres archivista. Redacta en español, en prosa formal y sin viñetas, la nota de presentación de la guía "
    "de un fondo de archivo histórico (2 a 4 párrafos). Usa ÚNICAMENTE los datos que se te entregan, que ya "
    "fueron validados por una archivista. No inventes fechas, cifras, instituciones, personas ni alcance "
    "documental que no estén en los datos. Si un dato no está, no lo menciones. No uses títulos ni markdown."
)


def datos_guia(db: Session, fondo: RecursoDocumental) -> dict:
    """Solo datos validados: los niveles superiores publicados, conteos,
    fechas extremas y los agentes y lugares más citados del vocabulario."""
    a = arbol(db, fondo)
    superiores = [a.nodos[n] for n in _en_orden(a, fondo.id)
                  if a.nodos[n].nivel in ("fondo", "seccion", "serie", "subserie")]
    conteo = defaultdict(int)
    for n in a.nodos.values():
        conteo[n.nivel] += 1
    # Los mecanismos (los programas que actúan en el sistema: el motor,
    # Siegfried, Ghostscript) no son puntos de acceso del contenido del
    # fondo: se administran en Vocabularios, pero no van al índice.
    activas = db.scalars(select(EntidadVocabulario).where(
        EntidadVocabulario.fondo_id == fondo.id, EntidadVocabulario.estado == "activa",
        or_(EntidadVocabulario.subtipo.is_(None), EntidadVocabulario.subtipo != "mecanismo"))).all()
    conexiones = vocabulario.conexiones_de(db, [e.id for e in activas])
    principales = defaultdict(list)
    for e in sorted(activas, key=lambda e: -conexiones[e.id]):
        if conexiones[e.id] and len(principales[e.clase]) < 6:
            n = conexiones[e.id]
            principales[e.clase].append(f"{e.nombre} ({n} documento{'s' if n != 1 else ''})")
    fechas = a.fechas.get(fondo.id)
    return {
        "fondo": fondo.titulo,
        "fechas_extremas": f"{fechas[0].year}–{fechas[1].year}" if fechas else fondo.fechas_extremas,
        "niveles": [{"nivel": NIVEL_NOMBRE[x.nivel], "titulo": x.titulo, "alcance_y_contenido": x.alcance_contenido}
                    for x in superiores],
        "conteos": {(NIVEL_NOMBRE[k] if v == 1 else NIVEL_PLURAL[k].capitalize()): v
                    for k, v in sorted(conteo.items(), key=lambda x: NIVEL_DESCRIPCION.index(x[0])) if k != "fondo"},
        "agentes_principales": principales.get("agente", []),
        "lugares_principales": principales.get("lugar", []),
        "formas_documentales": principales.get("forma_documental", []),
    }


def _borrador_sin_motor(datos: dict) -> str:
    partes = [f"El fondo {datos['fondo']}"
              + (f", con fechas extremas {datos['fechas_extremas']}," if datos["fechas_extremas"] else "")
              + " reúne la documentación descrita en este sistema."]
    if datos["conteos"]:
        partes.append("Comprende " + ", ".join(f"{v} {k.lower()}" for k, v in datos["conteos"].items()) + ".")
    for n in datos["niveles"]:
        if n["alcance_y_contenido"]:
            partes.append(f"{n['nivel']} «{n['titulo']}»: {n['alcance_y_contenido']}")
    if datos["agentes_principales"]:
        partes.append("Entre sus productores y corresponsales figuran " + "; ".join(datos["agentes_principales"]) + ".")
    return "\n\n".join(partes)


def redactar_guia(db: Session, fondo: RecursoDocumental) -> dict:
    import json

    datos = datos_guia(db, fondo)
    m = motor.motor_activo()
    if m is not None and hasattr(m, "redactar"):
        try:
            texto = m.redactar(INSTRUCCION_GUIA, "Datos validados del fondo:\n" + json.dumps(datos, ensure_ascii=False, indent=1))
            if texto:
                return {"texto": texto, "redactado_por_motor": True, "aviso": None}
        except motor.MotorError as exc:
            aviso = f"{exc} Se armó un borrador básico con los datos validados."
            return {"texto": _borrador_sin_motor(datos), "redactado_por_motor": False, "aviso": aviso}
    return {"texto": _borrador_sin_motor(datos), "redactado_por_motor": False,
            "aviso": "No hay motor de análisis configurado: se armó un borrador básico con los datos validados."}


def guia_docx(db: Session, fondo: RecursoDocumental, texto: str) -> bytes:
    """El documento lleva exactamente el texto que dejó el archivista, más
    un cuadro de datos del fondo tomado del sistema."""
    from docx import Document
    from docx.shared import Pt

    datos = datos_guia(db, fondo)
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Calibri", Pt(11)
    doc.add_heading(f"Guía del fondo {fondo.titulo}", level=0)
    doc.add_heading("Nota de presentación", level=1)
    for parrafo in [p.strip() for p in texto.replace("\r\n", "\n").split("\n\n")]:
        if parrafo:
            doc.add_paragraph(parrafo)
    doc.add_heading("Datos del fondo", level=1)
    tabla = doc.add_table(rows=0, cols=2)
    tabla.style = "Light Grid Accent 1"
    filas = [("Título", fondo.titulo), ("Fechas extremas", datos["fechas_extremas"] or "—")]
    filas += [(k, str(v)) for k, v in datos["conteos"].items()]
    for k, v in filas:
        c = tabla.add_row().cells
        c[0].text, c[1].text = k, v
    if datos["niveles"][1:]:
        doc.add_heading("Estructura", level=1)
        for n in datos["niveles"][1:]:
            doc.add_paragraph(f"{n['nivel']}: {n['titulo']}", style="List Bullet")
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


# --- Índice de términos -------------------------------------------------------------------------


def indice(db: Session, fondo: RecursoDocumental) -> dict:
    # Los mecanismos (los programas que actúan en el sistema: el motor,
    # Siegfried, Ghostscript) no son puntos de acceso del contenido del
    # fondo: se administran en Vocabularios, pero no van al índice.
    activas = db.scalars(select(EntidadVocabulario).where(
        EntidadVocabulario.fondo_id == fondo.id, EntidadVocabulario.estado == "activa",
        or_(EntidadVocabulario.subtipo.is_(None), EntidadVocabulario.subtipo != "mecanismo"))).all()
    conexiones = vocabulario.conexiones_de(db, [e.id for e in activas])
    grupos = []
    for clase in ("agente", "lugar", "forma_documental", "actividad", "tipo_actividad", "mandato"):
        entidades = sorted((e for e in activas if e.clase == clase), key=lambda e: (e.nombre_normalizado, e.nombre))
        letras = defaultdict(list)
        for e in entidades:
            inicial = (e.nombre_normalizado[:1] or "#").upper()
            letras[inicial if inicial.isalpha() else "#"].append(
                {"id": str(e.id), "nombre": e.nombre, "subtipo": e.subtipo, "documentos": conexiones[e.id]})
        grupos.append({"clase": clase, "total": len(entidades),
                       "letras": [{"letra": k, "entidades": v} for k, v in sorted(letras.items())]})
    return sin_campos_internos({"fondo": {"id": str(fondo.id), "titulo": fondo.titulo}, "grupos": grupos})


# --- Grafo (vista visual de solo lectura) ---------------------------------------------------------

ETIQUETA_RELACION = {
    "has_creator": "producido por", "has_sender": "remitido por", "has_addressee": "dirigido a",
    "has_or_had_subject": "trata de", "is_creation_date_of": "fecha de creación de", "documents": "documenta",
    "includes_or_included": "incluye", "has_or_had_instantiation": "tiene instanciación",
    "migrated_into": "migrada a", "forma_documental": "forma documental",
    "has_activity_type": "es del tipo", "performs_or_performed": "ejerce", "regulates_or_regulated": "regula",
    "authorizes": "autoriza a", "is_date_associated_with": "fecha de", "has_or_had_subordinate": "tiene como subordinado",
    "has_successor": "tiene como sucesor", "is_agent_associated_with_agent": "asociado con",
    # Versiones 3 de descripción y 2 de vocabularios
    "has_or_had_holder": "custodiado por", "precedes_or_preceded": "precede a", "has_or_had_constituent": "tiene la parte",
    "has_direct_subevent": "tiene la sub-actividad", "contains_or_contained": "contiene",
    "is_or_was_location_of": "lugar de", "issued_by": "expedido por", "is_related_to": "produce la serie",
    "affects_or_affected": "afecta a", "occupies_or_occupied": "ocupa",
}
MAX_NODOS_GRAFO = 150
TIPOS_NODO_GRAFO = ("recurso_documental", "entidad_vocabulario", "fecha", "actividad", "instanciacion")


def _nodo_grafo(db: Session, tipo: str, nodo_id: uuid.UUID, fondo_id: uuid.UUID) -> dict | None:
    """Datos públicos de un nodo (lista cerrada; nada de procedencia).
    None si no pertenece al fondo o no es visible (borrador, fusionada)."""
    from app.models.descripcion import Actividad

    if tipo == "recurso_documental":
        r = db.get(RecursoDocumental, nodo_id)
        if r is None or (r.fondo_id or r.id) != fondo_id or (r.nivel != "fondo" and r.publicado_en is None):
            return None
        return {"tipo": tipo, "clase": r.nivel, "etiqueta": r.titulo, "subtitulo": NIVEL_NOMBRE[r.nivel]}
    if tipo == "entidad_vocabulario":
        e = db.get(EntidadVocabulario, nodo_id)
        if e is None or e.fondo_id != fondo_id or e.estado != "activa":
            return None
        return {"tipo": tipo, "clase": e.clase, "etiqueta": e.nombre, "subtitulo": e.subtipo}
    if tipo == "fecha":
        f = db.get(Fecha, nodo_id)
        from app.servicios import fechas

        return {"tipo": tipo, "clase": "fecha", "etiqueta": fechas.legible(f.edtf, f.expresion),
                "subtitulo": f.edtf} if f else None
    if tipo == "actividad":
        a = db.get(Actividad, nodo_id)
        return {"tipo": tipo, "clase": "actividad", "etiqueta": a.nombre, "subtitulo": None} \
            if a and a.fondo_id == fondo_id else None
    if tipo == "instanciacion":
        i = db.get(Instanciacion, nodo_id)
        if i is None or i.fondo_id != fondo_id or i.estado != "listo_para_descripcion":
            return None
        return {"tipo": tipo, "clase": "instanciacion", "etiqueta": i.nombre_original, "subtitulo": i.formato_puid}
    return None


def grafo(db: Session, fondo: RecursoDocumental, centro_tipo: str | None, centro_id: uuid.UUID | None,
          profundidad: int = 1) -> dict:
    """Vecindario de un nodo hasta `profundidad` saltos, leído de las mismas
    relaciones RiC que guardó la descripción. Por defecto, el fondo."""
    from app.models.enums import URI_RICO

    centro_tipo, centro_id = (centro_tipo, centro_id) if centro_id else ("recurso_documental", fondo.id)
    if centro_tipo not in TIPOS_NODO_GRAFO:
        raise ErrorInstrumento("Tipo de nodo desconocido.")
    inicial = _nodo_grafo(db, centro_tipo, centro_id, fondo.id)
    if inicial is None:
        raise ErrorInstrumento("Ese nodo no existe en el fondo o no es visible.", 404)
    nodos = {f"{centro_tipo}:{centro_id}": {**inicial, "id": str(centro_id)}}
    aristas: dict[tuple, dict] = {}
    frontera, truncado = [(centro_tipo, centro_id)], False
    for _ in range(max(1, min(profundidad, 3))):
        siguiente = []
        ids = [i for _, i in frontera]
        relaciones = db.scalars(select(Relacion).where(
            Relacion.estado == "vigente", (Relacion.origen_id.in_(ids)) | (Relacion.destino_id.in_(ids)))).all()
        vecinos = [(r.origen_tipo, r.origen_id, r.destino_tipo, r.destino_id, r.codigo_ric) for r in relaciones]
        # La forma documental es un atributo (RiC-A17) que se dibuja como arista.
        for t, i in frontera:
            if t == "recurso_documental":
                r = db.get(RecursoDocumental, i)
                if r is not None and r.forma_documental_id:
                    vecinos.append((t, i, "entidad_vocabulario", r.forma_documental_id, "forma_documental"))
                # La jerarquía también desde el árbol (incluido_en), como la recorre el catálogo.
                if r is not None and r.incluido_en_id and r.id != fondo.id:
                    vecinos.append((t, r.incluido_en_id, t, i, "includes_or_included"))
                for hijo in db.scalars(select(RecursoDocumental.id).where(RecursoDocumental.incluido_en_id == i,
                                                                          RecursoDocumental.id != i)):
                    vecinos.append((t, i, t, hijo, "includes_or_included"))
            if t == "entidad_vocabulario":
                for rid in db.scalars(select(RecursoDocumental.id).where(RecursoDocumental.forma_documental_id == i)):
                    vecinos.append(("recurso_documental", rid, t, i, "forma_documental"))
        for ot, oi, dt, di, codigo in vecinos:
            claves = []
            for t, i in ((ot, oi), (dt, di)):
                clave = f"{t}:{i}"
                if clave not in nodos:
                    if len(nodos) >= MAX_NODOS_GRAFO:
                        truncado = True
                        break
                    datos = _nodo_grafo(db, t, i, fondo.id)
                    if datos is None:
                        break
                    nodos[clave] = {**datos, "id": str(i)}
                    siguiente.append((t, i))
                claves.append(clave)
            if len(claves) == 2:
                aristas[(claves[0], claves[1], codigo)] = {
                    "desde": claves[0], "hacia": claves[1], "codigo_ric": codigo,
                    "uri_rico": URI_RICO.get(codigo) or ("rico:hasOrHadDocumentaryFormType" if codigo == "forma_documental" else None),
                    "etiqueta": ETIQUETA_RELACION.get(codigo, codigo.replace("_", " "))}
        frontera = siguiente
        if not frontera:
            break
    conexiones = vocabulario.conexiones_de(db, [uuid.UUID(n["id"]) for n in nodos.values() if n["tipo"] == "entidad_vocabulario"])
    for clave, n in nodos.items():
        n["clave"] = clave
        if n["tipo"] == "entidad_vocabulario":
            n["documentos"] = conexiones.get(uuid.UUID(n["id"]), 0)
    return sin_campos_internos({"fondo": {"id": str(fondo.id), "titulo": fondo.titulo},
                                "centro": f"{centro_tipo}:{centro_id}", "nodos": list(nodos.values()),
                                "aristas": list(aristas.values()), "truncado": truncado})
