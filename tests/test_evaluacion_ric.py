"""
Evaluación ciega en RiC (versión 2 del módulo de evaluación).

- La descripción se hace con entidades y relaciones de RiC-O; la propuesta
  del motor se convierte a lo mismo con la regla de la publicación.
- Se miden las relaciones (por propiedad) y la confusión de propiedades.
- Cada descripción se revisa contra RiC-O 1.1: el mapeo único y el OWL
  oficial. Una descripción que no es RiC válida no se envía.
"""

import io

from openpyxl import load_workbook

from app.servicios import evaluacion_ric as ric
from app.servicios import ric_o
from tests.test_evaluacion import PROPUESTA, anotar, evaluadora, motor_eval, preparar  # noqa: F401
from tests.test_descripcion import fondo  # noqa: F401


def _ric(entidades, relaciones, titulo="Oficio 210 de 1948"):
    return {"titulo": titulo, "entidades": entidades, "relaciones": relaciones}


ALCALDIA = {"clave": "e1", "tipo": "agente", "subtipo": "entidad_corporativa", "valor": "Alcaldía Municipal de Tunja"}
CONCEJO = {"clave": "e2", "tipo": "agente", "subtipo": "entidad_corporativa", "valor": "Concejo Municipal"}
FECHA = {"clave": "e3", "tipo": "fecha", "valor": "1948", "edtf": "1948"}


# --- Conversión de la propuesta del motor ----------------------------------------------------------


def test_la_propuesta_del_motor_se_convierte_en_las_relaciones_que_crearia_la_publicacion():
    m = ric.desde_roles(PROPUESTA)
    assert [(r["origen"], r["propiedad"], r["destino"]) for r in m["relaciones"]] == [
        ("documento", "has_creator", "e1"), ("documento", "has_addressee", "e2"),
        ("documento", "has_or_had_subject", "e3"), ("e4", "is_creation_date_of", "documento"),
        ("documento", "forma_documental", "e5")]
    assert m["entidades"][0]["subtipo"] == "entidad_corporativa" and "rol" not in m["entidades"][0]


def test_una_entidad_con_dos_papeles_es_un_nodo_con_dos_relaciones_y_la_actividad_trae_su_cadena():
    propuesta = {"titulo": "Acta", "entidades": [
        {"clave": "p1", "tipo": "agente", "subtipo": "cargo", "rol": "productor", "valor": "Alcalde de Tunja"},
        {"clave": "p2", "tipo": "agente", "subtipo": "cargo", "rol": "remitente", "valor": "Alcalde de  Tunja"},
        {"clave": "p3", "tipo": "actividad", "valor": "Expedición de permisos", "tipo_clave": "p3t",
         "agente_clave": "p1", "mandato_clave": "p3m"},
        {"clave": "p3t", "tipo": "tipo_actividad", "valor": "Control de espectáculos"},
        {"clave": "p3m", "tipo": "mandato", "valor": "Acuerdo 12 de 1940"},
        {"clave": "p4", "tipo": "mandato", "valor": "Ley 4 de 1913"}]}
    m = ric.desde_roles(propuesta)
    assert sum(1 for e in m["entidades"] if e["tipo"] == "agente") == 1
    trip = {(r["origen"], r["propiedad"], r["destino"]) for r in m["relaciones"]}
    alcalde, act, tipo, mandato, ley = "e1", "e2", "e3", "e4", "e5"
    assert {("documento", "has_creator", alcalde), ("documento", "has_sender", alcalde),
            ("documento", "documents", act), (act, "has_activity_type", tipo),
            (alcalde, "performs_or_performed", act), (mandato, "regulates_or_regulated", act),
            (mandato, "authorizes", alcalde), ("documento", "has_or_had_subject", ley)} == trip
    assert ric.conformidad(m)["conforme"], ric.conformidad(m)["problemas"]


# --- Conformidad con RiC-O 1.1 ---------------------------------------------------------------------


def test_la_conformidad_detecta_lo_que_rico_no_admite_con_el_mapeo_y_con_el_owl_oficial():
    lugar = {"clave": "e4", "tipo": "lugar", "valor": "Tunja"}
    r = ric.conformidad(_ric([ALCALDIA, lugar], [{"origen": "e4", "propiedad": "has_creator", "destino": "documento"},
                                                  {"origen": "documento", "propiedad": "has_creator", "destino": "e1"}]))
    assert r["conforme"] is False and r["relaciones"] == 2 and r["relaciones_validas"] == 1
    assert any("no admite Place como origen" in p for p in r["problemas"])
    assert any(p.startswith("OWL oficial:") and "dominio" in p for p in r["problemas"])


def test_un_agente_sin_subtipo_vale_solo_donde_vale_cualquier_agente():
    agente = {"clave": "e1", "tipo": "agente", "valor": "Juan Pérez"}
    assert ric.conformidad(_ric([agente], [{"origen": "documento", "propiedad": "has_creator", "destino": "e1"}]))["conforme"]
    r = ric.conformidad(_ric([agente], [{"origen": "documento", "propiedad": "has_author", "destino": "e1"}]))
    assert not r["conforme"] and "indique el subtipo del agente" in r["problemas"][0]


def test_reglas_del_perfil_titulo_una_forma_y_ninguna_entidad_suelta():
    forma1 = {"clave": "f1", "tipo": "forma_documental", "valor": "Oficio"}
    forma2 = {"clave": "f2", "tipo": "forma_documental", "valor": "Carta"}
    r = ric.conformidad(_ric([forma1, forma2, CONCEJO],
                             [{"origen": "documento", "propiedad": "forma_documental", "destino": "f1"},
                              {"origen": "documento", "propiedad": "forma_documental", "destino": "f2"}], titulo=""))
    problemas = " | ".join(r["problemas"])
    assert "Falta el título" in problemas and "más de una forma documental" in problemas
    assert "«Concejo Municipal» no está relacionada con nada" in problemas


def test_el_catalogo_ofrece_solo_propiedades_del_mapeo_verificado_y_sus_extremos():
    for codigo in ric.PROPIEDADES:
        assert codigo == ric.FORMA or codigo in ric_o.PROPIEDADES
    cat = {c["codigo"]: c for c in ric.catalogo()}
    assert cat["has_creator"]["rico"] == "rico:hasCreator" and cat["has_creator"]["origen"] == ["documento"]
    assert cat["has_addressee"]["papel"] == "Destinatario" and set(ric.PAPEL) == set(ric.PROPIEDADES)
    assert "agente:persona" in cat["has_creator"]["destino"] and "lugar" not in cat["has_creator"]["destino"]
    assert "agente" not in cat["has_author"]["destino"]  # el autor exige saber qué clase de agente es
    assert cat["occupies_or_occupied"]["origen"] == ["agente:persona"]
    assert cat["forma_documental"]["rico"] == "rico:hasDocumentaryFormType"
    # Las relaciones estructurales no se describen leyendo un documento.
    assert "includes_or_included" not in cat and "has_or_had_instantiation" not in cat


# --- Comparación de relaciones ---------------------------------------------------------------------


def test_relacion_acierta_con_misma_propiedad_y_extremos_y_la_confusion_de_papel_se_cuenta_aparte():
    motor = _ric([ALCALDIA, CONCEJO, FECHA], [{"origen": "documento", "propiedad": "has_addressee", "destino": "e1"},
                                              {"origen": "documento", "propiedad": "has_addressee", "destino": "e2"},
                                              {"origen": "e3", "propiedad": "is_creation_date_of", "destino": "documento"}])
    archivista = _ric([ALCALDIA, CONCEJO, FECHA], [{"origen": "documento", "propiedad": "has_creator", "destino": "e1"},
                                                   {"origen": "documento", "propiedad": "has_addressee", "destino": "e2"},
                                                   {"origen": "e3", "propiedad": "is_creation_date_of",
                                                    "destino": "documento"}])
    cuentas, confusiones = ric.conteo_por_propiedad(ric.tripletas(motor), ric.tripletas(archivista), 1.0)
    assert cuentas["has_addressee"] == {"vp": 1, "fp": 1, "fn": 0}
    assert cuentas["has_creator"] == {"vp": 0, "fp": 0, "fn": 1}
    assert cuentas["is_creation_date_of"] == {"vp": 1, "fp": 0, "fn": 0}
    # La Alcaldía la encontró el motor, pero como destinataria: la archivista dijo productora.
    assert confusiones == [("has_addressee", "has_creator")]
    # Una relación con un extremo distinto no es acierto aunque la propiedad coincida.
    otra = _ric([ALCALDIA, {**FECHA, "edtf": "1949", "valor": "1949"}],
                [{"origen": "e3", "propiedad": "is_creation_date_of", "destino": "documento"}])
    assert ric.emparejar(ric.tripletas(motor), ric.tripletas(otra), 1.0)[0] == 0


# --- Flujo por la API ------------------------------------------------------------------------------


def test_la_archivista_describe_en_ric_y_no_envia_una_descripcion_que_rico_no_admite(cliente, db, fondo, cabeceras_admin,
                                                                                     motor_eval):
    ev_id, d = preparar(cliente, db, fondo, cabeceras_admin)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    luis = evaluadora(cliente, db, "luis@correo.com", "Luis Pardo")
    cat = cliente.get("/api/evaluacion/propiedades", headers=ana)
    assert cat.status_code == 200 and any(c["codigo"] == "documents" for c in cat.json())
    a = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "ciega"}).json()
    url = f"/api/evaluacion/anotaciones/{a['id']}"
    # Una propiedad que no es de RiC-O, o una relación hacia una entidad que no está: 422.
    assert cliente.put(url, headers=ana, json=_ric([ALCALDIA], [{"origen": "documento", "propiedad": "fue_hecho_por",
                                                                  "destino": "e1"}])).status_code == 422
    assert cliente.put(url, headers=ana, json=_ric([ALCALDIA], [{"origen": "documento", "propiedad": "has_creator",
                                                                  "destino": "e9"}])).status_code == 422
    # El borrador con una relación que RiC-O no admite se guarda y lo dice; enviarlo, no.
    mala = _ric([ALCALDIA, FECHA], [{"origen": "documento", "propiedad": "has_creator", "destino": "e1"},
                                    {"origen": "documento", "propiedad": "has_creator", "destino": "e3"}])
    r = cliente.put(url, headers=ana, json=mala)
    assert r.status_code == 200 and r.json()["conformidad"]["relaciones_validas"] == 1
    r = cliente.put(url, headers=ana, json=mala | {"enviar": True})
    assert r.status_code == 422 and "rico:hasCreator no admite Date como destino" in r.json()["detail"]
    buena = _ric([ALCALDIA, CONCEJO, FECHA], [{"origen": "documento", "propiedad": "has_creator", "destino": "e1"},
                                              {"origen": "documento", "propiedad": "has_creator", "destino": "e2"},
                                              {"origen": "e3", "propiedad": "is_creation_date_of", "destino": "documento"}])
    r = cliente.put(url, headers=ana, json=buena | {"enviar": True})
    assert r.status_code == 200 and r.json()["conformidad"]["conforme"] and len(r.json()["datos"]["relaciones"]) == 3
    # Luis anota con el formato anterior (roles): se convierte a RiC.
    assert anotar(cliente, luis, ev_id, d.id, "ciega", {"titulo": "Oficio", "entidades": [
        {"tipo": "agente", "rol": "productor", "valor": "Alcaldía Municipal de Tunja"},
        {"tipo": "fecha", "valor": "1948", "edtf": "1948"}]}).status_code == 200

    res = cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=cabeceras_admin).json()
    rel = res["relaciones"]
    # Ana dijo que el Concejo produjo el oficio; el motor, que lo recibió.
    assert {"motor": "has_addressee", "archivista": "has_creator", "veces": 1} in rel["confusiones"]
    assert rel["estricto"]["por_tipo"]["has_creator"]["vp"] == 2  # la Alcaldía, frente a Ana y frente a Luis
    assert res["conformidad"]["ciega"] == res["conformidad"]["ciega"] | {"descripciones": 2, "conformes": 2}
    hoja = cliente.get(f"/api/evaluacion/{ev_id}/resultados/hoja-de-calculo", headers=cabeceras_admin)
    libro = load_workbook(io.BytesIO(hoja.content))
    filas = [[c.value for c in f] for f in libro["Relaciones RiC"].iter_rows()]
    # Ana también dijo que el Concejo produjo el oficio: un falso negativo.
    assert ["estricto", "rico:hasCreator", 2, 0, 1, 1.0, 0.6667, 0.8] in filas
    assert any(f[:3] == [None, "rico:hasAddressee", "rico:hasCreator"] for f in filas)


def test_una_asistida_anulada_se_reabre_con_la_propuesta_del_motor(cliente, db, fondo, cabeceras_admin, motor_eval):
    ev_id, d = preparar(cliente, db, fondo, cabeceras_admin)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    a = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "asistida"}).json()
    assert cliente.post(f"/api/evaluacion/anotaciones/{a['id']}/anular", headers=cabeceras_admin).status_code == 200
    otra = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "asistida"}).json()
    assert otra["id"] == a["id"] and len(otra["datos"]["relaciones"]) == 5 and otra["conformidad"]["conforme"]
