"""Rediseño de navegación: historial reciente con exportación completa a
Excel, visor de documentos con nivel de acceso y datos de apoyo de las
vistas cortas (sin endpoints que cambien contratos existentes)."""

import io

import pytest
from openpyxl import load_workbook

from app.models.preservacion import DeclaracionDerechos
from app.servicios import formato
from app.servicios.auditoria import registrar
from tests import archivos
from tests.conftest import crear_usuario, ingresar
from tests.test_preservacion import archivista, describir, fondo, ingresar_archivo, jpeg  # noqa: F401


@pytest.fixture(scope="module", autouse=True)
def herramientas():
    formato.herramienta.cache_clear()
    formato.herramienta()


@pytest.fixture()
def consulta(cliente, db):
    crear_usuario(db, "lector@correo.com", "consulta")
    return ingresar(cliente, "lector@correo.com")


def _filas(contenido: bytes, hoja: str | None = None) -> list[tuple]:
    libro = load_workbook(io.BytesIO(contenido))
    h = libro[hoja] if hoja else libro.active
    return [f for f in h.iter_rows(min_row=2, values_only=True) if any(v is not None for v in f)]


# --- Historial reciente con exportación completa ---------------------------------------------------


def test_mi_trazabilidad_muestra_quince_y_exporta_todo_lo_filtrado(cliente, db, admin, cabeceras_admin):
    for n in range(22):
        registrar(db, modulo="vocabularios", accion="entidad_enriquecida", usuario_id=admin.id, entidad_tipo="entidad_vocabulario",
                  entidad_id=f"e-{n}", anterior={"nombre": f"Antes {n}"}, nuevo={"nombre": f"Después {n}"})
    for n in range(4):
        registrar(db, modulo="vocabularios", accion="entidad_fusionada", usuario_id=admin.id,
                  entidad_tipo="entidad_vocabulario", entidad_id=f"f-{n}")
    db.commit()
    r = cliente.get("/api/auditoria/mi-trazabilidad", headers=cabeceras_admin,
                    params={"accion": "entidad_enriquecida", "limite": 15}).json()
    assert len(r["eventos"]) == 15 and r["total"] == 22
    x = cliente.get("/api/auditoria/trazabilidad/exportar", headers=cabeceras_admin, params={"accion": "entidad_enriquecida"})
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    filas = _filas(x.content, "Trazabilidad")
    assert len(filas) == 22  # todo lo que cumple el filtro, no solo lo visible
    encabezado = next(load_workbook(io.BytesIO(x.content))["Trazabilidad"].iter_rows(max_row=1, values_only=True))
    assert "Identificador del evento" in encabezado and "Cambios (antes → después)" in encabezado
    assert any("nombre: Antes 3 → Después 3" in (f[10] or "") for f in filas)
    # Sin filtro, las 26 de esta persona.
    assert len(_filas(cliente.get("/api/auditoria/trazabilidad/exportar", headers=cabeceras_admin).content)) >= 26


def test_cada_quien_exporta_solo_su_trazabilidad(cliente, db, admin, archivista):
    registrar(db, modulo="ingesta", accion="documento_cargado", usuario_id=admin.id, entidad_tipo="instanciacion",
              entidad_id="de-la-administradora")
    db.commit()
    x = cliente.get("/api/auditoria/trazabilidad/exportar", headers=archivista)
    assert x.status_code == 200
    assert all(f[7] != "de-la-administradora" for f in _filas(x.content))


def test_panel_consolidado_y_decisiones_se_exportan_completos(cliente, db, admin, cabeceras_admin, archivista):
    x = cliente.get("/api/auditoria/panel-consolidado/exportar", headers=cabeceras_admin)
    assert x.status_code == 200
    libro = load_workbook(io.BytesIO(x.content))
    assert libro.sheetnames[0].startswith("Semana ") and "Sesiones" in libro.sheetnames
    assert any(f[0] == admin.nombre for f in _filas(x.content, libro.sheetnames[0]))
    d = cliente.get("/api/auditoria/decisiones-ia/exportar", headers=cabeceras_admin)
    assert d.status_code == 200 and load_workbook(io.BytesIO(d.content)).sheetnames == ["Resumen", "Decisiones"]
    # La vista muestra 25 por defecto (el servidor acepta el límite).
    assert cliente.get("/api/auditoria/decisiones-ia", headers=cabeceras_admin, params={"limite": 25}).status_code == 200
    for ruta in ("/api/auditoria/panel-consolidado/exportar", "/api/auditoria/decisiones-ia/exportar"):
        assert cliente.get(ruta, headers=archivista).status_code == 403
        assert cliente.get(ruta).status_code == 401


# --- Visor de documentos ---------------------------------------------------------------------------


def test_previsualizar_en_ingesta_y_descripcion_sin_entregar_el_original(cliente, db, fondo, archivista):
    inst = ingresar_archivo(cliente, db, archivista, fondo, "oficio.pdf", archivos.pdf_con_texto())
    for base in (f"/api/ingesta/{inst.id}/previsualizar", f"/api/descripcion/{inst.id}/previsualizar"):
        info = cliente.get(base, headers=archivista).json()
        assert info["admite"] is True and info["total"] == 1 and "ACTA" in (info["texto"] or "").upper()
        assert "ruta" not in info
        pagina = cliente.get(f"{base}/1", headers=archivista)
        assert pagina.status_code == 200 and pagina.headers["content-type"] == "image/png"
        assert pagina.content[:4] == b"\x89PNG"  # una imagen generada, no el PDF
        assert cliente.get(f"{base}/2", headers=archivista).status_code == 422
    assert cliente.get(f"/api/ingesta/{inst.id}/previsualizar").status_code == 401


def test_quien_no_es_archivista_no_previsualiza_lo_reservado(cliente, db, fondo, admin, archivista, consulta):
    inst = ingresar_archivo(cliente, db, archivista, fondo, "reservado.jpg", jpeg())
    recurso = describir(db, fondo, inst)
    db.add(DeclaracionDerechos(fondo_id=fondo.id, entidad_tipo="recurso_documental", entidad_id=recurso.id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()
    for base in (f"/api/ingesta/{inst.id}/previsualizar", f"/api/descripcion/{inst.id}/previsualizar",
                 f"/api/instrumentos/previsualizar/{inst.id}"):
        assert cliente.get(base, headers=consulta).status_code == 403, base
        assert cliente.get(f"{base}/1", headers=consulta).status_code == 403, base
    # El equipo de archivo sí lo ve, también en la ficha del catálogo.
    assert cliente.get(f"/api/instrumentos/previsualizar/{inst.id}", headers=archivista).json()["admite"] is True


def test_en_el_catalogo_solo_se_ve_el_archivo_de_una_descripcion_publicada(cliente, db, fondo, archivista, consulta):
    inst = ingresar_archivo(cliente, db, archivista, fondo, "sin_describir.jpg", jpeg())
    assert cliente.get(f"/api/instrumentos/previsualizar/{inst.id}", headers=consulta).status_code == 404
    describir(db, fondo, inst)
    assert cliente.get(f"/api/instrumentos/previsualizar/{inst.id}", headers=consulta).status_code == 200


# --- Datos de apoyo de las vistas cortas ------------------------------------------------------------


def test_actividad_reciente_y_resumen_de_ingesta(cliente, db, fondo, archivista):
    vacio = cliente.get("/api/ingesta/recientes", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert vacio == []  # la pantalla muestra entonces la guía de tres pasos
    insts = [ingresar_archivo(cliente, db, archivista, fondo, f"doc{n}.pdf", archivos.pdf_con_texto(f"Acta {n}"))
             for n in range(6)]
    describir(db, fondo, insts[-1])
    recientes = cliente.get("/api/ingesta/recientes", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert len(recientes) == 5 and recientes[0]["nombre"] == "doc5.pdf"
    assert recientes[0]["descrito_en"] == "Oficio 114" and recientes[1]["descrito_en"] is None
    resumen = cliente.get("/api/ingesta/resumen", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert resumen["documentos"] == 6 and resumen["con_texto"] == 6 and resumen["en_riesgo"] >= 0


def test_productividad_y_descritas_en_excel(cliente, db, fondo, admin, archivista):
    inst = ingresar_archivo(cliente, db, archivista, fondo, "oficio.pdf", archivos.pdf_con_texto())
    describir(db, fondo, inst)
    usuario = crear_usuario(db, "otra@correo.com", "archivista")
    otra = ingresar(cliente, "otra@correo.com")
    registrar(db, modulo="descripcion", accion="descripcion_publicada", usuario_id=usuario.id,
              entidad_tipo="recurso_documental", entidad_id="x")
    db.commit()
    p = cliente.get("/api/descripcion/productividad", headers=otra, params={"fondo_id": str(fondo.id)}).json()
    assert p == {"hoy_por_mi": 1, "total_fondo": 1}
    assert cliente.get("/api/descripcion/productividad", headers=archivista,
                       params={"fondo_id": str(fondo.id)}).json()["hoy_por_mi"] == 0
    x = cliente.get("/api/descripcion/publicadas/exportar", headers=archivista, params={"fondo_id": str(fondo.id)})
    assert x.status_code == 200 and [f[0] for f in _filas(x.content)] == ["Oficio 114"]


def test_linea_de_tiempo_de_preservacion(cliente, db, fondo, archivista):
    e = cliente.get("/api/preservacion/eventos-recientes", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert e == {"verificacion": None, "migracion": None, "restauracion": None, "simulacro_base_de_datos": None}
    inst = ingresar_archivo(cliente, db, archivista, fondo, "oficio.pdf", archivos.pdf_con_texto())
    assert cliente.post(f"/api/preservacion/instanciacion/{inst.id}/verificar", headers=archivista).status_code == 200
    e = cliente.get("/api/preservacion/eventos-recientes", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert e["verificacion"]["resultados"] == {"integra": 1}
