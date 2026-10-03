"""
Pruebas del Módulo 2 actualizado (secciones 11 y 12 del prompt): fechas en
EDTF con sus tres subtipos, la cadena documento → actividad → tipo de
actividad → mandato (RiC-R033, hasActivityType, R060, R063, R067), los seis
tipos reutilizables del vocabulario, las relaciones entre agentes, y un
evento de decisión de IA por cada propuesta del motor.
"""

import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.auditoria import RegistroAuditoria
from app.models.descripcion import EntidadVocabulario, Fecha, Relacion
from app.servicios import fechas, motor, vocabulario
from app.servicios.consulta import CAMPOS_PROHIBIDOS
from tests.conftest import crear_usuario, ingresar
from tests.test_descripcion import MotorDePrueba, aceptar_todo, archivista, claves, documento, fondo, iniciar  # noqa: F401

OFICIO = ("La Alcaldía Municipal, en ejercicio de la policía local que le confiere el Acuerdo 7 de 1946 del Concejo, "
          "informa al Señor Gobernador del Departamento que durante 1948 se vigilaron los mercados de Tunja. "
          "Expedido hacia 1948, con copia a las sesiones del 15 de enero y del 2 de marzo de 1948.")

RESPUESTA = {
    "titulo": "Oficio de la Alcaldía sobre la vigilancia de los mercados",
    "alcance_contenido": "La Alcaldía informa al Gobernador sobre la vigilancia de los mercados en 1948.",
    "confianza_alcance": 0.8,
    "entidades": [
        {"tipo": "agente", "subtipo": "entidad_corporativa", "rol": "productor", "valor": "Alcaldía Municipal",
         "fragmento": "La Alcaldía Municipal", "documento": 1, "confianza": 0.95},
        {"tipo": "agente", "subtipo": "cargo", "rol": "destinatario", "valor": "Gobernador del Departamento",
         "fragmento": "Señor Gobernador del Departamento", "documento": 1, "confianza": 0.9},
        {"tipo": "mandato", "subtipo": "acuerdo", "valor": "Acuerdo 7 de 1946", "edtf": "1946",
         "fragmento": "el Acuerdo 7 de 1946 del Concejo", "documento": 1, "confianza": 0.8},
        {"tipo": "actividad", "valor": "Vigilancia de mercados por la Alcaldía, 1948", "tipo_actividad": "Policía local",
         "ejercida_por": "Alcaldía Municipal", "mandato": "Acuerdo 7 de 1946", "edtf": "1948",
         "fragmento": "en ejercicio de la policía local", "documento": 1, "confianza": 0.75},
        {"tipo": "lugar", "valor": "Tunja", "fragmento": "los mercados de Tunja", "documento": 1, "confianza": 0.9},
        {"tipo": "fecha", "valor": "hacia 1948", "subtipo_fecha": "simple", "edtf": "1948~",
         "fragmento": "Expedido hacia 1948", "documento": 1, "confianza": 0.7},
        {"tipo": "forma_documental", "valor": "Oficio", "fragmento": "informa al Señor", "documento": 1, "confianza": 0.6},
    ],
}


@pytest.fixture()
def motor_contexto(monkeypatch):
    m = MotorDePrueba(RESPUESTA)
    monkeypatch.setattr(motor, "motor_activo", lambda: m)
    return m


def por_tipo(espacio, tipo):
    return [e for e in espacio["propuesta"]["entidades"] if e["tipo"] == tipo]


def decisiones(db, recurso_id):
    return db.scalars(select(RegistroAuditoria).where(RegistroAuditoria.accion == "decision_ia",
                                                      RegistroAuditoria.entidad_id == str(recurso_id))).all()


# --- Fechas EDTF -----------------------------------------------------------------------------


@pytest.mark.parametrize("edtf, subtipo, legible, inicio, fin", [
    ("1948-03-15", "simple", "15 de marzo de 1948", "1948-03-15", "1948-03-15"),
    ("1948~", "simple", "c. 1948", "1948-01-01", "1948-12-31"),
    ("194X", "simple", "década de 1940", "1940-01-01", "1949-12-31"),
    ("1948?/1952", "rango", "de 1948 (incierta) a 1952", "1948-01-01", "1952-12-31"),
    ("/1952", "rango", "hasta 1952 (inicio desconocido)", None, "1952-12-31"),
    ("{1948-01-15,1948-03-02}", "conjunto", "15 de enero de 1948 y 2 de marzo de 1948", "1948-01-15", "1948-03-02"),
])
def test_fecha_se_interpreta_y_normaliza(edtf, subtipo, legible, inicio, fin):
    i = fechas.interpretar(edtf, subtipo)
    assert (i.subtipo, i.legible) == (subtipo, legible)
    assert (i.inicio.isoformat() if i.inicio else None, i.fin.isoformat() if i.fin else None) == (inicio, fin)


@pytest.mark.parametrize("edtf", ["1948-21", "1948-02-30", "1952/1948", "../1952", "[1948,1949]", "Y170000002", "-0100"])
def test_fecha_fuera_del_subconjunto_se_rechaza(edtf):
    with pytest.raises(fechas.FechaInvalida):
        fechas.interpretar(edtf)


def test_fechas_aproximada_rango_y_conjunto_quedan_en_edtf(cliente, db, fondo, archivista, monkeypatch):
    monkeypatch.setattr(motor, "motor_activo", lambda: None)
    d = documento(db, fondo, "acta.pdf", OFICIO)
    espacio = iniciar(cliente, archivista, [d.id])
    entidades = [
        {"tipo": "fecha", "valor": "hacia 1948", "edtf": "1948~", "fecha_subtipo": "simple"},
        {"tipo": "fecha", "valor": "entre 1948 y 1952", "edtf": "1948/1952", "fecha_subtipo": "rango"},
        {"tipo": "fecha", "valor": "sesiones de enero y marzo", "edtf": "{1948-01-15,1948-03-02}", "fecha_subtipo": "conjunto"},
    ]
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json={
        "trabajo_id": espacio["trabajo_id"], "titulo": "Acta", "entidades": entidades})
    assert r.status_code == 201, r.text
    guardadas = {f.edtf: f for f in db.scalars(select(Fecha)).all()}
    assert {k: (f.subtipo, str(f.inicio), str(f.fin)) for k, f in guardadas.items()} == {
        "1948~": ("simple", "1948-01-01", "1948-12-31"),
        "1948/1952": ("rango", "1948-01-01", "1952-12-31"),
        "{1948-01-15,1948-03-02}": ("conjunto", "1948-01-15", "1948-03-02")}
    legibles = {e["edtf"]: e["fecha_legible"] for e in r.json()["entidades"]}
    assert legibles["1948~"] == "c. 1948" and legibles["1948/1952"] == "de 1948 a 1952"
    # Un subtipo que no corresponde a la expresión, o una fecha sin completar, no se publica.
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "b.pdf", OFICIO).id])
    for mala in ({"tipo": "fecha", "valor": "1948", "edtf": "1948/1950", "fecha_subtipo": "simple"},
                 {"tipo": "fecha", "valor": "mediados de siglo"}):
        assert cliente.post("/api/descripcion/publicar", headers=archivista, json={
            "trabajo_id": espacio["trabajo_id"], "titulo": "Acta", "entidades": [mala]}).status_code == 422


# --- Cadena documento → actividad → tipo → mandato -----------------------------------------------------


def test_la_actividad_queda_conectada_a_su_tipo_agente_y_mandato_y_se_ve_en_el_catalogo(cliente, db, fondo, archivista,
                                                                                         motor_contexto):
    d = documento(db, fondo, "Oficio_210_1948.pdf", OFICIO)
    espacio = iniciar(cliente, archivista, [d.id])
    [act] = por_tipo(espacio, "actividad")
    [tipo] = por_tipo(espacio, "tipo_actividad")  # el motor dio el tipo dentro de la actividad: queda como propuesta propia
    [mandato] = por_tipo(espacio, "mandato")
    [alcaldia] = [e for e in por_tipo(espacio, "agente") if e["rol"] == "productor"]
    assert (act["tipo_clave"], act["agente_clave"], act["mandato_clave"]) == (tipo["clave"], alcaldia["clave"],
                                                                              mandato["clave"])
    assert espacio["propuesta"]["version_prompt"] == motor.VERSION_PROMPT
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json=aceptar_todo(espacio))
    assert r.status_code == 201, r.text
    recurso_id = uuid.UUID(r.json()["id"])

    nodos = {e.clase: e for e in db.scalars(select(EntidadVocabulario)).all() if e.clase != "agente"}
    agentes = {e.nombre: e for e in db.scalars(select(EntidadVocabulario).where(EntidadVocabulario.clase == "agente"))}
    rel = {(x.origen_id, x.codigo_ric, x.destino_id) for x in db.scalars(select(Relacion).where(Relacion.estado == "vigente"))}
    a, t, m, alc = nodos["actividad"].id, nodos["tipo_actividad"].id, nodos["mandato"].id, agentes["Alcaldía Municipal"].id
    assert (recurso_id, "documents", a) in rel  # RiC-R033
    assert (a, "has_activity_type", t) in rel  # rico:hasActivityType
    assert (alc, "performs_or_performed", a) in rel  # RiC-R060i
    assert (m, "regulates_or_regulated", a) in rel  # RiC-R063
    assert (m, "authorizes", alc) in rel  # RiC-R067: del mandato al agente, como en RiC-O 1.1
    assert nodos["mandato"].subtipo == "acuerdo"
    # Un documento nunca se conecta directo al tipo de actividad.
    assert not any(o == recurso_id and dst == t for o, _, dst in rel)

    # Desde el catálogo se recupera como una sola cadena, sin campos internos.
    ficha = cliente.get(f"/api/instrumentos/catalogo/{recurso_id}", headers=archivista)
    assert ficha.status_code == 200, ficha.text
    [actividad] = [e for e in ficha.json()["entidades"] if e["tipo"] == "actividad"]
    cadena = actividad["contexto"]
    assert cadena["tipo_actividad"]["nombre"] == "Policía local"
    assert [x["nombre"] for x in cadena["ejercida_por"]] == ["Alcaldía Municipal"]
    [norma] = cadena["regulada_por"]
    assert (norma["nombre"], norma["subtipo"], norma["expedicion"]["edtf"]) == ("Acuerdo 7 de 1946", "acuerdo", "1946")
    assert cadena["periodo"]["fecha_legible"] == "1948"
    # Los conteos son de documentos, no de actividades ni mandatos: aquí hay uno solo.
    assert {e["valor"]: e["documentos"] for e in ficha.json()["entidades"] if e["en_vocabulario"]} == {
        "Alcaldía Municipal": 1, "Gobernador del Departamento": 1, "Tunja": 1,
        "Vigilancia de mercados por la Alcaldía, 1948": 1}
    assert vocabulario.conexiones_de(db, [t, m]) == {t: 1, m: 1}  # alcanzados a través de la actividad
    # La fecha del documento se muestra con su precisión en el catálogo, no como un día inventado.
    nodo = cliente.get("/api/instrumentos/catalogo", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert nodo["hijos"][0]["fechas_extremas"] == "c. 1948"
    publica = cliente.get(f"/api/catalogo/registros/{recurso_id}", headers=archivista).json()
    assert not set(claves(publica)) & CAMPOS_PROHIBIDOS
    assert [e for e in publica["entidades"] if e["tipo"] == "actividad"][0]["contexto"]["tipo_actividad"]

    # Otro documento de la misma actividad la reutiliza: la cadena no se duplica.
    d2 = documento(db, fondo, "Oficio_211_1948.pdf", OFICIO)
    espacio2 = iniciar(cliente, archivista, [d2.id])
    datos = aceptar_todo(espacio2)
    for e in [x for x in datos["entidades"] if x["tipo"] != "fecha"]:
        existente = db.scalar(select(EntidadVocabulario).where(EntidadVocabulario.nombre == e["valor"],
                                                               EntidadVocabulario.clase == e["tipo"]))
        if existente is not None:
            e |= {"reutilizar_id": str(existente.id), "crear_nueva": False}
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=datos).status_code == 201
    cadena_db = db.scalars(select(Relacion).where(Relacion.codigo_ric.in_(
        ("has_activity_type", "performs_or_performed", "regulates_or_regulated", "authorizes")))).all()
    assert len(cadena_db) == 4
    assert len(db.scalars(select(Relacion).where(Relacion.codigo_ric == "documents", Relacion.destino_id == a)).all()) == 2


def test_tipo_de_actividad_suelto_no_se_publica(cliente, db, fondo, archivista, motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    datos = aceptar_todo(espacio)
    datos["entidades"] = [e for e in datos["entidades"] if e["tipo"] != "actividad"]  # se descartó la actividad
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json=datos)
    assert r.status_code == 422 and "no está asignado a ninguna actividad" in r.json()["detail"]


def test_publicar_sin_actividad_ni_mandato_es_valido(cliente, db, fondo, archivista, motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    datos = aceptar_todo(espacio)
    datos["entidades"] = [e for e in datos["entidades"] if e["tipo"] not in ("actividad", "tipo_actividad", "mandato")]
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=datos).status_code == 201


# --- Seis tipos reutilizables del vocabulario --------------------------------------------------------------


@pytest.mark.parametrize("clase, subtipo", [("agente", "entidad_corporativa"), ("lugar", None),
                                            ("forma_documental", None), ("actividad", None),
                                            ("tipo_actividad", None), ("mandato", "decreto")])
def test_verificacion_de_vocabulario_reutilizar_y_crear_para_cada_tipo(cliente, db, fondo, archivista, admin, clase,
                                                                         subtipo):
    existente = vocabulario.crear(db, fondo_id=fondo.id, clase=clase, nombre="Policía Local de Tunja", subtipo=subtipo,
                                  origen="persona", confianza=None, motor=None, usuario_id=admin.id)
    db.commit()
    r = cliente.post("/api/descripcion/verificar-vocabulario", headers=archivista,
                     json={"fondo_id": str(fondo.id), "tipo": clase, "valor": "policia local de Tunja"})
    assert r.status_code == 200 and r.json()[0]["id"] == str(existente.id)
    # Otra clase con el mismo nombre no se confunde.
    otra = "lugar" if clase != "lugar" else "agente"
    assert cliente.post("/api/descripcion/verificar-vocabulario", headers=archivista,
                        json={"fondo_id": str(fondo.id), "tipo": otra, "valor": "policia local de Tunja"}).json() == []


def test_reutilizar_y_crear_nueva_actividad_tipo_y_mandato(cliente, db, fondo, archivista, admin, motor_contexto):
    previos = {c: vocabulario.crear(db, fondo_id=fondo.id, clase=c, nombre=n, subtipo=s, origen="persona", confianza=None,
                                    motor=None, usuario_id=admin.id)
               for c, n, s in (("tipo_actividad", "Policía local", None), ("mandato", "Acuerdo 7 de 1946", "acuerdo"))}
    db.commit()
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    datos = aceptar_todo(espacio)
    # Sin decir si es la misma, el servidor no crea un duplicado: pregunta.
    for e in datos["entidades"]:
        e["crear_nueva"] = False
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json=datos)
    assert r.status_code == 409 and r.json()["coincidencias"]
    for e in datos["entidades"]:
        if e["tipo"] in previos:
            e["reutilizar_id"] = str(previos[e["tipo"]].id)
        else:
            e["crear_nueva"] = True
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=datos).status_code == 201
    for clase in previos:
        assert len(db.scalars(select(EntidadVocabulario).where(EntidadVocabulario.clase == clase)).all()) == 1


# --- Decisiones de IA en auditoría (definición de terminado) -------------------------------------------------


def test_un_evento_de_decision_por_cada_propuesta_aceptada_corregida_rechazada_y_agregada(cliente, db, fondo, archivista,
                                                                                            motor_contexto):
    d = documento(db, fondo, "Oficio_212_1948.pdf", OFICIO)
    espacio = iniciar(cliente, archivista, [d.id])
    datos = aceptar_todo(espacio)
    por_valor = {e["valor"]: e for e in datos["entidades"]}
    por_valor["Gobernador del Departamento"]["valor"] = "Gobernación de Boyacá"  # corregida
    por_valor["Gobernador del Departamento"]["subtipo"] = "entidad_corporativa"
    por_valor["hacia 1948"]["edtf"] = "1948?"  # corregida: misma expresión, otra precisión
    datos["entidades"].remove(por_valor["Oficio"])  # rechazada
    datos["entidades"].remove(por_valor["Tunja"])  # rechazada
    datos["entidades"].append({"tipo": "lugar", "valor": "Mercado de Tunja", "crear_nueva": True})  # agregada
    datos["titulo"] = "Oficio sobre la vigilancia de los mercados de Tunja"  # corregido
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json=datos)
    assert r.status_code == 201, r.text
    eventos = [e.valor_nuevo for e in decisiones(db, r.json()["id"])]
    propuestas = espacio["propuesta"]["entidades"]
    de_entidades = [e for e in eventos if e["tipo"] not in ("titulo", "alcance")]
    # Exactamente un evento por propuesta, más uno por lo agregado a mano.
    assert sorted(e["clave"] for e in de_entidades if e["decision"] != "agregada") == sorted(p["clave"] for p in propuestas)
    decision = {(e["propuesto"] or e["final"])["valor"]: e for e in de_entidades}
    assert decision["Alcaldía Municipal"]["decision"] == "aceptada"
    gob = decision["Gobernador del Departamento"]
    assert gob["decision"] == "corregida" and gob["final"]["valor"] == "Gobernación de Boyacá"
    assert gob["propuesto"]["subtipo"] == "cargo" and gob["final"]["subtipo"] == "entidad_corporativa"
    fecha = decision["hacia 1948"]
    assert fecha["decision"] == "corregida" and (fecha["propuesto"]["edtf"], fecha["final"]["edtf"]) == ("1948~", "1948?")
    assert decision["Oficio"]["decision"] == "rechazada" and decision["Oficio"]["final"] is None
    assert decision["Tunja"]["decision"] == "rechazada"
    assert decision["Mercado de Tunja"]["decision"] == "agregada" and decision["Mercado de Tunja"]["propuesto"] is None
    for e in ("Acuerdo 7 de 1946", "Policía local", "Vigilancia de mercados por la Alcaldía, 1948"):
        assert decision[e]["decision"] == "aceptada", e
    assert all(e["modelo"] == "motor-de-prueba" and e["version_prompt"] == motor.VERSION_PROMPT for e in eventos)
    [titulo] = [e for e in eventos if e["tipo"] == "titulo"]
    assert titulo["decision"] == "corregida" and titulo["final"]["valor"].endswith("de Tunja")
    [alcance] = [e for e in eventos if e["tipo"] == "alcance"]
    assert alcance["decision"] == "aceptada"


def test_sin_motor_no_hay_decisiones_de_ia(cliente, db, fondo, archivista, monkeypatch):
    monkeypatch.setattr(motor, "motor_activo", lambda: None)
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    r = cliente.post("/api/descripcion/publicar", headers=archivista, json={
        "trabajo_id": espacio["trabajo_id"], "titulo": "Oficio", "entidades": [{"tipo": "lugar", "valor": "Tunja",
                                                                                "crear_nueva": True}]})
    assert r.status_code == 201 and decisiones(db, r.json()["id"]) == []


def test_panel_de_decisiones_de_ia_solo_administrador_con_cifras_y_hoja(cliente, db, fondo, archivista, cabeceras_admin,
                                                                        motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    datos = aceptar_todo(espacio)
    datos["entidades"] = [e for e in datos["entidades"] if e["tipo"] != "forma_documental"]  # 1 rechazada
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=datos).status_code == 201
    assert cliente.get("/api/auditoria/decisiones-ia", headers=archivista).status_code == 403
    crear_usuario(db, "julian@correo.com", "revisor")
    assert cliente.get("/api/auditoria/decisiones-ia", headers=ingresar(cliente, "julian@correo.com")).status_code == 403
    r = cliente.get("/api/auditoria/decisiones-ia", headers=cabeceras_admin)
    assert r.status_code == 200, r.text
    total_propuestas = len(espacio["propuesta"]["entidades"]) + 2  # más título y alcance
    resumen = r.json()["resumen"]
    assert resumen["propuestas"] == total_propuestas and resumen["rechazada"] == 1
    assert resumen["aceptada"] == total_propuestas - 1 and r.json()["versiones_prompt"] == [motor.VERSION_PROMPT]
    formas = cliente.get("/api/auditoria/decisiones-ia", headers=cabeceras_admin,
                         params={"tipo": "forma_documental"}).json()
    assert [f["decision"] for f in formas["filas"]] == ["rechazada"] and formas["filas"][0]["propuesto"] == "Oficio"
    hoja = cliente.get("/api/auditoria/decisiones-ia/hoja-de-calculo", headers=cabeceras_admin)
    assert hoja.status_code == 200
    libro = load_workbook(io.BytesIO(hoja.content))
    assert libro.sheetnames == ["Resumen", "Decisiones"] and libro["Decisiones"].max_row == total_propuestas + 1
    assert cliente.get("/api/auditoria/decisiones-ia/hoja-de-calculo", headers=archivista).status_code == 403


# --- Relaciones entre agentes y fusión en los dos sentidos ---------------------------------------------------


def test_relaciones_entre_agentes_con_su_inversa(cliente, db, fondo, admin, cabeceras_admin):
    def agente(nombre):
        return vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre=nombre, subtipo="entidad_corporativa",
                                 origen="persona", confianza=None, motor=None, usuario_id=admin.id)

    concejo, alcaldia, junta, lugar = agente("Concejo Municipal"), agente("Alcaldía Municipal"), agente("Junta de Ornato"), \
        vocabulario.crear(db, fondo_id=fondo.id, clase="lugar", nombre="Tunja", subtipo=None, origen="persona",
                          confianza=None, motor=None, usuario_id=admin.id)
    db.commit()
    url = f"/api/vocabulario/{alcaldia.id}/relaciones-agente"
    assert cliente.post(url, headers=cabeceras_admin, json={"destino_id": str(junta.id), "tipo": "subordinado"}).status_code == 201
    assert cliente.post(f"/api/vocabulario/{concejo.id}/relaciones-agente", headers=cabeceras_admin,
                        json={"destino_id": str(alcaldia.id), "tipo": "asociado"}).status_code == 201
    # Repetida (también al revés si es asociativa), consigo misma o con un lugar: no.
    for destino, tipo, desde in ((junta, "subordinado", alcaldia), (concejo, "asociado", alcaldia),
                                 (alcaldia, "sucesor", alcaldia), (lugar, "sucesor", alcaldia)):
        assert cliente.post(f"/api/vocabulario/{desde.id}/relaciones-agente", headers=cabeceras_admin,
                            json={"destino_id": str(destino.id), "tipo": tipo}).status_code == 422
    desde_junta = cliente.get(f"/api/vocabulario/{junta.id}", headers=cabeceras_admin).json()["relaciones_agente"]
    assert [(r["etiqueta"], r["uri_rico"], r["con"]["nombre"]) for r in desde_junta] == [
        ("está o estuvo subordinado a", "rico:isOrWasSubordinateTo", "Alcaldía Municipal")]
    desde_alcaldia = cliente.get(f"/api/vocabulario/{alcaldia.id}", headers=cabeceras_admin).json()["relaciones_agente"]
    assert {r["uri_rico"] for r in desde_alcaldia} == {"rico:hasOrHadSubordinate", "rico:isAgentAssociatedWithAgent"}
    # Una sola fila por relación: la inversa se lee, no se duplica.
    assert len(db.scalars(select(Relacion).where(Relacion.codigo_ric.in_(
        ("has_or_had_subordinate", "is_agent_associated_with_agent")))).all()) == 2


def test_fusionar_mueve_tambien_las_relaciones_donde_la_entidad_es_origen(cliente, db, fondo, archivista, admin,
                                                                          motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=aceptar_todo(espacio)).status_code == 201
    alcaldia = db.scalar(select(EntidadVocabulario).where(EntidadVocabulario.nombre == "Alcaldía Municipal"))
    definitiva = vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre="Alcaldía Municipal de Tunja",
                                   subtipo="entidad_corporativa", origen="persona", confianza=None, motor=None,
                                   usuario_id=admin.id)
    db.flush()
    vocabulario.fusionar(db, definitiva=definitiva, absorbida=alcaldia, usuario_id=admin.id)
    db.commit()
    ejerce = db.scalar(select(Relacion).where(Relacion.codigo_ric == "performs_or_performed"))
    autoriza = db.scalar(select(Relacion).where(Relacion.codigo_ric == "authorizes"))
    assert ejerce.origen_id == definitiva.id and ejerce.origen_original_id == alcaldia.id
    assert autoriza.destino_id == definitiva.id and autoriza.destino_original_id == alcaldia.id
