"""
Auditoría versión 7: hallazgos de conformidad, columna «propiedad RiC-O»,
identificador determinístico de la instrucción y su etiqueta legible.
"""

import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.auditoria import RegistroAuditoria
from app.models.hallazgo import HallazgoConformidad
from app.servicios import motor
from tests.conftest import crear_usuario, ingresar
from tests.test_descripcion_contexto import OFICIO, motor_contexto  # noqa: F401
from tests.test_descripcion import aceptar_todo, archivista, documento, fondo, iniciar  # noqa: F401


def eventos(db, accion, entidad_id=None):
    q = select(RegistroAuditoria).where(RegistroAuditoria.accion == accion)
    if entidad_id is not None:
        q = q.where(RegistroAuditoria.entidad_id == str(entidad_id))
    return db.scalars(q.order_by(RegistroAuditoria.id)).all()


# --- Hallazgos --------------------------------------------------------------------------------------


def test_el_panel_nace_con_los_hallazgos_de_la_revision_y_su_estado_real(cliente, db, cabeceras_admin):
    datos = cliente.get("/api/auditoria/hallazgos", headers=cabeceras_admin).json()
    hallazgos = {h["numero"]: h for h in datos["hallazgos"]}
    assert list(hallazgos) == list(range(1, 15))  # los diez de la revisión y cuatro de esta depuración
    # El punto de acceso sin sesión no se da por cerrado: en el servidor sigue apagado.
    assert hallazgos[1]["estado"] == "en_correccion" and "HTTPS" in hallazgos[1]["accion"]
    assert hallazgos[4]["estado"] == "en_correccion" and "25 de los 85" in hallazgos[4]["accion"]
    assert hallazgos[8]["estado"] == "en_correccion"  # falta el fundamento normativo que confirma la autora
    assert {n for n, h in hallazgos.items() if h["estado"] == "cerrado"} == {2, 3, 5, 6, 7, 9, 10, 11, 12, 13, 14}
    assert all((h["estado"] == "cerrado") == (h["cerrado_en"] is not None) for h in hallazgos.values())
    assert datos["conteo"] == {"abierto": 0, "en_correccion": 3, "cerrado": 11}
    # Cada uno con su evento de creación en la auditoría.
    creados = {e.entidad_id for e in eventos(db, "hallazgo_creado")}
    assert {h["id"] for h in hallazgos.values()} <= creados


def test_crear_y_cerrar_un_hallazgo_deja_cada_cambio_en_auditoria(cliente, db, cabeceras_admin):
    r = cliente.post("/api/auditoria/hallazgos", headers=cabeceras_admin, json={
        "titulo": "Sin exportación EAD", "descripcion": "El fondo no se puede entregar en EAD 3.",
        "componentes": ["instrumentos"]})
    assert r.status_code == 201, r.text
    h = r.json()
    assert h["numero"] == 15 and h["estado"] == "abierto" and h["cerrado_en"] is None
    [creado] = eventos(db, "hallazgo_creado", h["id"])
    assert creado.valor_nuevo["estado"] == "abierto"
    ruta = f"/api/auditoria/hallazgos/{h['id']}"
    assert cliente.patch(ruta, headers=cabeceras_admin, json={"estado": "en_correccion",
                                                              "accion": "Se evalúa la librería."}).status_code == 200
    cerrado = cliente.patch(ruta, headers=cabeceras_admin, json={"estado": "cerrado"}).json()
    assert cerrado["cerrado_en"] is not None  # la fecha de cierre se pone sola
    cambios = eventos(db, "hallazgo_actualizado", h["id"])
    assert [(e.valor_anterior["estado"], e.valor_nuevo["estado"]) for e in cambios] == [
        ("abierto", "en_correccion"), ("en_correccion", "cerrado")]
    assert cambios[0].valor_nuevo["accion"] == "Se evalúa la librería."
    # Reabrir quita la fecha de cierre; la historia queda.
    reabierto = cliente.patch(ruta, headers=cabeceras_admin, json={"estado": "abierto"}).json()
    assert reabierto["cerrado_en"] is None and len(eventos(db, "hallazgo_actualizado", h["id"])) == 3
    # El título y la descripción originales no se editan.
    assert cliente.patch(ruta, headers=cabeceras_admin, json={"titulo": "Otro"}).status_code == 422
    assert cliente.patch(ruta, headers=cabeceras_admin,
                         json={"estado": "abierto", "cerrado_en": "2026-10-03"}).status_code == 422
    assert db.get(HallazgoConformidad, uuid.UUID(h["id"])).titulo == "Sin exportación EAD"
    # Sin cambios, sin evento.
    cliente.patch(ruta, headers=cabeceras_admin, json={"estado": "abierto"})
    assert len(eventos(db, "hallazgo_actualizado", h["id"])) == 3


def test_filtros_y_hoja_de_calculo_de_hallazgos(cliente, db, cabeceras_admin):
    abiertos = cliente.get("/api/auditoria/hallazgos", headers=cabeceras_admin, params={"estado": "en_correccion"}).json()
    assert [h["numero"] for h in abiertos["hallazgos"]] == [1, 4, 8]
    pres = cliente.get("/api/auditoria/hallazgos", headers=cabeceras_admin, params={"componente": "preservacion"}).json()
    assert all("preservacion" in h["componentes"] for h in pres["hallazgos"]) and 8 in [h["numero"] for h in pres["hallazgos"]]
    hoja = cliente.get("/api/auditoria/hallazgos/hoja-de-calculo", headers=cabeceras_admin,
                       params={"estado": "en_correccion"})
    filas = list(load_workbook(io.BytesIO(hoja.content)).active.iter_rows(values_only=True))
    assert filas[0][0] == "N.º" and [f[0] for f in filas[1:]] == [1, 4, 8]
    assert cliente.get("/api/auditoria/hallazgos", headers=cabeceras_admin,
                       params={"componente": "inventado"}).status_code == 422


def test_el_archivista_no_ve_ni_toca_los_hallazgos(cliente, db, archivista, cabeceras_admin):
    h = cliente.get("/api/auditoria/hallazgos", headers=cabeceras_admin).json()["hallazgos"][0]
    assert cliente.get("/api/auditoria/hallazgos", headers=archivista).status_code == 403
    assert cliente.post("/api/auditoria/hallazgos", headers=archivista, json={
        "titulo": "x" * 5, "descripcion": "y" * 5, "componentes": ["ingesta"]}).status_code == 403
    assert cliente.patch(f"/api/auditoria/hallazgos/{h['id']}", headers=archivista,
                         json={"estado": "cerrado"}).status_code == 403
    assert cliente.get("/api/auditoria/hallazgos/hoja-de-calculo", headers=archivista).status_code == 403
    assert cliente.get("/api/auditoria/versiones-prompt", headers=archivista).status_code == 403
    assert db.scalar(select(func.count()).select_from(HallazgoConformidad)) == 14


# --- Versión de la instrucción -------------------------------------------------------------------------


def test_el_identificador_de_la_instruccion_es_deterministico():
    esquema = {"type": "object"}
    assert motor.version_de("Describe el documento.", esquema) == motor.version_de("Describe el documento.", esquema)
    assert motor.version_de("Describe el documento.", esquema) != motor.version_de("Describe el documento!", esquema)
    assert motor.version_de("Describe el documento.", esquema) != motor.version_de("Describe el documento.", {"type": "array"})
    assert len(motor.VERSION_PROMPT) == 8 and motor.VERSION_PROMPT == motor.version_de(motor.INSTRUCCION, motor.ESQUEMA)


def test_la_etiqueta_legible_acompana_al_identificador_sin_reemplazarlo(cliente, db, fondo, archivista, cabeceras_admin,
                                                                        motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=aceptar_todo(espacio)).status_code == 201
    [vigente] = [v for v in cliente.get("/api/auditoria/versiones-prompt", headers=cabeceras_admin).json() if v["vigente"]]
    assert vigente["version"] == motor.VERSION_PROMPT and vigente["decisiones"] > 0 and vigente["etiqueta"] is None
    r = cliente.put(f"/api/auditoria/versiones-prompt/{motor.VERSION_PROMPT}", headers=cabeceras_admin,
                    json={"etiqueta": "v3", "nota": "Con el contexto de vocabulario."})
    assert r.status_code == 200
    filas = cliente.get("/api/auditoria/decisiones-ia", headers=cabeceras_admin).json()["filas"]
    assert all(f["version_prompt"] == motor.VERSION_PROMPT and f["version_etiqueta"] == "v3" for f in filas)
    [evento] = eventos(db, "version_prompt_etiquetada")
    assert evento.valor_anterior == {"etiqueta": None} and evento.valor_nuevo == {"etiqueta": "v3"}
    # Solo versiones que existen, y una etiqueta no nombra dos versiones.
    assert cliente.put("/api/auditoria/versiones-prompt/ffffffff", headers=cabeceras_admin,
                       json={"etiqueta": "v9"}).status_code == 422


# --- Propiedad RiC-O ---------------------------------------------------------------------------------


def test_cada_decision_muestra_su_clase_y_su_propiedad_de_ric_o(cliente, db, fondo, archivista, cabeceras_admin,
                                                               motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    datos = aceptar_todo(espacio)
    por_valor = {e["valor"]: e for e in datos["entidades"]}
    datos["entidades"].remove(por_valor["Tunja"])  # rechazada: muestra igual la propiedad que habría tenido
    assert cliente.post("/api/descripcion/publicar", headers=archivista, json=datos).status_code == 201
    filas = cliente.get("/api/auditoria/decisiones-ia", headers=cabeceras_admin).json()["filas"]
    por = {(f["tipo"], (f["propuesto"] or f["final"] or "").split(" · ")[0]): f for f in filas}
    productor = por[("agente", "Alcaldía Municipal")]
    assert productor["clase_rico"] == "rico:CorporateBody"
    assert productor["propiedad_rico"]["nombre"] == "rico:hasCreator" and productor["propiedad_rico"]["codigo_cm"] == "RiC-R027"
    assert por[("agente", "Gobernador del Departamento")]["propiedad_rico"]["nombre"] == "rico:hasAddressee"
    assert por[("lugar", "Tunja")]["decision"] == "rechazada"
    assert por[("lugar", "Tunja")]["propiedad_rico"]["nombre"] == "rico:hasOrHadSubject"
    assert por[("actividad", "Vigilancia de mercados por la Alcaldía, 1948")]["propiedad_rico"]["nombre"] == "rico:documents"
    assert por[("mandato", "Acuerdo 7 de 1946")]["propiedad_rico"]["nombre"] == "rico:regulatesOrRegulated"
    assert por[("tipo_actividad", "Policía local")]["propiedad_rico"]["nombre"] == "rico:hasActivityType"
    assert por[("forma_documental", "Oficio")]["propiedad_rico"]["nombre"] == "rico:hasDocumentaryFormType"
    [titulo] = [f for f in filas if f["tipo"] == "titulo"]
    assert titulo["propiedad_rico"] == {"nombre": "rico:title", "estado": "verificada"}
    fecha = next(f for f in filas if f["tipo"] == "fecha")
    assert fecha["clase_rico"] == "rico:Date" and fecha["propiedad_rico"]["nombre"] == "rico:isCreationDateOf"


def test_la_trazabilidad_marca_la_propiedad_o_el_literal_pendiente(cliente, db, fondo, archivista, motor_contexto):
    espacio = iniciar(cliente, archivista, [documento(db, fondo, "a.pdf", OFICIO).id])
    recurso_id = cliente.post("/api/descripcion/publicar", headers=archivista, json=aceptar_todo(espacio)).json()["id"]
    trabajo = cliente.post(f"/api/descripcion/registros/{recurso_id}/reabrir", headers=archivista).json()["trabajo_id"]
    r = cliente.patch(f"/api/descripcion/{recurso_id}", headers=archivista, json={
        "trabajo_id": trabajo, "titulo": "Oficio sobre los mercados", "control": {"caja": "4"}})
    assert r.status_code == 200, r.text
    historia = cliente.get(f"/api/auditoria/entidad/{recurso_id}", headers=archivista,
                           params={"tipo": "recurso_documental"}).json()["eventos"]
    cambios = {c["campo"]: c for e in historia if e["accion"] == "descripcion_editada" for c in e["cambios"]}
    assert cambios["titulo"]["propiedad_rico"] == {"nombre": "rico:title", "estado": "verificada"}
    # La caja (dato de control del inventario) no tiene propiedad confirmada: se dice, no se inventa.
    assert cambios["caja"]["propiedad_rico"] == {"nombre": None, "estado": "literal_pendiente"}


def test_un_vinculo_entre_agentes_muestra_su_propiedad(cliente, db, fondo, archivista):
    from app.servicios import vocabulario

    a = vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre="Alcaldía", subtipo="entidad_corporativa",
                          origen="persona", confianza=None, motor=None, usuario_id=None)
    b = vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre="Secretaría de Gobierno",
                          subtipo="entidad_corporativa", origen="persona", confianza=None, motor=None, usuario_id=None)
    db.commit()
    r = cliente.post(f"/api/vocabulario/{a.id}/vinculos", headers=archivista,
                     json={"tipo": "subordinado", "con_tipo": "entidad_vocabulario", "con_id": str(b.id)})
    assert r.status_code in (200, 201), r.text
    historia = cliente.get(f"/api/auditoria/entidad/{a.id}", headers=archivista,
                           params={"tipo": "entidad_vocabulario"}).json()["eventos"]
    [vinculo] = [e for e in historia if e["accion"] == "vinculo_declarado"]
    assert vinculo["propiedad_rico"]["nombre"] == "rico:hasOrHadSubordinate"
    assert vinculo["propiedad_rico"]["codigo_cm"] == "RiC-R045"


def test_todo_campo_con_propiedad_existe_en_el_mapeo_verificado():
    from app.servicios import ric_o

    nombres = {n for n, _ in ric_o.ATRIBUTOS.values()} | {n for n, _, _ in ric_o.APOYO.values()}
    assert all(v in nombres or v.startswith("skos:") for v in ric_o.CAMPO_RICO.values())
