"""
Cierre de los dos riesgos abiertos antes del piloto:

1. La clave del segundo factor se guarda cifrada (AES-256-GCM), atada a su
   cuenta y con la clave fuera de la base: un volcado o un paquete de
   recuperación no basta para generar códigos. Sin la clave correcta, ningún
   código de la aplicación sirve (falla cerrado) y los de respaldo sí.
2. El FUID, la guía y el índice se aprueban como documentos fijos con huella,
   acta y responsable; se reemplazan por versiones y se anulan sin borrar; en
   RiC-O son documentos que describen su nivel (RiC-R021), conformes.
"""

import hashlib
import io

import pytest
from openpyxl import load_workbook
from rdflib import Namespace, URIRef
from sqlalchemy import select, text

from app.core import cifrado, doble_factor
from app.core.config import settings
from app.models.auditoria import RegistroAuditoria
from app.models.instrumento_aprobado import InstrumentoAprobado
from app.models.usuario import Usuario
from app.servicios import almacen, exportacion_rico
from tests.conftest import CONTRASENA, crear_usuario, ingresar
from tests.test_brechas_lote1 import _activar
from tests.test_especializacion_ric import exportar_conforme
from tests.test_instrumentos import archivista, fondo_descrito  # noqa: F401  (fixtures)

RICO = Namespace("https://www.ica.org/standards/RiC/ontology#")


# --- 1 · Segundo factor cifrado ---------------------------------------------------------------------


def _usuario(db, correo):
    db.expire_all()
    return db.scalar(select(Usuario).where(Usuario.correo == correo))


def test_la_clave_del_segundo_factor_queda_cifrada_y_atada_a_la_cuenta(cliente, db):
    crear_usuario(db, "cifra@ricora.prueba", "archivista")
    crear_usuario(db, "otra@ricora.prueba", "archivista")
    clave, _ = _activar(cliente, ingresar(cliente, "cifra@ricora.prueba"))
    u = _usuario(db, "cifra@ricora.prueba")
    assert u.mfa_secreto.startswith("c1:") and clave not in u.mfa_secreto
    assert doble_factor.leer(u.mfa_secreto, u.id) == clave
    # Copiar el valor cifrado a otra cuenta no sirve: el dato asociado es la cuenta.
    otra = _usuario(db, "otra@ricora.prueba")
    assert doble_factor.leer(u.mfa_secreto, otra.id) == ""
    # Alterado, tampoco.
    alterado = u.mfa_secreto[:-4] + ("AAAA" if not u.mfa_secreto.endswith("AAAA") else "BBBB")
    assert doble_factor.leer(alterado, u.id) == ""


def test_con_otra_clave_falla_cerrado_y_los_codigos_de_respaldo_siguen_sirviendo(cliente, db, monkeypatch):
    crear_usuario(db, "llave@ricora.prueba", "archivista")
    clave, respaldo = _activar(cliente, ingresar(cliente, "llave@ricora.prueba"))
    monkeypatch.setattr(settings, "clave_cifrado", "otra-clave-de-cifrado-distinta-con-mas-de-32-caracteres")
    with pytest.raises(cifrado.ErrorCifrado, match="otra clave"):
        cifrado.descifrar(_usuario(db, "llave@ricora.prueba").mfa_secreto, "x")
    desafio = cliente.post("/api/auth/login", json={"correo": "llave@ricora.prueba",
                                                    "contrasena": CONTRASENA}).json()["desafio"]
    codigo = doble_factor.codigo(clave, doble_factor.paso_actual() + 1)
    assert cliente.post("/api/auth/login/segundo-factor", json={"desafio": desafio, "codigo": codigo}).status_code == 401
    r = cliente.post("/api/auth/login/segundo-factor", json={"desafio": desafio, "codigo": respaldo[0]})
    assert r.status_code == 200, r.text


def test_la_migracion_cifra_las_claves_que_estaban_en_claro(db):
    crear_usuario(db, "vieja@ricora.prueba", "archivista")
    u = _usuario(db, "vieja@ricora.prueba")
    en_claro = doble_factor.generar_secreto()
    db.execute(text("UPDATE usuarios SET mfa_secreto = :v WHERE id = :i"), {"v": en_claro, "i": u.id})
    assert doble_factor.leer(en_claro, u.id) == ""  # sin migrar, no se acepta en claro
    assert doble_factor.cifrar_en_claro(db.connection()) == 1
    assert doble_factor.cifrar_en_claro(db.connection()) == 0  # repetirla no cambia nada
    u = _usuario(db, "vieja@ricora.prueba")
    assert u.mfa_secreto.startswith("c1:") and doble_factor.leer(u.mfa_secreto, u.id) == en_claro


# --- 2 · Instrumentos aprobados ---------------------------------------------------------------------


def _aprobar(cliente, cab, **datos):
    return cliente.post("/api/instrumentos/aprobados", headers=cab,
                        json={k: str(v) if hasattr(v, "hex") else v for k, v in datos.items()})


def test_un_fuid_con_campos_obligatorios_vacios_no_se_aprueba(cliente, fondo_descrito, archivista):  # noqa: F811
    r = _aprobar(cliente, archivista, tipo="inventario", recurso_id=fondo_descrito["exp48"].id)
    assert r.status_code == 422 and "campos obligatorios del FUID sin dato" in r.json()["detail"]


def test_aprobar_guarda_el_archivo_con_su_huella_y_su_acta(cliente, db, fondo_descrito, archivista):  # noqa: F811
    exp49 = fondo_descrito["exp49"]
    r = _aprobar(cliente, archivista, tipo="inventario", recurso_id=exp49.id, acta="Acta 07 del Comité",
                 fecha_acta="2026-10-20")
    assert r.status_code == 201, r.text
    a = r.json()
    assert (a["version"], a["estado"], a["acta"], a["fecha_acta"]) == (1, "vigente", "Acta 07 del Comité", "2026-10-20")
    assert a["aprobado_por"] == "Catalina Torres" and a["tipo_nombre"] == "Inventario documental (FUID)"
    d = cliente.get(f"/api/instrumentos/aprobados/{a['id']}/archivo", headers=archivista)
    assert d.status_code == 200 and hashlib.sha256(d.content).hexdigest() == a["huella"] == d.headers["X-Huella-SHA256"]
    filas = [[c.value for c in f] for f in load_workbook(io.BytesIO(d.content)).active.iter_rows()]
    assert any(f[1] == "CO-49" for f in filas)
    assert db.scalar(select(RegistroAuditoria).where(RegistroAuditoria.accion == "instrumento_aprobado"))
    # Si el archivo cambia en el disco, no se entrega como si fuera el aprobado.
    i = db.get(InstrumentoAprobado, a["id"])
    almacen.ruta_absoluta(i.ruta).write_bytes(d.content + b"x")
    r = cliente.get(f"/api/instrumentos/aprobados/{a['id']}/archivo", headers=archivista)
    assert r.status_code == 409 and "alterado" in r.json()["detail"]


def test_version_nueva_reemplaza_y_anular_devuelve_la_anterior_sin_borrar(cliente, db, fondo_descrito,  # noqa: F811
                                                                          archivista):
    exp49 = fondo_descrito["exp49"]
    v1 = _aprobar(cliente, archivista, tipo="inventario", recurso_id=exp49.id).json()
    v2 = _aprobar(cliente, archivista, tipo="inventario", recurso_id=exp49.id).json()
    assert v2["version"] == 2
    estados = {x["id"]: x["estado"] for x in cliente.get(
        "/api/instrumentos/aprobados", headers=archivista, params={"fondo_id": str(fondo_descrito["fondo"].id)}).json()}
    assert estados == {v1["id"]: "reemplazado", v2["id"]: "vigente"}
    corto = cliente.post(f"/api/instrumentos/aprobados/{v2['id']}/anular", headers=archivista, json={"motivo": "error"})
    assert corto.status_code == 422
    r = cliente.post(f"/api/instrumentos/aprobados/{v2['id']}/anular", headers=archivista,
                     json={"motivo": "Se aprobó con el expediente equivocado."})
    assert r.status_code == 200 and r.json()["estado"] == "anulado"
    db.expire_all()
    assert db.get(InstrumentoAprobado, v1["id"]).estado == "vigente"
    anulado = db.get(InstrumentoAprobado, v2["id"])
    assert almacen.ruta_absoluta(anulado.ruta).exists()  # nada se borra
    assert cliente.post(f"/api/instrumentos/aprobados/{v2['id']}/anular", headers=archivista,
                        json={"motivo": "Otra vez, por si acaso."}).status_code == 409


def test_guia_e_indice_del_fondo_y_quien_puede_aprobar(cliente, db, fondo_descrito, archivista):  # noqa: F811
    f = fondo_descrito
    assert _aprobar(cliente, archivista, tipo="guia", recurso_id=f["fondo"].id, texto="  ").status_code == 422
    assert _aprobar(cliente, archivista, tipo="guia", recurso_id=f["serie"].id, texto="Nota").status_code == 422
    assert _aprobar(cliente, archivista, tipo="guia", recurso_id=f["fondo"].id,
                    texto="El fondo reúne la correspondencia de la Alcaldía.").status_code == 201
    indice = _aprobar(cliente, archivista, tipo="indice", recurso_id=f["fondo"].id)
    assert indice.status_code == 201, indice.text
    contenido = cliente.get(f"/api/instrumentos/aprobados/{indice.json()['id']}/archivo", headers=archivista).content
    filas = [[c.value for c in x] for x in load_workbook(io.BytesIO(contenido)).active.iter_rows()]
    assert ["Agentes", "Alcaldía Municipal", "Oficio N.º 114", "Unidad documental"] in filas
    crear_usuario(db, "lector@ricora.prueba", "consulta")
    lector = ingresar(cliente, "lector@ricora.prueba")
    assert _aprobar(cliente, lector, tipo="indice", recurso_id=f["fondo"].id).status_code == 403


def test_en_ric_o_cada_instrumento_aprobado_describe_su_nivel(cliente, db, fondo_descrito, archivista):  # noqa: F811
    f = fondo_descrito
    v1 = _aprobar(cliente, archivista, tipo="inventario", recurso_id=f["exp49"].id, acta="Acta 07").json()
    v2 = _aprobar(cliente, archivista, tipo="inventario", recurso_id=f["exp49"].id).json()
    v3 = _aprobar(cliente, archivista, tipo="inventario", recurso_id=f["exp49"].id).json()
    cliente.post(f"/api/instrumentos/aprobados/{v3['id']}/anular", headers=archivista,
                 json={"motivo": "Duplicado por error al aprobar."})
    g = exportar_conforme(db, f["fondo"]).grafo
    u = exportacion_rico.uri
    doc1, doc2 = u(v1["id"], "documento"), u(v2["id"], "documento")
    for doc in (doc1, doc2):
        assert (doc, RICO.describesOrDescribed, u(f["exp49"].id)) in g
        assert any(g.objects(doc, RICO.hasOrHadInstantiation))
    assert (doc1, RICO.precedesOrPreceded, doc2) in g
    assert not any(g.triples((u(v3["id"], "documento"), None, None)))  # el anulado no sale
    notas = [str(o) for o in g.objects(u(v1["id"]), RICO.authenticityNote)]
    assert notas == [f"Huella SHA-256 del archivo aprobado: {v1['huella']}."]
    assert "Acta 07" in str(next(g.objects(doc1, RICO.generalDescription)))
    assert URIRef(u(v1["id"])) in set(g.subjects(RICO.identifier, None))
