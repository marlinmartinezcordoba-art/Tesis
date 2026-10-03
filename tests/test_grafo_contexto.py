"""Grafo de contexto del fondo: recorrido acotado, filtros, acceso, ficha,
exportación del fragmento visible y prueba de carga."""

import io
import time
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from rdflib import RDF, Graph, URIRef

from app.db.base import ahora
from app.models.descripcion import EntidadVocabulario, Fecha, Relacion, TrabajoDescripcion
from app.models.preservacion import DeclaracionDerechos
from app.models.recurso_documental import RecursoDocumental
from app.servicios import exportacion_rico, grafo
from tests.conftest import crear_usuario, ingresar
from tests.test_instrumentos import _sin_nada_interno, archivista, fondo_descrito  # noqa: F401


@pytest.fixture()
def consulta(cliente, db):
    crear_usuario(db, "lector@correo.com", "consulta")
    return ingresar(cliente, "lector@correo.com")


def pedir(cliente, quien, tipo, ident, **params):
    r = cliente.get(f"/api/grafo/{tipo}/{ident}", headers=quien, params=params)
    assert r.status_code == 200, r.text
    return r.json()


def etiquetas(g):
    return {n["etiqueta"] for n in g["nodos"]}


FONDO, SERIE, EXP48, EXP49 = "Correspondencia municipal", "Correspondencia", "Correspondencia 1948", "Correspondencia 1949"
O114, O115, ALCALDIA = "Oficio N.º 114", "Oficio N.º 115", "Alcaldía Municipal"


# --- Recorrido acotado: exactamente lo alcanzable, ni más ni menos ---------------------------------------


def test_un_dos_y_tres_saltos_devuelven_exactamente_lo_alcanzable(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    uno = pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=1)
    assert etiquetas(uno) == {FONDO, SERIE}
    assert [(a["codigo_ric"], a["dirigida"]) for a in uno["aristas"]] == [("includes_or_included", True)]
    _sin_nada_interno(uno)
    dos = pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=2)
    assert etiquetas(dos) == {FONDO, SERIE, EXP48, EXP49}
    tres = pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=3)
    # Expediente 1949: sus dos fechas y la Alcaldía. Expediente 1948: sus dos oficios (el borrador no).
    assert etiquetas(tres) == {FONDO, SERIE, EXP48, EXP49, O114, O115, ALCALDIA, "1949-01-10", "1949-12-02"}
    assert "Borrador" not in etiquetas(tres)
    assert {n["salto"] for n in tres["nodos"]} == {0, 1, 2, 3}
    claves = {n["clave"] for n in tres["nodos"]}
    assert all(a["desde"] in claves and a["hacia"] in claves for a in tres["aristas"])
    # Nunca más de tres saltos.
    assert cliente.get(f"/api/grafo/recurso_documental/{f['fondo'].id}", headers=archivista,
                       params={"saltos": 4}).status_code == 422


def test_desde_un_documento_sus_relaciones_de_primer_nivel(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    g = pedir(cliente, archivista, "recurso_documental", f["o114"].id)
    assert etiquetas(g) == {O114, EXP48, ALCALDIA, "Gobernador del Departamento", "Boyacá", "Oficio",
                            "Oficio_114_1948.pdf", "1948-03-15"}
    familias = {n["etiqueta"]: n["familia"] for n in g["nodos"]}
    assert familias[O114] == "Record" and familias[EXP48] == "RecordSet" and familias[ALCALDIA] == "Agent"
    assert familias["Boyacá"] == "Place" and familias["Oficio"] == "DocumentaryFormType"
    assert familias["Oficio_114_1948.pdf"] == "Instantiation"
    assert any(n["familia"] == "Date" for n in g["nodos"])
    assert any(a["uri_rico"] == "rico:hasCreator" and a["etiqueta"] == "producido por" for a in g["aristas"])
    # La raíz va marcada y es el punto de partida.
    assert g["centro"] == f"recurso_documental:{f['o114'].id}"


# --- Filtros ----------------------------------------------------------------------------------------


def test_filtro_por_tipo_de_entidad_excluye_nodos_y_sus_relaciones(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    g = pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=3, tipo_entidad=["RecordSet", "Record"])
    assert etiquetas(g) == {FONDO, SERIE, EXP48, EXP49, O114, O115}
    assert {n["familia"] for n in g["nodos"]} == {"RecordSet", "Record"}
    assert {a["codigo_ric"] for a in g["aristas"]} == {"includes_or_included"}
    assert g["filtros_activos"] == 1
    # El recorrido no atraviesa lo filtrado: sin agrupaciones, desde el fondo no se llega a nada.
    solo_agentes = pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=3, tipo_entidad=["Agent"])
    assert etiquetas(solo_agentes) == {FONDO} and solo_agentes["aristas"] == []
    assert cliente.get(f"/api/grafo/recurso_documental/{f['fondo'].id}", headers=archivista,
                       params={"tipo_entidad": "Persona"}).status_code == 422


def test_filtro_por_tipo_de_relacion(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    g = pedir(cliente, archivista, "recurso_documental", f["o114"].id, tipo_relacion=["has_creator", "has_addressee"])
    assert etiquetas(g) == {O114, ALCALDIA, "Gobernador del Departamento"}


def test_filtro_por_rango_de_fechas_excluye_lo_que_cae_fuera(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    # El oficio 114 es del 15 de marzo de 1948; el 115, del 18.
    g = pedir(cliente, archivista, "recurso_documental", f["exp48"].id, saltos=2, fecha_desde="1948-03-17")
    assert O115 in etiquetas(g) and O114 not in etiquetas(g)
    assert "Oficio_114_1948.pdf" not in etiquetas(g)  # solo se llegaba por el 114
    # Sin fecha no se excluye (no hay evidencia de que caiga fuera); la raíz siempre se muestra.
    assert {SERIE, EXP48} <= etiquetas(g)
    g = pedir(cliente, archivista, "recurso_documental", f["exp49"].id, fecha_hasta="1949-06-30")
    assert "1949-01-10" in etiquetas(g) and "1949-12-02" not in etiquetas(g)
    assert cliente.get(f"/api/grafo/recurso_documental/{f['exp49'].id}", headers=archivista,
                       params={"fecha_desde": "1950-01-01", "fecha_hasta": "1949-01-01"}).status_code == 422


def test_filtro_por_estado_de_la_descripcion(cliente, db, fondo_descrito, archivista, admin):
    f = fondo_descrito
    db.add(TrabajoDescripcion(usuario_id=admin.id, fondo_id=f["fondo"].id, nivel="unidad_documental",
                              recurso_id=f["o115"].id, estado="abierto"))
    db.commit()
    g = pedir(cliente, archivista, "recurso_documental", f["exp48"].id, estado=["en_edicion"])
    assert etiquetas(g) == {EXP48, O115}
    assert next(n for n in g["nodos"] if n["etiqueta"] == O115)["estado"] == "en_edicion"
    publicadas = pedir(cliente, archivista, "recurso_documental", f["exp48"].id, estado=["publicada"])
    assert O114 in etiquetas(publicadas) and O115 not in etiquetas(publicadas)


# --- Acceso -----------------------------------------------------------------------------------------


def test_quien_no_es_archivista_nunca_recibe_lo_reservado(cliente, db, fondo_descrito, archivista, consulta, admin):
    f = fondo_descrito
    db.add(DeclaracionDerechos(fondo_id=f["fondo"].id, entidad_tipo="recurso_documental", entidad_id=f["exp48"].id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()
    lector = pedir(cliente, consulta, "recurso_documental", f["fondo"].id, saltos=3)
    assert not ({EXP48, O114, O115} & etiquetas(lector))  # el expediente y lo que hereda su reserva
    assert EXP49 in etiquetas(lector)
    for ruta in (f"/api/grafo/recurso_documental/{f['o114'].id}", f"/api/grafo/recurso_documental/{f['o114'].id}/ficha"):
        assert cliente.get(ruta, headers=consulta).status_code == 404
    assert cliente.get(f"/api/grafo/recurso_documental/{f['o114'].id}/exportar", headers=consulta).status_code == 404
    # La vista antigua aplica la misma regla.
    viejo = cliente.get("/api/instrumentos/grafo", headers=consulta,
                        params={"fondo_id": str(f["fondo"].id), "profundidad": 3}).json()
    assert O114 not in etiquetas(viejo)
    # Con escritura en el catálogo sí se ve, como en la exportación RiC-O.
    assert O114 in etiquetas(pedir(cliente, archivista, "recurso_documental", f["fondo"].id, saltos=3))
    assert cliente.get(f"/api/grafo/recurso_documental/{f['fondo'].id}").status_code == 401


# --- Ficha y relaciones -------------------------------------------------------------------------------


def test_ficha_con_resumen_atributos_y_relaciones(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    r = cliente.get(f"/api/grafo/recurso_documental/{f['o114'].id}/ficha", headers=archivista)
    assert r.status_code == 200
    ficha = r.json()
    _sin_nada_interno(ficha)
    resumen = {c["campo"]: c["valor"] for c in ficha["resumen"]}
    assert resumen["Identificador"] == "CO-AM-114" and resumen["Título"] == O114
    assert resumen["Alcance y contenido"].startswith("Informa")
    atributos = {c["campo"]: c["valor"] for c in ficha["atributos"]}
    assert atributos["Forma documental"] == "Oficio" and atributos["Folios"] == "2"
    assert ficha["familia"] == "Record" and ficha["relaciones_total"] == 7
    assert {x["otro"]["etiqueta"] for x in ficha["relaciones"]} >= {ALCALDIA, EXP48}
    hacia_exp = next(x for x in ficha["relaciones"] if x["otro"]["etiqueta"] == EXP48)
    assert hacia_exp["sentido"] == "entra" and hacia_exp["uri_rico"] == "rico:includesOrIncluded"
    agente = cliente.get(f"/api/grafo/entidad_vocabulario/{f['alcaldia'].id}/ficha", headers=archivista).json()
    assert {c["campo"]: c["valor"] for c in agente["resumen"]}["Tipo"] == "entidad corporativa"
    assert agente["relaciones_total"] == 3  # los tres documentos publicados que produjo


def test_muchas_relaciones_muestran_las_recientes_y_exportan_todas(cliente, db, fondo_descrito, archivista, admin):
    f = fondo_descrito
    for n in range(20):
        lugar = EntidadVocabulario(fondo_id=f["fondo"].id, clase="lugar", nombre=f"Vereda {n:02d}",
                                   nombre_normalizado=f"vereda {n:02d}", origen="persona")
        db.add(lugar)
        db.flush()
        db.add(Relacion(origen_tipo="recurso_documental", origen_id=f["o115"].id, destino_tipo="entidad_vocabulario",
                        destino_id=lugar.id, tipo_relacion="asociacion", codigo_ric="has_or_had_subject", origen="persona"))
    db.commit()
    ficha = cliente.get(f"/api/grafo/recurso_documental/{f['o115'].id}/ficha", headers=archivista).json()
    assert ficha["relaciones_total"] == 23 and len(ficha["relaciones"]) == grafo.RELACIONES_VISIBLES
    r = cliente.get(f"/api/grafo/recurso_documental/{f['o115'].id}/relaciones/exportar", headers=archivista)
    assert r.status_code == 200
    hoja = load_workbook(io.BytesIO(r.content)).active
    filas = [x for x in hoja.iter_rows(min_row=4, values_only=True) if x[0]]
    assert len(filas) == 23 and hoja.cell(3, 2).value == "Propiedad RiC-O"


# --- Exportación del fragmento visible ----------------------------------------------------------------


def _entidades_en(g: Graph, todas: set[URIRef]) -> set[URIRef]:
    return {s for s in g.subjects(RDF.type, None) if s in todas}


def test_exportar_produce_exactamente_lo_visible(cliente, db, fondo_descrito, archivista):
    f = fondo_descrito
    params = {"saltos": 1, "tipo_entidad": ["Record", "RecordSet", "Agent"]}
    visible = pedir(cliente, archivista, "recurso_documental", f["o114"].id, **params)
    r = cliente.get(f"/api/grafo/recurso_documental/{f['o114'].id}/exportar", headers=archivista,
                     params={**params, "formato": "turtle"})
    assert r.status_code == 200
    assert "fragmento" in r.headers["content-disposition"] and "1-saltos" in r.headers["content-disposition"]
    g = Graph().parse(data=r.content, format="turtle")
    ex = exportacion_rico.exportar(db, db.get(RecursoDocumental, f["fondo"].id), True)
    todas = {exportacion_rico.uri(i) for i in ex.nodos}
    esperadas = {exportacion_rico.uri(n["id"]) for n in visible["nodos"]}
    assert _entidades_en(g, todas) == esperadas  # ni el fondo entero ni otro subconjunto
    # Toda tripleta entre dos entidades corresponde a una relación visible.
    pares = {frozenset((exportacion_rico.uri(a["desde"].split(":")[1]), exportacion_rico.uri(a["hacia"].split(":")[1])))
             for a in visible["aristas"]}
    entre = {frozenset((s, o)) for s, _, o in g if s in todas and o in todas}
    assert entre and entre <= pares
    assert exportacion_rico.uri(f["exp49"].id) not in set(g.all_nodes())
    # JSON-LD dice lo mismo.
    j = cliente.get(f"/api/grafo/recurso_documental/{f['o114'].id}/exportar", headers=archivista,
                     params={**params, "formato": "jsonld"})
    assert len(Graph().parse(data=j.content, format="json-ld")) == len(g)


# --- Prueba de carga ----------------------------------------------------------------------------------


def test_carga_con_un_fondo_de_mas_de_quinientas_entidades(cliente, db, archivista, admin):
    """Fondo › 5 series › 50 expedientes › 250 documentos, 150 agentes,
    50 lugares y 250 fechas: 756 entidades y más de 1.000 relaciones."""
    fondo = RecursoDocumental(id=uuid.uuid4(), nivel="fondo", titulo="Fondo de carga")
    fondo.fondo_id = fondo.id
    db.add(fondo)
    db.flush()
    agentes = [EntidadVocabulario(id=uuid.uuid4(), fondo_id=fondo.id, clase="agente", subtipo="persona",
                                  nombre=f"Persona {i}", nombre_normalizado=f"persona {i}", origen="persona")
               for i in range(150)]
    lugares = [EntidadVocabulario(id=uuid.uuid4(), fondo_id=fondo.id, clase="lugar", nombre=f"Lugar {i}",
                                  nombre_normalizado=f"lugar {i}", origen="persona") for i in range(50)]
    db.add_all(agentes + lugares)
    filas, relaciones, documentos = [], [], 0
    for s in range(5):
        serie = RecursoDocumental(id=uuid.uuid4(), nivel="serie", titulo=f"Serie {s}", fondo_id=fondo.id,
                                  incluido_en_id=fondo.id, publicado_en=ahora())
        filas.append(serie)
        for e in range(10):
            exp = RecursoDocumental(id=uuid.uuid4(), nivel="expediente", titulo=f"Expediente {s}.{e}", fondo_id=fondo.id,
                                    incluido_en_id=serie.id, publicado_en=ahora())
            filas.append(exp)
            for d in range(5):
                doc = RecursoDocumental(id=uuid.uuid4(), nivel="unidad_documental", titulo=f"Documento {s}.{e}.{d}",
                                        fondo_id=fondo.id, incluido_en_id=exp.id, publicado_en=ahora())
                fecha = Fecha(id=uuid.uuid4(), expresion="1950", edtf="1950", inicio=date(1950, 1, 1),
                              fin=date(1950, 12, 31), origen="persona")
                filas += [doc, fecha]
                for destino, codigo in ((agentes[documentos % 150], "has_creator"),
                                        (agentes[(documentos * 7) % 150], "has_addressee"),
                                        (lugares[documentos % 50], "has_or_had_subject")):
                    relaciones.append(Relacion(origen_tipo="recurso_documental", origen_id=doc.id,
                                               destino_tipo="entidad_vocabulario", destino_id=destino.id,
                                               tipo_relacion="asociacion", codigo_ric=codigo, origen="persona"))
                relaciones.append(Relacion(origen_tipo="fecha", origen_id=fecha.id, destino_tipo="recurso_documental",
                                           destino_id=doc.id, tipo_relacion="temporal",
                                           codigo_ric="is_creation_date_of", origen="persona"))
                documentos += 1
    db.add_all(filas)
    db.flush()
    db.add_all(relaciones)
    db.commit()
    tiempos = {}
    for nombre, tipo, ident in (("fondo", "recurso_documental", fondo.id),
                                ("agente", "entidad_vocabulario", agentes[0].id)):
        inicio = time.perf_counter()
        g = pedir(cliente, archivista, tipo, ident, saltos=3)
        tiempos[nombre] = (time.perf_counter() - inicio, len(g["nodos"]), g["truncado"])
    print(f"\nCarga de 3 saltos (segundos, nodos, truncado): {tiempos}")
    for segundos, nodos, _ in tiempos.values():
        assert segundos < 3.0 and nodos <= grafo.MAX_NODOS
    # Desde el fondo, a 3 saltos hay 306 descripciones: se corta en el máximo y se avisa.
    assert tiempos["fondo"][1] == grafo.MAX_NODOS and tiempos["fondo"][2] is True


def test_la_parte_documental_es_constitutiva_como_en_la_exportacion(cliente, db, fondo_descrito, archivista, admin):
    """Una parte cuelga del documento por R003 (tiene la parte), no por
    inclusión: el grafo no puede decir algo distinto de la exportación."""
    f = fondo_descrito
    parte = RecursoDocumental(id=uuid.uuid4(), nivel="parte_documental", titulo="Sello de la Alcaldía", fondo_id=f["fondo"].id,
                              incluido_en_id=f["o114"].id, publicado_en=ahora())
    db.add(parte)
    db.flush()
    db.add(Relacion(origen_tipo="recurso_documental", origen_id=f["o114"].id, destino_tipo="recurso_documental",
                    destino_id=parte.id, tipo_relacion="inclusion", codigo_ric="has_or_had_constituent", origen="persona"))
    db.commit()
    g = pedir(cliente, archivista, "recurso_documental", parte.id)
    hacia_padre = [a for a in g["aristas"] if a["desde"] == f"recurso_documental:{f['o114'].id}"]
    assert [(a["codigo_ric"], a["uri_rico"]) for a in hacia_padre] == [("has_or_had_constituent", "rico:hasOrHadConstituent")]
    assert next(n for n in g["nodos"] if n["etiqueta"] == "Sello de la Alcaldía")["familia"] == "RecordPart"


def test_el_panel_de_filtros_ofrece_solo_las_relaciones_que_el_fondo_usa(cliente, fondo_descrito, archivista):
    o = cliente.get("/api/grafo/opciones", headers=archivista, params={"fondo_id": str(fondo_descrito["fondo"].id)}).json()
    codigos = {r["clave"] for r in o["relaciones"]}
    assert codigos == {"includes_or_included", "has_creator", "has_addressee", "has_or_had_subject",
                       "has_or_had_instantiation", "is_creation_date_of", "forma_documental"}
    assert all(r["uri_rico"] for r in o["relaciones"])
    assert {f["clave"] for f in o["familias"]} >= {"RecordSet", "Record", "Agent", "Place", "Activity", "Date", "Instantiation"}
    raices = [r["etiqueta"] for r in o["raices"]]
    assert raices[0] == "Correspondencia municipal" and "Borrador" not in raices


# --- Resumen del fondo (portada del catálogo) -----------------------------------------------------


def test_resumen_del_fondo_y_coherencia_de_fechas(cliente, db, fondo_descrito, archivista, consulta, admin):
    f = fondo_descrito
    fondo = db.get(RecursoDocumental, f["fondo"].id)
    fondo.fechas_extremas = "1968–1975"  # declaradas al registrar; los documentos son de 1948 y 1949
    db.commit()
    r = cliente.get("/api/instrumentos/resumen", headers=archivista, params={"fondo_id": str(fondo.id)}).json()
    assert {n["nombre"]: n["cantidad"] for n in r["niveles"]} == {"Serie": 1, "Expediente": 2, "Unidad documental": 2}
    assert r["archivos"] == 1 and r["descripciones"] == 5
    assert r["productores"][0] == {"id": str(f["alcaldia"].id), "nombre": ALCALDIA, "documentos": 3}
    assert [x["nombre"] for x in r["lugares"]] == ["Boyacá"] and [x["nombre"] for x in r["formas"]] == ["Oficio"]
    assert r["fechas"] == {"declaradas": "1968–1975", "documentos": "1948–1949", "coinciden": False}
    # Lo reservado no cuenta para quien no es archivista.
    db.add(DeclaracionDerechos(fondo_id=fondo.id, entidad_tipo="recurso_documental", entidad_id=f["exp48"].id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()
    lector = cliente.get("/api/instrumentos/resumen", headers=consulta, params={"fondo_id": str(fondo.id)}).json()
    assert lector["descripciones"] == 2 and lector["archivos"] == 0


def test_el_indice_de_consulta_no_cuenta_lo_reservado(cliente, db, fondo_descrito, archivista, consulta, admin):
    f = fondo_descrito
    db.add(DeclaracionDerechos(fondo_id=f["fondo"].id, entidad_tipo="recurso_documental", entidad_id=f["exp48"].id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()

    def alcaldia(quien):
        datos = cliente.get("/api/instrumentos/indice", headers=quien, params={"fondo_id": str(f["fondo"].id)}).json()
        return next(e for g in datos["grupos"] for x in g["letras"] for e in x["entidades"] if e["nombre"] == ALCALDIA)

    assert alcaldia(archivista)["documentos"] == 3
    lector = alcaldia(consulta)
    assert lector["documentos"] == 1 and [d["titulo"] for d in lector["descripciones"]] == [EXP49]


# --- Lo clasificado o reservado no sale al perfil de consulta por ningún camino -------------------


def test_el_perfil_de_consulta_no_ve_lo_reservado_en_el_catalogo_ni_en_las_fichas(
        cliente, db, fondo_descrito, archivista, consulta, admin):
    """Ley 1712 de 2014, art. 18 y 19: la reserva del expediente 1948 la
    heredan sus dos oficios. El catálogo, la ficha, la ficha pública, los
    conteos de documentos y el selector del grafo la respetan."""
    f = fondo_descrito
    db.add(DeclaracionDerechos(fondo_id=f["fondo"].id, entidad_tipo="recurso_documental", entidad_id=f["exp48"].id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()
    serie = {"fondo_id": str(f["fondo"].id), "nodo_id": str(f["serie"].id)}
    hijos = lambda quien: [h["titulo"] for h in cliente.get("/api/instrumentos/catalogo", headers=quien, params=serie).json()["hijos"]]
    assert hijos(archivista) == [EXP48, EXP49] and hijos(consulta) == [EXP49]
    for ruta in (f"/api/instrumentos/catalogo/{f['o114'].id}", f"/api/catalogo/registros/{f['o114'].id}"):
        assert cliente.get(ruta, headers=consulta).status_code == 404, ruta
        assert cliente.get(ruta, headers=archivista).status_code == 200, ruta
    assert cliente.get("/api/instrumentos/catalogo", headers=consulta,
                       params={"fondo_id": str(f["fondo"].id), "nodo_id": str(f["exp48"].id)}).status_code == 404
    # «N documentos de esta entidad» no cuenta lo reservado.
    ficha = cliente.get(f"/api/instrumentos/catalogo/{f['exp49'].id}", headers=consulta).json()
    assert next(e for e in ficha["entidades"] if e["valor"] == ALCALDIA)["documentos"] == 1
    ficha = cliente.get(f"/api/instrumentos/catalogo/{f['exp49'].id}", headers=archivista).json()
    assert next(e for e in ficha["entidades"] if e["valor"] == ALCALDIA)["documentos"] == 3
    # Quien solo aparece en lo reservado (el Gobernador, destinatario del oficio 114) no se ofrece ni se dibuja.
    raices = lambda quien: {r["etiqueta"] for r in cliente.get("/api/grafo/opciones", headers=quien,
                                                               params={"fondo_id": str(f["fondo"].id)}).json()["raices"]}
    assert "Gobernador del Departamento" in raices(archivista)
    assert "Gobernador del Departamento" not in raices(consulta) and ALCALDIA in raices(consulta)
    gobernador = db.scalar(select(EntidadVocabulario).where(EntidadVocabulario.nombre == "Gobernador del Departamento"))
    assert cliente.get(f"/api/grafo/entidad_vocabulario/{gobernador.id}", headers=consulta).status_code == 404


def test_la_version_de_conservacion_cuelga_de_su_original_y_no_parece_un_duplicado(cliente, db, fondo_descrito, archivista):
    """Una migración a PDF/A crea una segunda instanciación. RiC-O la declara
    del documento y del original; el grafo la dibuja solo desde el original
    («migrada a») y la rotula como versión de conservación."""
    from app.models.instanciacion import Instanciacion

    f = fondo_descrito
    original = db.scalar(select(Instanciacion).where(Instanciacion.nombre_original == "Oficio_114_1948.pdf"))
    copia = Instanciacion(id=uuid.uuid4(), fondo_id=f["fondo"].id, nombre_original="Oficio_114_1948 (PDF/A-2b).pdf",
                          ruta="x/z.pdf", tamano_bytes=12, estado="listo_para_descripcion", huella="b" * 64,
                          formato_puid="fmt/476", derivada_de_id=original.id)
    db.add(copia)
    db.flush()
    for o_tipo, o_id, codigo in (("recurso_documental", f["o114"].id, "has_or_had_instantiation"),
                                 ("instanciacion", original.id, "migrated_into")):
        db.add(Relacion(origen_tipo=o_tipo, origen_id=o_id, destino_tipo="instanciacion", destino_id=copia.id,
                        tipo_relacion="asociacion", codigo_ric=codigo, origen="persona"))
    db.commit()
    g = pedir(cliente, archivista, "recurso_documental", f["o114"].id, saltos=2)
    hacia_copia = [(a["desde"].split(":")[0], a["codigo_ric"]) for a in g["aristas"] if a["hacia"] == f"instanciacion:{copia.id}"]
    assert hacia_copia == [("instanciacion", "migrated_into")]
    assert next(n for n in g["nodos"] if n["id"] == str(copia.id))["subtitulo"] == "Versión de conservación"
