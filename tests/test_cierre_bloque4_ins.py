"""
Cierre de la auditoría RiC, bloque 4 (instrumentos): INS-01, INS-03,
INS-04, INS-05, INS-06 e INS-08. Cada salida pública se prueba con un
documento reservado para comprobar que no lo deja pasar.
"""

import io
import json
from datetime import date, timedelta

import pytest
from lxml import etree
from openpyxl import load_workbook
from rdflib import Graph

from app.core.config import settings
from app.models.preservacion import DeclaracionDerechos
from app.servicios import instrumentos, intercambio, parametros
from tests.test_cierre_ins02 import _reservar
from tests.test_exportacion_rico import fondo_rico, u  # noqa: F401
from tests.test_instrumentos import archivista, fondo_descrito  # noqa: F401

SECRETO = "Oficio N.º 115"  # el título del documento que se reserva en cada prueba


@pytest.fixture()
def con_reserva(db, fondo_rico, admin):
    f = fondo_rico
    _reservar(db, f, "recurso_documental", f["o115"].id, admin)
    db.commit()
    return f


def _encender(db, admin):
    parametros.cambiar(db, "rdf_uris_publicas", True, admin.id, "instrumentos")
    db.commit()


# --- INS-01 · descarga sin sesión del subconjunto público y URI estables ---------------------------


def test_rdf_publico_sin_sesion_solo_si_esta_encendido_y_sin_lo_reservado(cliente, db, con_reserva, admin):
    f = con_reserva
    ruta = f"/api/publico/rdf?fondo_id={f['fondo'].id}"
    assert cliente.get(ruta).status_code == 401  # apagado por defecto (decisión INS-01)
    _encender(db, admin)
    r = cliente.get(ruta)
    assert r.status_code == 200 and r.headers["x-ricora-conformidad"] == "conforme"
    g = Graph().parse(data=r.content, format="turtle")
    assert (u(f["o114"].id), None, None) in g
    assert (u(f["o115"].id), None, None) not in g and SECRETO not in r.text
    r = cliente.get(ruta + "&formato=jsonld")
    assert r.headers["content-type"].startswith("application/ld+json") and SECRETO not in r.text


def test_publicar_uris_exige_dominio_y_no_cambia_la_base_sin_confirmarlo(cliente, db, cabeceras_admin, monkeypatch):
    monkeypatch.setattr(settings, "url_publica", "http://203.0.113.7")
    r = cliente.put("/api/exportacion/uris-publicas", headers=cabeceras_admin, json={"publicas": True})
    assert r.status_code == 422 and "dominio" in r.json()["detail"]
    monkeypatch.setattr(settings, "url_publica", "https://archivo.tunja.gov.co")
    assert cliente.put("/api/exportacion/uris-publicas", headers=cabeceras_admin, json={"publicas": True}).status_code == 200
    assert parametros.leer(db, "rdf_base_publicada") == "https://archivo.tunja.gov.co/id/"
    monkeypatch.setattr(settings, "url_publica", "https://otro.dominio.co")
    r = cliente.put("/api/exportacion/uris-publicas", headers=cabeceras_admin, json={"publicas": True})
    assert r.status_code == 409 and "archivo.tunja.gov.co" in r.json()["detail"]
    assert cliente.put("/api/exportacion/uris-publicas", headers=cabeceras_admin,
                       json={"publicas": True, "confirmar_nueva_base": True}).status_code == 200


# --- INS-03 · la guía no lleva lo reservado; el FUID lo marca ------------------------------------------


def test_la_guia_no_incluye_lo_reservado_ni_lo_envia_al_motor(db, fondo_rico, admin):
    f = fondo_rico
    serie_secreta = f["subserie"]
    serie_secreta.titulo, serie_secreta.alcance_contenido = "Subserie SECRETA", "Contenido reservado"
    _reservar(db, f, "recurso_documental", serie_secreta.id, admin)
    db.commit()
    datos = instrumentos.datos_guia(db, f["fondo"])
    texto = json.dumps(datos, ensure_ascii=False)
    assert "SECRETA" not in texto and "Contenido reservado" not in texto
    assert any("Correspondencia" == n["titulo"] for n in datos["niveles"])


def test_el_fuid_marca_lo_clasificado_o_reservado(cliente, db, con_reserva, archivista):
    f = con_reserva
    previa = cliente.post("/api/instrumentos/inventario/vista-previa", headers=archivista,
                          json={"recurso_id": str(f["exp48"].id)}).json()
    acceso = {x["valores"]["codigo"]: x["valores"]["acceso"] for x in previa["filas"]}
    assert acceso == {"CO-AM-114": "Pública", "CO-AM-115": "Reservada"}


# --- INS-04 · SPARQL de solo lectura sobre el subconjunto público -------------------------------------


def test_sparql_responde_sobre_lo_publico_y_no_ve_lo_reservado(cliente, db, con_reserva, admin):
    f = con_reserva
    ruta = f"/api/publico/sparql?fondo_id={f['fondo'].id}"
    fondo = {"fondo_id": str(f["fondo"].id)}
    consulta = "PREFIX rico: <https://www.ica.org/standards/RiC/ontology#> SELECT ?t WHERE { ?r rico:title ?t }"
    assert cliente.get("/api/publico/sparql", params=fondo | {"query": consulta}).status_code == 401
    _encender(db, admin)
    r = cliente.get("/api/publico/sparql", params=fondo | {"query": consulta})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/sparql-results+json")
    titulos = {b["t"]["value"] for b in r.json()["results"]["bindings"]}
    assert "Oficio N.º 114" in titulos and SECRETO not in titulos
    # Por POST, como manda el protocolo SPARQL.
    r = cliente.post(ruta, content=consulta.encode(), headers={"Content-Type": "application/sparql-query"})
    assert r.status_code == 200 and SECRETO not in r.text
    ask = cliente.get("/api/publico/sparql", params=fondo | {"query": f'ASK {{ ?r ?p "{SECRETO}" }}'}).json()
    assert ask["boolean"] is False


@pytest.mark.parametrize("consulta", [
    "SELECT * WHERE { SERVICE <http://example.org/sparql> { ?s ?p ?o } }",
    "INSERT DATA { <http://x> <http://y> <http://z> }",
    "LOAD <http://example.org/datos.ttl>",
    "DROP ALL",
])
def test_sparql_rechaza_lo_que_no_es_lectura(cliente, db, fondo_rico, admin, consulta):
    _encender(db, admin)
    r = cliente.get("/api/publico/sparql", params={"fondo_id": str(fondo_rico["fondo"].id), "query": consulta})
    assert r.status_code in (400, 403), r.text


# --- INS-05 · EAD3, EAC-CPF e IIIF ----------------------------------------------------------------------


def test_ead3_del_fondo_valida_contra_el_esquema_y_no_lleva_lo_reservado(cliente, db, con_reserva, archivista):
    f = con_reserva
    r = cliente.get(f"/api/publico/ead3?fondo_id={f['fondo'].id}", headers=archivista)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/xml")
    assert intercambio.validar(r.content, "ead3.xsd") == []
    raiz = etree.fromstring(r.content)
    ns = {"e": intercambio.EAD}
    titulos = [t.text for t in raiz.iterfind(".//e:unittitle", ns)]
    assert "Correspondencia municipal" in titulos and "Oficio N.º 114" in titulos and SECRETO not in titulos
    assert raiz.find("e:archdesc", ns).get("level") == "fonds"


def test_eac_cpf_del_agente_valida_contra_el_esquema(cliente, db, fondo_rico, archivista):
    f = fondo_rico
    r = cliente.get(f"/api/publico/eac/{f['alcaldia'].id}", headers=archivista)
    assert r.status_code == 200
    assert intercambio.validar(r.content, "eac.xsd") == []
    raiz = etree.fromstring(r.content)
    ns = {"c": intercambio.EAC}
    nombres = [p.text for p in raiz.iterfind(".//c:nameEntry/c:part", ns)]
    assert nombres[0] == "Alcaldía Municipal" and "Cabildo de Tunja" in nombres
    assert raiz.find(".//c:entityType", ns).get("value") == "corporateBody"
    assert raiz.find(".//c:identityId", ns).text == "Q1000000"
    # Un mecanismo no es una autoridad EAC-CPF.
    from app.servicios import vocabulario

    m = vocabulario.mecanismo(db, fondo_id=f["fondo"].id, nombre="Ghostscript", version="10.02.1")
    db.commit()
    assert cliente.get(f"/api/publico/eac/{m.id}", headers=archivista).status_code == 422


def test_un_validador_de_esquema_detecta_un_documento_mal_formado():
    assert intercambio.validar(b'<ead xmlns="http://ead3.archivists.org/schema/"><control/></ead>', "ead3.xsd")


def test_manifiesto_iiif_de_una_descripcion_publica(cliente, db, fondo_descrito, admin):
    from app.models.instanciacion import Instanciacion
    from app.servicios import almacen
    from tests import archivos

    f = fondo_descrito
    inst = db.query(Instanciacion).filter(Instanciacion.nombre_original == "Oficio_114_1948.pdf").one()
    destino = almacen.raiz() / "iiif-prueba.png"
    destino.write_bytes(archivos.png_con_texto())
    inst.ruta, inst.nombre_original, inst.formato_puid = "iiif-prueba.png", "Oficio_114.png", "fmt/11"
    db.commit()
    _encender(db, admin)
    r = cliente.get(f"/api/publico/iiif/{f['o114'].id}/manifest")
    assert r.status_code == 200
    m = r.json()
    assert m["@context"] == "http://iiif.io/api/presentation/3/context.json" and m["type"] == "Manifest"
    assert m["label"] == {"es": ["Oficio N.º 114"]}
    [lienzo] = m["items"]
    cuerpo = lienzo["items"][0]["items"][0]["body"]
    assert lienzo["type"] == "Canvas" and lienzo["width"] > 0 and cuerpo["format"] == "image/png"
    imagen = cliente.get(cuerpo["id"].replace(settings.url_publica, ""))
    assert imagen.status_code == 200 and imagen.content[:8] == b"\x89PNG\r\n\x1a\n"
    # Reservado: ni manifiesto ni imagen.
    db.add(DeclaracionDerechos(fondo_id=f["fondo"].id, entidad_tipo="instanciacion", entidad_id=inst.id,
                               base="estatuto", acceso="reservado", reproduccion="no_permitida",
                               fundamento="Ley 1712 de 2014, art. 19", creada_por_id=admin.id))
    db.commit()
    assert cliente.get(f"/api/publico/iiif/{f['o114'].id}/manifest").json()["items"] == []
    assert cliente.get(cuerpo["id"].replace(settings.url_publica, "")).status_code == 404


# --- INS-06 · FUID completo -------------------------------------------------------------------------------


def test_fuid_con_unidades_de_conservacion_frecuencia_encabezado_y_firmas(cliente, db, fondo_descrito, archivista):
    f = fondo_descrito
    f["o115"].tomo, f["o115"].frecuencia_consulta = "T-2", "baja"
    db.commit()
    r = cliente.post("/api/instrumentos/inventario", headers=archivista, json={
        "recurso_id": str(f["exp48"].id), "encabezado": {
            "entidad_remitente": "Secretaría de Gobierno", "entidad_productora": "Alcaldía de Tunja",
            "unidad_administrativa": "Despacho", "oficina_productora": "Correspondencia",
            "objeto": "transferencia_secundaria", "registro_entrada": "2026 10 03 T-014",
            "elaborado_por": {"nombre": "Catalina Torres", "cargo": "Archivista", "lugar": "Tunja", "fecha": "2026-10-03"},
            "recibido_por": {"nombre": "Pedro Gómez", "cargo": "Jefe de archivo"}}})
    assert r.status_code == 200
    hoja = load_workbook(io.BytesIO(r.content)).active
    valores = [[c.value for c in fila] for fila in hoja.iter_rows()]
    planos = {v for fila in valores for v in fila if v is not None}
    for esperado in ("Secretaría de Gobierno", "Alcaldía de Tunja", "Despacho", "Correspondencia",
                     "Transferencia secundaria", "2026 10 03 T-014", "Elaborado por", "Entregado por", "Recibido por",
                     "Catalina Torres", "Pedro Gómez", "Tomo", "Otro", "Frecuencia de consulta"):
        assert esperado in planos, esperado
    titulos = next(fila for fila in valores if fila[0] == "N.º de orden")
    fila_115 = next(fila for fila in valores if fila[1] == "CO-AM-115")
    assert fila_115[titulos.index("Tomo")] == "T-2" and fila_115[titulos.index("Frecuencia de consulta")] == "Baja"
    # Con tomo, la caja que falta ya no es un pendiente.
    assert r.headers["X-Campos-Pendientes"] == "0"


# --- INS-08 · índice de información clasificada y reservada ---------------------------------------------


def test_indice_ley_1712_lista_lo_reservado_con_fundamento_y_plazo(cliente, db, fondo_rico, admin, archivista):
    f = fondo_rico
    _reservar(db, f, "recurso_documental", f["o115"].id, admin, hasta=date.today() + timedelta(days=365))
    _reservar(db, f, "recurso_documental", f["exp49"].id, admin, hasta=date.today() - timedelta(days=1))
    db.commit()
    r = cliente.get(f"/api/publico/ley1712?fondo_id={f['fondo'].id}", headers=archivista).json()
    por_titulo = {x["titulo"]: x for x in r["filas"]}
    assert por_titulo[SECRETO]["fundamento"] == "Ley 1712 de 2014, art. 19"
    assert por_titulo[SECRETO]["categoria"].startswith("Información pública reservada")
    assert por_titulo[SECRETO]["estado"] == "Vigente" and por_titulo[SECRETO]["excepcion"] == "Total"
    assert por_titulo["Correspondencia 1949"]["estado"].startswith("Vencida")
    assert "Informa sobre" not in json.dumps(r, ensure_ascii=False)  # no revela el contenido
    x = cliente.get(f"/api/publico/ley1712?fondo_id={f['fondo'].id}&formato=xlsx", headers=archivista)
    hoja = load_workbook(io.BytesIO(x.content)).active
    assert hoja["A4"].value == "Nombre o título de la categoría de información"
