"""
El CCD y la TRD en el formato en que los entrega la entidad.

- Las columnas se reconocen por su nombre: una fila por serie, subserie o tipo
  documental; la jerarquía de dependencias en varias columnas (sección,
  subsección 1, 2, 3…); el encabezado en una o dos filas; celdas combinadas.
- Se leen Excel, PDF con texto, Word, CSV y un .zip con varios de ellos.
- Lo que el archivo no dice o dice de dos maneras (una dependencia sin código,
  una serie con dos códigos) queda marcado en la vista previa, no se inventa.
- Una estructura de tres niveles (Secretaría General › Subdirección › Grupo)
  llega al cuadro aplicado con subsecciones anidadas, conforme a RiC-O.
"""

import io
import zipfile

from openpyxl import Workbook
from sqlalchemy import select

from app.models.recurso_documental import RecursoDocumental
from app.servicios import cuadro
from tests.test_descripcion import archivista, eventos, fondo, motor_prueba  # noqa: F401
from tests.test_parametrizacion_trd import aplicar_trd

# Como el CCD del Ministerio de Ambiente: FONDO, SUBFONDO, cuatro niveles de dependencia, un «CÓDIGO» suelto
# (el de la dependencia productora), SERIE, «CÓDIG» (truncado en el PDF), SUBSERIE, CÓDIGO.
CCD_JERARQUICO = """FONDO;SUBFONDO;SECCIÓN / UNIDAD;SUBSECCIÓN 1 /;SUBSECCIÓN 2 /;SUBSECCIÓN 3 /;CÓDIGO;SERIE;CÓDIG;SUBSERIE;CÓDIGO
MINISTERIO;N/A;DESPACHO DEL MINISTRO;;;;1000;ACTAS;02;Actas de Comités Interinstitucionales;04
MINISTERIO;N/A;DESPACHO DEL MINISTRO;;;;1000;DERECHOS DE PETICIÓN;14;;
MINISTERIO;N/A;SECRETARÍA GENERAL;;;;4000;INFORMES;34;Informes de Gestión;02
MINISTERIO;N/A;SECRETARÍA GENERAL;SUBDIRECCIÓN ADMINISTRATIVA Y FINANCIERA;;;4100;PLANES;33;Plan Anual de Adquisiciones;05
MINISTERIO;N/A;SECRETARÍA GENERAL;SUBDIRECCIÓN ADMINISTRATIVA Y FINANCIERA;;GRUPO DE GESTIÓN DOCUMENTAL;4106;PLANES;33;Planes de Transferencias Documentales Primarias;20
MINISTERIO;N/A;SECRETARÍA GENERAL;SUBDIRECCIÓN ADMINISTRATIVA Y FINANCIERA;;GRUPO DE GESTIÓN DOCUMENTAL;4106;PLANES;33;Planes de Preservación Digital a Largo Plazo;16
"""


def _por(vista, serie, subserie=None):
    return next(f for f in vista["filas"] if f["serie"]["nombre"] == serie
                and (f["subserie"] or {}).get("nombre") == subserie)


def test_ccd_jerarquico_con_varios_niveles_de_dependencias():
    v = cuadro.vista_previa(CCD_JERARQUICO.encode("utf-8"), "ccd.csv")
    assert v["origen"].startswith("csv · columnas reconocidas")
    assert v["resumen"]["con_problemas"] == 0 and v["resumen"]["filas"] == 6
    grupo = _por(v, "PLANES", "Planes de Transferencias Documentales Primarias")
    assert grupo["seccion"] == {"codigo": "4000", "nombre": "SECRETARÍA GENERAL"}
    assert grupo["intermedias"] == [{"codigo": "4100", "nombre": "SUBDIRECCIÓN ADMINISTRATIVA Y FINANCIERA"}]
    assert grupo["subseccion"] == {"codigo": "4106", "nombre": "GRUPO DE GESTIÓN DOCUMENTAL"}
    # El código de la oficina ya es completo (4106, no 4000.4100.4106); la serie se compone desde él.
    assert grupo["codigos"] == {"seccion": "4000", "subseccion": "4106", "serie": "4106.33", "subserie": "4106.33.20"}
    peticiones = _por(v, "DERECHOS DE PETICIÓN")
    assert peticiones["subseccion"] is None and peticiones["codigos"]["serie"] == "1000.14"


def test_lo_que_el_archivo_no_dice_queda_marcado():
    sin_codigos = CCD_JERARQUICO.replace(";4106;", ";;").replace("PLANES;33;Plan Anual", "PLANES;47;Plan Anual")
    sin_codigos += "MINISTERIO;N/A;SECRETARÍA GENERAL;;;;4000;INFORMES;27;Informes de Gestión;02\n"
    v = cuadro.vista_previa(sin_codigos.encode("utf-8"), "ccd.csv")
    grupo = _por(v, "PLANES", "Planes de Preservación Digital a Largo Plazo")
    assert grupo["subseccion"]["codigo"] is None
    assert any("GRUPO DE GESTIÓN DOCUMENTAL» no tiene código" in p for p in grupo["problemas"])
    informes = _por(v, "INFORMES", "Informes de Gestión")
    assert informes["serie"]["codigo"] is None
    assert any("varios códigos (27, 34)" in p for p in informes["problemas"])
    assert _por(v, "ACTAS", "Actas de Comités Interinstitucionales")["problemas"] == []


def _trd_excel_libre() -> bytes:
    """Una TRD en Excel a la manera de una entidad: encabezado en dos filas, una fila por tipo documental y
    celdas combinadas (la dependencia y la serie solo en la primera fila de su bloque)."""
    libro = Workbook()
    h = libro.active
    h.append(["TABLA DE RETENCIÓN DOCUMENTAL"])
    h.append(["Entidad productora: Concejo Municipal"])
    h.append(["Oficina productora", "Código oficina", "Código serie", "Serie documental", "Código subserie",
              "Subserie documental", "Tipos documentales", "Retención", None, "Disposición final", None, None, None,
              "Procedimiento"])
    h.append([None, None, None, None, None, None, None, "AG", "AC", "CT", "E", "M", "S", None])
    h.append(["Secretaría", "110", "2", "Actas", "1", "Actas de sesión", "Citación", 2, 8, "X", None, "X", None,
              "Conservar por su valor histórico."])
    h.append([None, None, None, None, None, None, "Acta", None, None, None, None, None, None, None])
    h.append([None, None, "15", "Derechos de petición", None, None, "Petición; Respuesta", 1, 4, None, "X", None,
              None, "Eliminar en el archivo central."])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_trd_en_excel_con_encabezado_en_dos_filas_y_celdas_combinadas():
    v = cuadro.vista_previa(_trd_excel_libre(), "TRD Secretaría.xlsx")
    assert v["resumen"]["con_problemas"] == 0, [f["problemas"] for f in v["filas"]]
    actas = _por(v, "Actas", "Actas de sesión")
    assert actas["seccion"] == {"codigo": "110", "nombre": "Secretaría"}
    assert actas["tipos"] == ["Citación", "Acta"]
    assert (actas["gestion"], actas["central"], actas["disposicion"], actas["otras_disposiciones"]) == (2, 8, "ct", ["m"])
    peticiones = _por(v, "Derechos de petición")
    assert peticiones["codigos"]["serie"] == "110.15" and peticiones["tipos"] == ["Petición", "Respuesta"]
    assert (peticiones["gestion"], peticiones["central"], peticiones["disposicion"]) == (1, 4, "e")


def _trd_word() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("TABLA DE RETENCIÓN DOCUMENTAL")
    d.add_paragraph("Dependencia productora: Oficina Jurídica   Código: 120")
    t = d.add_table(rows=1, cols=7)
    for celda, texto in zip(t.rows[0].cells, ("Código", "Serie", "Subserie", "Tipos documentales", "AG", "AC", "CT")):
        celda.text = texto
    fila = t.add_row().cells
    for celda, texto in zip(fila, ("5", "Procesos judiciales", "", "Demanda\nFallo", "5", "15", "X")):
        celda.text = texto
    salida = io.BytesIO()
    d.save(salida)
    return salida.getvalue()


def test_trd_en_word_con_la_dependencia_en_el_parrafo_del_encabezado():
    v = cuadro.vista_previa(_trd_word(), "trd-juridica.docx")
    assert v["origen"].startswith("word")
    f = _por(v, "Procesos judiciales")
    assert f["seccion"] == {"codigo": "120", "nombre": "Oficina Jurídica"}
    assert f["codigos"]["serie"] == "120.5" and f["tipos"] == ["Demanda", "Fallo"] and f["disposicion"] == "ct"
    assert f["problemas"] == []


def test_un_zip_con_una_trd_por_dependencia():
    paquete = io.BytesIO()
    with zipfile.ZipFile(paquete, "w") as z:
        z.writestr("TRD/01 Secretaria.xlsx", _trd_excel_libre())
        z.writestr("TRD/02 Juridica.docx", _trd_word())
        z.writestr("TRD/Leame.txt", "Tablas aprobadas.")
        z.writestr("__MACOSX/._01 Secretaria.xlsx", b"x")
    v = cuadro.vista_previa(paquete.getvalue(), "TRD-V2.zip")
    assert v["origen"] == "zip · 2 archivo(s) leídos"
    assert {f["seccion"]["nombre"] for f in v["filas"]} == {"Secretaría", "Oficina Jurídica"}
    assert v["resumen"]["con_problemas"] == 0


def test_formatos_que_no_se_leen_lo_dicen():
    for contenido, nombre, texto in ((b"abc", "trd.xls", "guárdelo como .xlsx"),
                                     (b"PK\x03\x04roto", "trd.zip", "dañado"),
                                     (b"Nada;que;ver\n1;2;3\n", "otra.csv", "columna de la serie")):
        try:
            cuadro.vista_previa(contenido, nombre)
        except cuadro.ErrorCuadro as exc:
            assert texto in str(exc), str(exc)
        else:
            raise AssertionError(nombre)


def test_tres_niveles_llegan_al_cuadro_aplicado(cliente, db, fondo, cabeceras_admin):  # noqa: F811
    v = cuadro.vista_previa(CCD_JERARQUICO.encode("utf-8"), "ccd.csv")
    # Un CCD no trae retención: el cuadro se arma igual; las reglas llegan con la TRD.
    r = aplicar_trd(cliente, cabeceras_admin, fondo, v["filas"])
    assert r.status_code == 200, r.text
    assert r.json()["validacion"]["conforme"]
    deps = {d["nombre"]: d for d in cliente.get("/api/trd/dependencias", headers=cabeceras_admin,
                                                 params={"fondo_id": str(fondo.id)}).json()}
    assert deps["GRUPO DE GESTIÓN DOCUMENTAL"]["codigo_completo"] == "4106"
    assert deps["GRUPO DE GESTIÓN DOCUMENTAL"]["superior_id"] == deps["SUBDIRECCIÓN ADMINISTRATIVA Y FINANCIERA"]["id"]
    grupo = db.scalar(select(RecursoDocumental).where(RecursoDocumental.fondo_id == fondo.id,
                                                      RecursoDocumental.codigo_referencia == "4106"))
    subdireccion = db.get(RecursoDocumental, grupo.incluido_en_id)
    secretaria = db.get(RecursoDocumental, subdireccion.incluido_en_id)
    assert (grupo.nivel, subdireccion.nivel, secretaria.nivel) == ("subseccion", "subseccion", "seccion")
    assert (subdireccion.codigo_referencia, secretaria.codigo_referencia) == ("4100", "4000")
    assert db.scalar(select(RecursoDocumental).where(RecursoDocumental.fondo_id == fondo.id,
                                                     RecursoDocumental.codigo_referencia == "4106.33.20")) is not None
