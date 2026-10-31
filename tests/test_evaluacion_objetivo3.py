"""
Objetivo 3 de la tesis convertido en lógica (documentacion/modulo-evaluacion.md, §25).

- El criterio de comparabilidad se fija antes de iniciar, queda congelado y
  el sistema lo aplica tipo por tipo, con el acuerdo entre archivistas como techo.
- Los documentos con información clasificada o reservada no entran a la evaluación.
- La supervisión de la condición asistida: aceptadas, corregidas, rechazadas y agregadas.
- Cada desacuerdo entre el motor y una referencia ciega tiene una clave y se
  clasifica con la taxonomía de errores: por separado y, si difieren, por consenso.
- Intervalos de confianza por remuestreo de documentos y prueba de Wilcoxon de los tiempos.
"""

import io
import math
from datetime import date, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.models.evaluacion import Anotacion
from app.servicios import evaluacion_analisis as an
from tests.conftest import crear_usuario, ingresar
from tests.test_descripcion import documento, fondo  # noqa: F401
from tests.test_evaluacion import CIEGA_A, CIEGA_B, TEXTO, anotar, evaluadora, motor_eval  # noqa: F401

# --- Funciones puras (calculadas a mano) -------------------------------------------------------------


def test_wilcoxon_calculado_a_mano():
    # Ocho diferencias positivas 1…8: W+ = 36, W− = 0; media 18, varianza 8·9·17/24 = 51.
    w = an.wilcoxon([(x + 1.0, 1.0) for x in range(1, 9)])
    assert (w["n"], w["w_mas"], w["w_menos"]) == (8, 36, 0)
    z = (0 - 18 + 0.5) / math.sqrt(51)
    assert w["z"] == round(z, 4) and abs(w["p"] - math.erfc(abs(z) / math.sqrt(2))) < 1e-4
    assert an.wilcoxon([(3, 2), (5, 1)])["p"] is None  # menos de 6 pares: no se calcula p
    assert an.wilcoxon([(2, 2)])["n"] == 0  # sin diferencias


def test_kappa_nominal_entre_clasificadoras():
    assert an.kappa_nominal([("papel", "papel"), ("omision", "omision")]) == 1.0
    assert an.kappa_nominal([("a", "a"), ("b", "b"), ("a", "b"), ("b", "a")]) == 0.0
    assert an.kappa_nominal([]) is None


def test_intervalo_por_remuestreo_de_documentos_es_reproducible():
    docs = [{"vp": 3, "fp": 1, "fn": 1}, {"vp": 1, "fp": 2, "fn": 0}, {"vp": 4, "fp": 0, "fn": 2}]
    a, b = an.intervalo_f(docs), an.intervalo_f(docs)
    assert a == b and a["documentos"] == 3 and a["inferior"] <= a["superior"]
    assert an.intervalo_f(docs[:1])["inferior"] is None  # con un documento no hay intervalo


def test_supervision_de_la_condicion_asistida():
    motor = [{"tipo": "agente", "valor": "Alcaldía Municipal de Tunja"}, {"tipo": "agente", "valor": "Concejo Municipal"},
             {"tipo": "lugar", "valor": "Tunja"}]
    asistida = [{"tipo": "agente", "valor": "Alcaldia Municipal de Tunja"},
                {"tipo": "agente", "valor": "Concejo Municipal de Tunja"}, {"tipo": "lugar", "valor": "Boyacá"}]
    s = an.supervision(motor, asistida, 0.75)
    assert s["agente"] == {"aceptada": 1, "corregida": 1, "rechazada": 0, "agregada": 0}
    assert s["lugar"] == {"aceptada": 0, "corregida": 0, "rechazada": 1, "agregada": 1}
    r = an.resumen_supervision(s)
    assert r["propuestas"] == 3 and r["pct_aceptadas"] == 33.3 and r["pct_agregadas"] == 33.3


def test_el_criterio_se_aplica_tipo_por_tipo():
    resultados = {
        "flexible": {"por_tipo": {"agente": {"vp": 9, "fp": 1, "fn": 3, "f1": 0.8},
                                  "fecha": {"vp": 10, "fp": 0, "fn": 0, "f1": 1.0},
                                  "mandato": {"vp": 1, "fp": 2, "fn": 9, "f1": 0.15},
                                  "lugar": {"vp": 2, "fp": 0, "fn": 1, "f1": 0.8}}},
        "acuerdo_por_tipo": {"agente": {"f1": 0.85}, "fecha": {"f1": 1.0}, "mandato": {"f1": 0.6}, "lugar": {"f1": 0.9}},
        "conformidad": {"asistida": {"descripciones": 4, "conformes": 4}},
        "rubrica": {"exactitud": {"media": 4}, "completitud": {"media": 3.5}, "pertinencia": {"media": 3}},
    }
    v = an.veredicto(resultados, None)  # criterio por defecto: margen 0,10; 10 casos; rúbrica 3; conformidad 100 %
    estados = {t["tipo"]: t["estado"] for t in v["tipos"]}
    assert estados == {"agente": "comparable", "fecha": "comparable", "mandato": "no_comparable",
                       "lugar": "datos_insuficientes"}
    mandato = next(t for t in v["tipos"] if t["tipo"] == "mandato")
    assert mandato["diferencia"] == 0.45 and "por debajo del acuerdo" in mandato["motivos"][0]
    # Si la rúbrica no alcanza, ningún tipo es comparable aunque la F sí lo sea.
    peor = resultados | {"rubrica": {"exactitud": {"media": 2.5}, "completitud": {"media": 4}, "pertinencia": {"media": 4}}}
    assert an.veredicto(peor, None)["comparables"] == []
    try:
        an.validar_criterio({"margen_f": 0.9})
        raise AssertionError("debió rechazarlo")
    except an.ErrorAnalisis as exc:
        assert "margen_f" in str(exc)


# --- Flujo completo -------------------------------------------------------------------------------------


def _preparar_con_criterio(cliente, db, fondo, adm, criterio):  # noqa: F811
    d = documento(db, fondo, "Oficio_210_1948.pdf", TEXTO)
    ev = cliente.post("/api/evaluacion", headers=adm, json={"fondo_id": str(fondo.id), "nombre": "Piloto objetivo 3",
                                                             "umbral_similitud": 0.75}).json()
    r = cliente.put(f"/api/evaluacion/{ev['id']}/criterio", headers=adm, json=criterio)
    assert r.status_code == 200 and r.json()["criterio"]["minimo_casos"] == criterio["minimo_casos"]
    cliente.post(f"/api/evaluacion/{ev['id']}/documentos", headers=adm, json={"instanciacion_ids": [str(d.id)]})
    cliente.post(f"/api/evaluacion/{ev['id']}/propuestas", headers=adm)
    assert cliente.post(f"/api/evaluacion/{ev['id']}/estado", headers=adm, json={"estado": "en_curso"}).status_code == 200
    return ev["id"], d


def test_un_documento_reservado_no_entra_a_la_evaluacion(cliente, db, fondo, cabeceras_admin, admin, motor_eval):  # noqa: F811
    from app.servicios import derechos

    d = documento(db, fondo, "Reservado.pdf", TEXTO)
    derechos.declarar(db, entidad_tipo="instanciacion", entidad_id=d.id, base="estatuto", acceso="reservado",
                      reproduccion="no_permitida", fundamento="Ley 1712 de 2014, art. 19", nota=None,
                      vigente_hasta=date(2040, 1, 1), usuario_id=admin.id)
    db.commit()
    ev = cliente.post("/api/evaluacion", headers=cabeceras_admin, json={"fondo_id": str(fondo.id), "nombre": "Con reserva"}).json()
    r = cliente.post(f"/api/evaluacion/{ev['id']}/documentos", headers=cabeceras_admin, json={"instanciacion_ids": [str(d.id)]})
    assert r.status_code == 422 and "reservada" in r.json()["detail"]


def test_objetivo_3_de_principio_a_fin(cliente, db, fondo, cabeceras_admin, motor_eval):  # noqa: F811
    adm = cabeceras_admin
    ev_id, d = _preparar_con_criterio(cliente, db, fondo, adm, {"margen_f": 0.10, "minimo_casos": 1,
                                                                 "rubrica_minima": 3, "conformidad_asistida": 1.0})
    # Con la evaluación iniciada, el criterio ya no cambia.
    r = cliente.put(f"/api/evaluacion/{ev_id}/criterio", headers=adm, json={"minimo_casos": 2})
    assert r.status_code == 409 and "antes de iniciar" in r.json()["detail"]

    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    luis = evaluadora(cliente, db, "luis@correo.com", "Luis Pardo")
    assert anotar(cliente, ana, ev_id, d.id, "ciega", CIEGA_A).status_code == 200
    assert anotar(cliente, luis, ev_id, d.id, "ciega", CIEGA_B).status_code == 200
    # Asistida de Ana: quita el lugar «Tunja», corrige el destinatario y agrega «Boyacá».
    asistida = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana,
                            json={"condicion": "asistida"}).json()["datos"]
    lugar = next(e for e in asistida["entidades"] if e["tipo"] == "lugar")
    prop_lugar = next(r["propiedad"] for r in asistida["relaciones"] if lugar["clave"] in (r["origen"], r["destino"]))
    entidades = [e for e in asistida["entidades"] if e["clave"] != lugar["clave"]]
    for e in entidades:
        if e["valor"] == "Concejo Municipal":
            e["valor"] = "Concejo Municipal de Tunja"
    entidades.append({"clave": "e9", "tipo": "lugar", "valor": "Boyacá", "subtipo": None, "edtf": None})
    relaciones = [r for r in asistida["relaciones"] if lugar["clave"] not in (r["origen"], r["destino"])]
    relaciones.append({"origen": "e9" if prop_lugar.startswith("is_or_was_location") else "documento",
                       "propiedad": prop_lugar, "destino": "documento" if prop_lugar.startswith("is_or_was_location") else "e9"})
    a = db.scalar(select(Anotacion).where(Anotacion.condicion == "asistida"))
    r = cliente.put(f"/api/evaluacion/anotaciones/{a.id}", headers=ana,
                    json=asistida | {"entidades": entidades, "relaciones": relaciones, "enviar": True})
    assert r.status_code == 200, r.text
    for cab, notas in ((ana, {"exactitud": 4, "completitud": 3, "pertinencia": 5}),
                       (luis, {"exactitud": 4, "completitud": 4, "pertinencia": 5})):
        cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/propuesta", headers=cab)
        cliente.put(f"/api/evaluacion/{ev_id}/documentos/{d.id}/calificacion", headers=cab, json=notas)
    ciega = db.scalar(select(Anotacion).where(Anotacion.condicion == "ciega", Anotacion.evaluador_id == a.evaluador_id))
    ciega.iniciada_en = ciega.enviada_en - timedelta(minutes=30)
    db.commit()

    # Clasificar antes de cerrar no se puede: los desacuerdos todavía pueden cambiar.
    res = cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=adm).json()
    clave = res["errores"]["lista"][0]["clave"]
    r = cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{clave}", headers=adm, json={"tipo_error": "papel"})
    assert r.status_code == 409

    # Supervisión de la asistida: Alcaldía, fecha y forma aceptadas; Concejo corregido; Tunja rechazado; Boyacá agregado.
    sup = res["supervision"]["total"]
    assert (sup["aceptada"], sup["corregida"], sup["rechazada"], sup["agregada"]) == (3, 1, 1, 1)
    # El criterio, tipo por tipo: el lugar no tiene techo (solo una archivista lo registró).
    comp = res["comparabilidad"]
    assert set(comp["comparables"]) == {"agente", "fecha", "forma_documental"}
    assert comp["sin_decidir"] == ["lugar"] and comp["criterio"]["minimo_casos"] == 1
    assert res["acuerdo_por_tipo"]["fecha"]["f1"] == 1.0
    # Tiempos pareados (Ana): un solo par, sin p.
    assert res["tiempos_por_documento"][0]["evaluadora"] == "Evaluadora 1"
    assert res["prueba_tiempos"]["n"] == 1 and res["prueba_tiempos"]["p"] is None
    assert res["intervalos"]["entidades_flexible"]["documentos"] == 1
    # Desacuerdos: Ana (variante Concejo, sobrante Tunja, faltante Boyacá, y las dos relaciones del lugar)
    # y Luis (sobrantes Concejo y Tunja, y sus dos relaciones).
    clases = res["errores"]["por_clase"]
    assert clases == {"variante": 1, "sobrante": 3, "faltante": 1, "relacion_sobrante": 3, "relacion_faltante": 1}

    assert cliente.post(f"/api/evaluacion/{ev_id}/estado", headers=adm, json={"estado": "cerrada"}).status_code == 200
    crear_usuario(db, "investigadora@correo.com", "administrador", nombre="Investigadora Dos")
    inv2 = ingresar(cliente, "investigadora@correo.com")
    lista = cliente.get(f"/api/evaluacion/{ev_id}/desacuerdos", headers=adm).json()
    assert len(lista["desacuerdos"]) == 9 and lista["resumen"]["sin_clasificar"] == 9
    claves = [x["clave"] for x in lista["desacuerdos"]]
    # Las claves no cambian entre consultas.
    assert claves == [x["clave"] for x in cliente.get(f"/api/evaluacion/{ev_id}/desacuerdos", headers=adm).json()["desacuerdos"]]

    # Primera investigadora: todo «omisión» salvo el primero, «papel».
    for i, k in enumerate(claves):
        assert cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{k}", headers=adm,
                           json={"tipo_error": "papel" if i == 0 else "omision"}).status_code == 200
    # La segunda no ve la clasificación de la primera hasta clasificar ella misma.
    vista = cliente.get(f"/api/evaluacion/{ev_id}/desacuerdos", headers=inv2).json()["desacuerdos"][0]
    assert vista["otra"] is None and vista["otra_clasifico"] is True
    # El consenso solo existe cuando dos difieren.
    r = cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{claves[1]}", headers=inv2,
                    json={"tipo_error": "omision", "ronda": "consenso"})
    assert r.status_code == 409
    for k in claves:
        cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{k}", headers=inv2, json={"tipo_error": "omision"})
    vista = cliente.get(f"/api/evaluacion/{ev_id}/desacuerdos", headers=inv2).json()
    primero = vista["desacuerdos"][0]
    assert primero["otra"] == "papel" and primero["requiere_consenso"] is True
    assert vista["resumen"]["pendientes_consenso"] == 1 and vista["resumen"]["clasificados"] == 8
    assert cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{claves[0]}", headers=inv2,
                       json={"tipo_error": "papel", "ronda": "consenso", "nota": "El motor confundió el papel."}).status_code == 200
    final = cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=adm).json()["errores"]
    assert final["clasificados"] == 9 and final["pendientes_consenso"] == 0 and final["pares"] == 9
    assert {x["tipo_error"]: x["veces"] for x in final["distribucion"]} == {"papel": 1, "omision": 8}
    assert final["kappa"] == 0.0  # una sola categoría para la segunda: el acuerdo no supera al azar
    assert cliente.put(f"/api/evaluacion/{ev_id}/desacuerdos/{claves[0]}", headers=adm,
                       json={"tipo_error": "inventado"}).status_code == 422

    hoja = load_workbook(io.BytesIO(cliente.get(f"/api/evaluacion/{ev_id}/resultados/hoja-de-calculo", headers=adm).content))
    assert {"Comparabilidad", "Supervisión (asistida)", "Tiempos por documento", "Desacuerdos"} <= set(hoja.sheetnames)
    assert [c.value for c in hoja["Desacuerdos"][2]][-1] in ("Papel equivocado", "Omisión")


def test_indicador_de_fragmento_no_hallado_en_las_decisiones():
    from app.servicios import decisiones_ia

    filas = [{"decision": "aceptada", "fragmento_localizado": True}, {"decision": "corregida", "fragmento_localizado": False},
             {"decision": "rechazada", "fragmento_localizado": False}, {"decision": "agregada", "fragmento_localizado": None}]
    c = decisiones_ia._conteos(filas)
    # Las decisiones sin el dato (anteriores al registro, o agregadas) no cuentan en la base.
    assert c["sin_fragmento"] == 2 and c["pct_sin_fragmento"] == 66.7
