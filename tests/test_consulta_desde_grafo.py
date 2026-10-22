"""
La consulta lee del grafo RiC-O (app/servicios/lectura.py).

- Si una tabla de trabajo se cambia por fuera del sistema, el catálogo, la
  ficha, el buscador y el FUID siguen mostrando lo que dice el grafo.
- Un cambio hecho por el sistema llega al grafo y de ahí a la consulta.
- La proyección solo se rehace cuando el grafo cambió.
- Si el almacén de tripletas no responde, la consulta no se cae.
"""

import io

from openpyxl import load_workbook
from sqlalchemy import text

from app.servicios import grafo_rdf, lectura
from tests.test_instrumentos import archivista, fondo_descrito  # noqa: F401  (fixtures)


def _por_fuera(db, recurso, titulo):
    """Un cambio directo en la tabla de trabajo, sin pasar por el sistema (no marca el grafo)."""
    db.execute(text("UPDATE public.recursos_documentales SET titulo = :t WHERE id = :i"),
               {"t": titulo, "i": recurso.id})


def test_si_las_tablas_difieren_la_consulta_muestra_el_grafo(cliente, db, fondo_descrito, archivista):  # noqa: F811
    f = fondo_descrito
    ficha = f"/api/instrumentos/catalogo/{f['o114'].id}"
    assert cliente.get(ficha, headers=archivista).json()["titulo"] == "Oficio N.º 114"
    _por_fuera(db, f["o114"], "Título cambiado por fuera")
    assert db.execute(text("SELECT titulo FROM public.recursos_documentales WHERE id = :i"),
                      {"i": f["o114"].id}).scalar() == "Título cambiado por fuera"
    # La ficha, el catálogo, el buscador y el FUID dicen lo del grafo.
    assert cliente.get(ficha, headers=archivista).json()["titulo"] == "Oficio N.º 114"
    nivel = cliente.get("/api/instrumentos/catalogo", headers=archivista,
                        params={"fondo_id": str(f["fondo"].id), "nodo_id": str(f["exp48"].id)}).json()
    assert "Oficio N.º 114" in [h["titulo"] for h in nivel["hijos"]]
    encontrados = cliente.get("/api/buscar", headers=archivista, params={"q": "cambiado por fuera"}).json()
    assert encontrados["total"] == 0
    assert cliente.get("/api/buscar", headers=archivista, params={"q": "Oficio 114"}).json()["total"] >= 1
    r = cliente.post("/api/instrumentos/inventario", headers=archivista, json={"recurso_id": str(f["exp48"].id)})
    filas = [[c.value for c in x] for x in load_workbook(io.BytesIO(r.content)).active.iter_rows()]
    assert any(x[2] == "Correspondencia / Oficio N.º 114" for x in filas)
    # La comparación periódica lo detecta: el grafo es la fuente y la tabla se corrige desde él.
    assert any(d["columna"] == "titulo" for d in grafo_rdf.comparar(db, f["fondo"].id))


def test_un_cambio_del_sistema_llega_a_la_consulta_y_solo_entonces_se_rehace(cliente, db, fondo_descrito,  # noqa: F811
                                                                             archivista):
    f = fondo_descrito
    assert lectura.preparar(db, f["fondo"])["rehecha"] in (True, False)
    assert lectura.preparar(db, f["fondo"])["rehecha"] is False  # el grafo no cambió
    f["o114"].titulo = "Oficio N.º 114 (corregido)"
    db.commit()  # por el sistema: marca el grafo del fondo
    r = lectura.preparar(db, f["fondo"])
    assert r["rehecha"] is True and r["segundos"] < 30
    ficha = cliente.get(f"/api/instrumentos/catalogo/{f['o114'].id}", headers=archivista).json()
    assert ficha["titulo"] == "Oficio N.º 114 (corregido)"


def test_sin_almacen_la_consulta_sigue_desde_las_tablas(cliente, db, fondo_descrito, archivista,  # noqa: F811
                                                        monkeypatch):
    f = fondo_descrito
    f["o114"].titulo = "Oficio N.º 114 bis"
    db.commit()

    def caido(*_a, **_k):
        raise grafo_rdf.ErrorGrafo("almacén caído")

    monkeypatch.setattr(grafo_rdf, "al_dia", caido)
    r = cliente.get(f"/api/instrumentos/catalogo/{f['o114'].id}", headers=archivista)
    assert r.status_code == 200 and r.json()["titulo"] == "Oficio N.º 114 bis"
