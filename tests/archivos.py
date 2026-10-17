"""Archivos de prueba generados en el momento (nada de datos reales)."""

import io
import os

from PIL import Image, ImageDraw, ImageFont

TEXTO_ACTA = "Acta del Concejo Municipal, sesion ordinaria del 20 de julio de 1948. Se aprobo el presupuesto."


def txt(texto: str = TEXTO_ACTA) -> bytes:
    return texto.encode("utf-8")


def pdf_con_texto(texto: str = TEXTO_ACTA) -> bytes:
    """PDF mínimo válido, de una página, con capa de texto."""
    flujo = f"BT /F1 12 Tf 50 750 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(flujo)).encode() + b" >>\nstream\n" + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = io.BytesIO()
    salida.write(b"%PDF-1.4\n")
    posiciones = []
    for i, obj in enumerate(objetos, start=1):
        posiciones.append(salida.tell())
        salida.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = salida.tell()
    salida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for p in posiciones:
        salida.write(f"{p:010d} 00000 n \n".encode())
    salida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return salida.getvalue()


def _imagen_con_texto(lineas=("ARCHIVO MUNICIPAL", "OFICIO NUMERO 114", "BOGOTA 1948")) -> Image.Image:
    img = Image.new("L", (1700, 700), 255)
    dibujo = ImageDraw.Draw(img)
    fuente = ImageFont.load_default(size=90)
    for i, linea in enumerate(lineas):
        dibujo.text((80, 80 + i * 180), linea, fill=0, font=fuente)
    return img


def png_con_texto() -> bytes:
    salida = io.BytesIO()
    _imagen_con_texto().save(salida, format="PNG")
    return salida.getvalue()


def png_degradado() -> bytes:
    """Imagen con el mismo texto, borrosa y con manchas, como una copia mala:
    Tesseract la lee, pero con confianza baja (alrededor de 56 sobre 100)."""
    import random

    from PIL import ImageFilter

    img = _imagen_con_texto().resize((850, 350)).filter(ImageFilter.GaussianBlur(2))
    azar = random.Random(7)
    px = img.load()
    for _ in range(15000):
        px[azar.randrange(850), azar.randrange(350)] = azar.choice((0, 255))
    salida = io.BytesIO()
    img.save(salida, format="PNG")
    return salida.getvalue()


def pdf_escaneado() -> bytes:
    """PDF sin capa de texto: una imagen de página, como sale de un escáner."""
    salida = io.BytesIO()
    _imagen_con_texto().convert("RGB").save(salida, format="PDF", resolution=150)
    return salida.getvalue()


def pdf_danado() -> bytes:
    return b"%PDF-1.4\n1 0 obj << /Type /Catalog basura sin terminar"


def desconocido() -> bytes:
    """Bytes sin ninguna firma conocida, con una extensión inventada."""
    return b"\x13\x37RICORA" + os.urandom(512)


def pdf_escaneado_paginas(paginas: list[list[str]]) -> bytes:
    """PDF escaneado de varias páginas (sin capa de texto): cada página es una
    imagen con sus líneas, desde arriba (RF-OCR-001)."""
    imagenes = []
    fuente = ImageFont.load_default(size=60)
    for lineas in paginas:
        img = Image.new("L", (1700, 2200), 255)
        dibujo = ImageDraw.Draw(img)
        for i, linea in enumerate(lineas):
            dibujo.text((120, 300 + i * 140), linea, fill=0, font=fuente)
        imagenes.append(img.convert("RGB"))
    salida = io.BytesIO()
    imagenes[0].save(salida, format="PDF", save_all=True, append_images=imagenes[1:], resolution=150)
    return salida.getvalue()


def pdf_con_paginas(textos: list[str]) -> bytes:
    """PDF válido de varias páginas con capa de texto: una línea por página,
    a 1/4 de la altura desde arriba."""
    objetos = [b"<< /Type /Catalog /Pages 2 0 R >>", None]
    hijos = []
    for texto in textos:
        flujo = f"BT /F1 14 Tf 72 594 Td ({texto}) Tj ET".encode("latin-1")
        n_pagina, n_flujo = len(objetos) + 1, len(objetos) + 2
        hijos.append(f"{n_pagina} 0 R")
        objetos.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n_flujo} 0 R "
                       f"/Resources << /Font << /F1 {len(textos) * 2 + 3} 0 R >> >> >>".encode())
        objetos.append(b"<< /Length " + str(len(flujo)).encode() + b" >>\nstream\n" + flujo + b"\nendstream")
    objetos[1] = f"<< /Type /Pages /Kids [{' '.join(hijos)}] /Count {len(textos)} >>".encode()
    objetos.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    salida = io.BytesIO()
    salida.write(b"%PDF-1.4\n")
    posiciones = []
    for i, obj in enumerate(objetos, start=1):
        posiciones.append(salida.tell())
        salida.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = salida.tell()
    salida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for p in posiciones:
        salida.write(f"{p:010d} 00000 n \n".encode())
    salida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return salida.getvalue()


def trd_agn_xlsx() -> bytes:
    """Una TRD en el formato habitual del AGN: encabezado con la dependencia,
    códigos D/S/Sb, la columna «Series, subseries y tipos documentales» con
    los tipos debajo de cada serie, AG/AC, CT/E/M/S y procedimiento. Con
    bordes, para que su versión en PDF (tests/fixtures/trd_agn.pdf, generada
    con LibreOffice) tenga una tabla con líneas."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side

    libro = Workbook()
    h = libro.active
    h.title = "TRD"
    h.append(["TABLA DE RETENCIÓN DOCUMENTAL"])
    h.append(["ENTIDAD PRODUCTORA: Alcaldía Municipal de Tunja"])
    h.append(["DEPENDENCIA PRODUCTORA: Secretaría General   CÓDIGO: 110"])
    h.append(["CÓDIGO", None, None, "SERIES, SUBSERIES Y TIPOS DOCUMENTALES", "RETENCIÓN", None,
              "DISPOSICIÓN FINAL", None, None, None, "PROCEDIMIENTO"])
    h.append(["D", "S", "Sb", None, "AG", "AC", "CT", "E", "M", "S", None])
    h.append(["110", "2", None, "ACTAS", None, None, None, None, None, None, None])
    h.append(["110", "2", "1", "Actas de Consejo de Gobierno\nCitación\nActa\nLista de asistencia", 2, 8, "X", None,
              "X", None, "Conservación total por su valor histórico; se digitaliza."])
    h.append(["110", "2", "2", "Actas de Comité de Archivo\nActa", 2, 3, None, None, None, "X",
              "Se selecciona una muestra del 10 %."])
    h.append(["110", "15", None, "CORRESPONDENCIA\nOficio\nMemorando", 2, 3, None, "X", None, None,
              "Se elimina al cumplir su tiempo, con acta aprobada por el comité."])
    h.append([None, None, None, "Circular", None, None, None, None, None, None, None])
    borde = Border(*(Side(style="thin"),) * 4)
    for fila in h.iter_rows(min_row=4, max_row=h.max_row, max_col=11):
        for c in fila:
            c.border = borde
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for letra, ancho in zip("ABCDEFGHIJK", (6, 6, 6, 42, 6, 6, 5, 5, 5, 5, 40)):
        h.column_dimensions[letra].width = ancho
    h.page_setup.orientation = "landscape"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _pdf_tabla(encabezado: list[str], columnas: list[float], filas: list[list[tuple[str, int]]]) -> bytes:
    """Un PDF de una página (horizontal) con líneas de encabezado y una tabla
    dibujada con bordes, como la de una TRD impresa. Cada celda es (texto,
    columnas que ocupa); el texto puede tener varias líneas. Sin programas
    externos: Helvetica con WinAnsiEncoding (admite tildes y eñes)."""
    def cadena(t: str) -> bytes:
        return b"(" + t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("cp1252") + b")"

    ops: list[bytes] = [b"0.5 w"]
    y = 560.0
    for linea in encabezado:
        ops.append(b"BT /F1 10 Tf 30 %.1f Td " % y + cadena(linea) + b" Tj ET")
        y -= 14
    y -= 6
    for fila in filas:
        alto = max(12 * max(len(t.split("\n")) for t, _ in fila), 14) + 6
        x, col = 30.0, 0
        for texto, ocupa in fila:
            ancho = sum(columnas[col:col + ocupa])
            ops.append(b"%.1f %.1f %.1f %.1f re S" % (x, y - alto, ancho, alto))
            for k, linea in enumerate(texto.split("\n")):
                if linea:
                    ops.append(b"BT /F1 8 Tf %.1f %.1f Td " % (x + 2, y - 11 - 12 * k) + cadena(linea) + b" Tj ET")
            x += ancho
            col += ocupa
        y -= alto
    contenido = b"\n".join(ops)
    objetos = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] /Resources << /Font << /F1 4 0 R >> >> "
               b"/Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
               b"<< /Length %d >>\nstream\n" % len(contenido) + contenido + b"\nendstream"]
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for i, o in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(salida)
    salida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    for p in posiciones:
        salida += b"%010d 00000 n \n" % p
    salida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref)
    return bytes(salida)


def trd_agn_pdf() -> bytes:
    """La misma TRD de trd_agn_xlsx, impresa en PDF con su tabla."""
    columnas = [28, 28, 28, 250, 30, 30, 26, 26, 26, 26, 260]
    vacias = lambda n: [("", 1)] * n  # noqa: E731
    filas = [
        [("CÓDIGO", 3), ("SERIES, SUBSERIES Y TIPOS DOCUMENTALES", 1), ("RETENCIÓN", 2), ("DISPOSICIÓN FINAL", 4),
         ("PROCEDIMIENTO", 1)],
        [("D", 1), ("S", 1), ("Sb", 1), ("", 1), ("AG", 1), ("AC", 1), ("CT", 1), ("E", 1), ("M", 1), ("S", 1), ("", 1)],
        [("110", 1), ("2", 1), ("", 1), ("ACTAS", 1)] + vacias(7),
        [("110", 1), ("2", 1), ("1", 1), ("Actas de Consejo de Gobierno\nCitación\nActa\nLista de asistencia", 1),
         ("2", 1), ("8", 1), ("X", 1), ("", 1), ("X", 1), ("", 1), ("Conservación total por su valor histórico.", 1)],
        [("110", 1), ("2", 1), ("2", 1), ("Actas de Comité de Archivo\nActa", 1), ("2", 1), ("3", 1), ("", 1), ("", 1),
         ("", 1), ("X", 1), ("Se selecciona una muestra del 10 %.", 1)],
        [("110", 1), ("15", 1), ("", 1), ("CORRESPONDENCIA\nOficio\nMemorando\nCircular", 1), ("2", 1), ("3", 1),
         ("", 1), ("X", 1), ("", 1), ("", 1), ("Se elimina con acta aprobada por el comité.", 1)],
    ]
    return _pdf_tabla(["TABLA DE RETENCIÓN DOCUMENTAL", "ENTIDAD PRODUCTORA: Alcaldía Municipal de Tunja",
                       "DEPENDENCIA PRODUCTORA: Secretaría General", "CÓDIGO: 110"], columnas, filas)
