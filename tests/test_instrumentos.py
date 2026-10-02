"""
Pruebas del Módulo 4 · Generación de instrumentos de descripción (sección
9 del prompt), contra PostgreSQL real. Los archivos exportados se abren
de verdad (openpyxl, python-docx) para revisar su contenido.
"""

import io
import json
import uuid
from datetime import date

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.base import ahora
from app.models.alerta import Alerta
from app.models.auditoria import RegistroAuditoria
from app.models.descripcion import Fecha, Relacion
from app.models.instanciacion import Instanciacion
from app.models.recurso_documental import RecursoDocumental
from app.servicios import motor, vocabulario
from tests.conftest import crear_usuario, ingresar

# Marcas únicas en los campos internos: si aparecen en una salida, se filtró algo.
MOTOR_SECRETO = "motor-interno-XYZ"
FRAGMENTO_SECRETO = "fragmento-citado-XYZ"
CLAVES_INTERNAS = {"origen", "confianza", "motor", "estado_revision", "origen_titulo", "origen_alcance",
                   "confianza_alcance", "fragmento", "documento_id"}


def _recurso(db, nivel, titulo, superior, fondo, **extra):
    r = RecursoDocumental(id=uuid.uuid4(), nivel=nivel, titulo=titulo, fondo_id=fondo.id, incluido_en_id=superior.id,
                          publicado_en=ahora(), origen_titulo="motor", motor=MOTOR_SECRETO, confianza_alcance=0.41, **extra)
    db.add(r)
    db.flush()
    db.add(Relacion(origen_tipo="recurso_documental", origen_id=superior.id, destino_tipo="recurso_documental",
                    destino_id=r.id, tipo_relacion="inclusion", codigo_ric="includes_or_included", origen="persona"))
    return r


def _fecha(db, recurso, dia):
    f = Fecha(id=uuid.uuid4(), expresion=dia.isoformat(), normalizada=dia, origen="motor", confianza=0.97, motor=MOTOR_SECRETO)
    db.add(f)
    db.flush()
    db.add(Relacion(origen_tipo="fecha", origen_id=f.id, destino_tipo="recurso_documental", destino_id=recurso.id,
                    tipo_relacion="temporal", codigo_ric="is_creation_date_of", origen="motor", confianza=0.97,
                    motor=MOTOR_SECRETO, fragmento=FRAGMENTO_SECRETO))


def _agente(db, recurso, entidad, rol="productor", codigo="has_creator"):
    db.add(Relacion(origen_tipo="recurso_documental", origen_id=recurso.id, destino_tipo="entidad_vocabulario",
                    destino_id=entidad.id, tipo_relacion="procedencia", codigo_ric=codigo, rol=rol, origen="motor",
                    confianza=0.93, motor=MOTOR_SECRETO, fragmento=FRAGMENTO_SECRETO))


@pytest.fixture()
def fondo_descrito(db, admin):
    """Fondo › Serie Correspondencia › Expediente 1948 (2 oficios) y
    Expediente 1949 descrito como un todo."""
    fondo = RecursoDocumental(id=uuid.uuid4(), nivel="fondo", titulo="Correspondencia municipal")
    fondo.fondo_id = fondo.id
    db.add(fondo)
    db.flush()
    serie = _recurso(db, "serie", "Correspondencia", fondo, fondo, codigo_referencia="CO",
                     alcance_contenido="Oficios remitidos y recibidos por la Alcaldía.")
    exp48 = _recurso(db, "expediente", "Correspondencia 1948", serie, fondo, codigo_referencia="CO-48")
    alcaldia = vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre="Alcaldía Municipal",
                                 subtipo="entidad_corporativa", origen="motor", confianza=0.9, motor=MOTOR_SECRETO, usuario_id=None)
    gobernador = vocabulario.crear(db, fondo_id=fondo.id, clase="agente", nombre="Gobernador del Departamento",
                                   subtipo="cargo", origen="persona", confianza=None, motor=None, usuario_id=None)
    boyaca = vocabulario.crear(db, fondo_id=fondo.id, clase="lugar", nombre="Boyacá", subtipo=None,
                               origen="persona", confianza=None, motor=None, usuario_id=None)
    oficio = vocabulario.crear(db, fondo_id=fondo.id, clase="forma_documental", nombre="Oficio", subtipo=None,
                               origen="persona", confianza=None, motor=None, usuario_id=None)
    db.flush()
    o114 = _recurso(db, "unidad_documental", "Oficio N.º 114", exp48, fondo, codigo_referencia="CO-AM-114",
                    caja="1", carpeta="3", folios=2, soporte="Papel", forma_documental_id=oficio.id,
                    alcance_contenido="Informa sobre el estado del archivo municipal.")
    o115 = _recurso(db, "unidad_documental", "Oficio N.º 115", exp48, fondo, codigo_referencia="CO-AM-115",
                    carpeta="3", folios=1, soporte="Papel")  # sin caja
    exp49 = _recurso(db, "expediente", "Correspondencia 1949", serie, fondo, codigo_referencia="CO-49",
                     caja="2", carpeta="1", folios=40, soporte="Papel")
    _fecha(db, o114, date(1948, 3, 15))
    _fecha(db, o115, date(1948, 3, 18))
    _fecha(db, exp49, date(1949, 1, 10))
    _fecha(db, exp49, date(1949, 12, 2))
    _agente(db, o114, alcaldia)
    _agente(db, o114, gobernador, rol="destinatario", codigo="has_addressee")
    _agente(db, o115, alcaldia)
    _agente(db, exp49, alcaldia)
    db.add(Relacion(origen_tipo="recurso_documental", origen_id=o114.id, destino_tipo="entidad_vocabulario",
                    destino_id=boyaca.id, tipo_relacion="asociacion", codigo_ric="has_or_had_subject", origen="motor",
                    confianza=0.4, motor=MOTOR_SECRETO, fragmento=FRAGMENTO_SECRETO))
    inst = Instanciacion(id=uuid.uuid4(), fondo_id=fondo.id, nombre_original="Oficio_114_1948.pdf", ruta="x/y.pdf",
                         tamano_bytes=10, estado="listo_para_descripcion", huella="a" * 64, formato_puid="fmt/18",
                         formato_nombre="Acrobat PDF 1.4")
    db.add(inst)
    db.flush()
    db.add(Relacion(origen_tipo="recurso_documental", origen_id=o114.id, destino_tipo="instanciacion", destino_id=inst.id,
                    tipo_relacion="asociacion", codigo_ric="has_or_had_instantiation", origen="persona"))
    # Una descripción sin publicar (en borrador) no aparece en ningún instrumento.
    borrador = RecursoDocumental(id=uuid.uuid4(), nivel="unidad_documental", titulo="Borrador", fondo_id=fondo.id,
                                 incluido_en_id=exp48.id)
    db.add(borrador)
    db.commit()
    return {"fondo": fondo, "serie": serie, "exp48": exp48, "exp49": exp49, "o114": o114, "o115": o115,
            "alcaldia": alcaldia}


@pytest.fixture()
def archivista(cliente, db):
    crear_usuario(db, "catalina@correo.com", "archivista", nombre="Catalina Torres")
    return ingresar(cliente, "catalina@correo.com")


def _claves(datos):
    if isinstance(datos, dict):
        for k, v in datos.items():
            yield k
            yield from _claves(v)
    elif isinstance(datos, list):
        for x in datos:
            yield from _claves(x)


def _sin_nada_interno(datos):
    assert not (set(_claves(datos)) & CLAVES_INTERNAS)
    texto = json.dumps(datos, ensure_ascii=False, default=str)
    assert MOTOR_SECRETO not in texto and FRAGMENTO_SECRETO not in texto


def _xlsx(contenido):
    hoja = load_workbook(io.BytesIO(contenido)).active
    return hoja, [[c.value for c in fila] for fila in hoja.iter_rows()]


# --- Catálogo -----------------------------------------------------------------------------------


def test_catalogo_navega_el_arbol_por_niveles(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    raiz = cliente.get("/api/instrumentos/catalogo", headers=archivista, params={"fondo_id": str(f["fondo"].id)}).json()
    assert raiz["actual"]["nivel"] == "fondo" and raiz["migas"] == []
    assert [h["titulo"] for h in raiz["hijos"]] == ["Correspondencia"]
    assert raiz["actual"]["fechas_extremas"] == "15/03/1948 – 02/12/1949"
    serie = cliente.get("/api/instrumentos/catalogo", headers=archivista,
                        params={"fondo_id": str(f["fondo"].id), "nodo_id": str(f["serie"].id)}).json()
    assert [h["titulo"] for h in serie["hijos"]] == ["Correspondencia 1948", "Correspondencia 1949"]
    assert [m["titulo"] for m in serie["migas"]] == ["Correspondencia municipal"]
    exp = cliente.get("/api/instrumentos/catalogo", headers=archivista,
                      params={"fondo_id": str(f["fondo"].id), "nodo_id": str(f["exp48"].id)}).json()
    assert [h["titulo"] for h in exp["hijos"]] == ["Oficio N.º 114", "Oficio N.º 115"]  # el borrador no sale
    assert [m["titulo"] for m in exp["migas"]] == ["Correspondencia municipal", "Correspondencia"]
    assert exp["actual"]["fechas_extremas"] == "15/03/1948 – 18/03/1948"
    for datos in (raiz, serie, exp):
        _sin_nada_interno(datos)


def test_ficha_con_entidades_vivas_y_preservacion_sin_campos_internos(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    ficha = cliente.get(f"/api/instrumentos/catalogo/{f['o114'].id}", headers=archivista).json()
    _sin_nada_interno(ficha)
    alcaldia = next(e for e in ficha["entidades"] if e["valor"] == "Alcaldía Municipal")
    assert alcaldia["en_vocabulario"] and alcaldia["documentos"] == 3 and alcaldia["entidad_id"] == str(f["alcaldia"].id)
    assert ficha["forma_documental"]["nombre"] == "Oficio"
    assert [m["titulo"] for m in ficha["migas"]] == ["Correspondencia municipal", "Correspondencia", "Correspondencia 1948"]
    [inst] = ficha["instanciaciones"]
    assert inst["preservacion"]["estado"] == "riesgo_obsolescencia" and inst["preservacion"]["puid"] == "fmt/18"
    assert ficha["control"]["caja"] == "1"
    assert cliente.get(f"/api/instrumentos/catalogo/{uuid.uuid4()}", headers=archivista).status_code == 404


# --- Inventario ---------------------------------------------------------------------------------


def test_inventario_un_renglon_por_unidad_o_expediente_con_columnas_fuid(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    r = cliente.post("/api/instrumentos/inventario", headers=archivista, json={"recurso_id": str(f["fondo"].id)})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    hoja, filas = _xlsx(r.content)
    titulos = next(fila for fila in filas if fila[0] == "N.º de orden")
    assert titulos == ["N.º de orden", "Código", "Nombre de la serie, subserie o asunto", "Fecha inicial", "Fecha final",
                       "Caja", "Carpeta", "N.º de folios", "Soporte", "Notas"]
    datos = filas[filas.index(titulos) + 1:]
    # Dos oficios del expediente 1948 (descrito por unidades) + el expediente 1949 (descrito como un todo).
    assert [d[1] for d in datos] == ["CO-AM-114", "CO-AM-115", "CO-49"]
    assert datos[0] == [1, "CO-AM-114", "Correspondencia / Oficio N.º 114", "15/03/1948", "15/03/1948", "1", "3", 2, "Papel", None]
    assert datos[2][3:5] == ["10/01/1949", "02/12/1949"] and datos[2][7] == 40


def test_campo_obligatorio_vacio_queda_pendiente_sin_bloquear(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    r = cliente.post("/api/instrumentos/inventario", headers=archivista, json={"recurso_id": str(f["exp48"].id)})
    assert r.status_code == 200 and r.headers["X-Campos-Pendientes"] == "1"
    hoja, filas = _xlsx(r.content)
    fila_115 = next(i for i, fila in enumerate(filas, start=1) if fila[1] == "CO-AM-115")
    celda = hoja.cell(row=fila_115, column=6)  # Caja
    assert celda.value == "Pendiente" and celda.font.italic and celda.comment is not None


def test_alerta_con_conteo_solo_cuando_hay_pendientes(cliente, db, fondo_descrito, archivista):
    f = fondo_descrito
    previa = cliente.post("/api/instrumentos/inventario/vista-previa", headers=archivista,
                          json={"recurso_id": str(f["exp48"].id)}).json()
    assert previa["pendientes"] == 1 and previa["pendientes_por_campo"] == {"caja": 1} and previa["alerta"]
    [alerta] = db.scalars(select(Alerta).where(Alerta.tipo == "inventario_campos_pendientes")).all()
    assert alerta.detalle["pendientes"] == 1 and alerta.entidad_id == str(f["exp48"].id)
    # Generar otra vez no duplica la alerta.
    cliente.post("/api/instrumentos/inventario", headers=archivista, json={"recurso_id": str(f["exp48"].id)})
    assert len(db.scalars(select(Alerta).where(Alerta.tipo == "inventario_campos_pendientes")).all()) == 1
    # Sin pendientes (expediente 1949 completo): ninguna alerta.
    completo = cliente.post("/api/instrumentos/inventario/vista-previa", headers=archivista,
                            json={"recurso_id": str(f["exp49"].id)}).json()
    assert completo["pendientes"] == 0 and completo["alerta"] is None
    assert db.scalars(select(Alerta).where(Alerta.entidad_id == str(f["exp49"].id))).all() == []
    # Al completar el dato, la siguiente generación deja la alerta atendida.
    f["o115"].caja = "1"
    db.commit()
    otra = cliente.post("/api/instrumentos/inventario/vista-previa", headers=archivista,
                        json={"recurso_id": str(f["exp48"].id)}).json()
    assert otra["pendientes"] == 0 and otra["alerta"] is None
    db.expire_all()
    assert db.get(Alerta, alerta.id).atendida_en is not None


# --- Guía ---------------------------------------------------------------------------------------


class MotorGuia:
    nombre = "motor-de-prueba"

    def __init__(self):
        self.pedidos = []

    def redactar(self, instruccion, pedido):
        self.pedidos.append(pedido)
        return "El fondo Correspondencia municipal reúne oficios de la Alcaldía.\n\nComprende dos expedientes."


def test_guia_redactada_por_el_motor_y_exportada_tal_como_se_edito(cliente, db, fondo_descrito, archivista, monkeypatch):
    f = fondo_descrito
    m = MotorGuia()
    monkeypatch.setattr(motor, "motor_activo", lambda: m)
    borrador = cliente.post("/api/instrumentos/guia", headers=archivista, json={"fondo_id": str(f["fondo"].id)}).json()
    assert borrador["redactado_por_motor"] and borrador["texto"].startswith("El fondo Correspondencia municipal")
    # El motor recibe solo datos validados: nada de procedencia.
    assert MOTOR_SECRETO not in m.pedidos[0] and FRAGMENTO_SECRETO not in m.pedidos[0]
    assert "Alcaldía Municipal (3 documentos)" in m.pedidos[0]
    editado = borrador["texto"].replace("dos expedientes", "dos expedientes, revisados por la archivista") + \
        "\n\nNota agregada a mano."
    r = cliente.post("/api/instrumentos/guia/exportar", headers=archivista,
                     json={"fondo_id": str(f["fondo"].id), "texto": editado})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    doc = Document(io.BytesIO(r.content))
    parrafos = [p.text for p in doc.paragraphs]
    inicio = parrafos.index("Nota de presentación") + 1
    assert parrafos[inicio:inicio + 3] == [
        "El fondo Correspondencia municipal reúne oficios de la Alcaldía.",
        "Comprende dos expedientes, revisados por la archivista.",
        "Nota agregada a mano."]
    assert db.scalars(select(RegistroAuditoria).where(RegistroAuditoria.accion == "guia_exportada")).one()


def test_guia_sin_motor_da_un_borrador_basico(cliente, fondo_descrito, archivista, monkeypatch):
    monkeypatch.setattr(motor, "motor_activo", lambda: None)
    b = cliente.post("/api/instrumentos/guia", headers=archivista, json={"fondo_id": str(fondo_descrito["fondo"].id)}).json()
    assert not b["redactado_por_motor"] and b["aviso"] and "Correspondencia municipal" in b["texto"]


# --- Índice -------------------------------------------------------------------------------------


def test_indice_agrupa_por_tipo_y_ordena_alfabeticamente(cliente, db, fondo_descrito, archivista):
    f = fondo_descrito
    for nombre in ("Ábrego, Luis", "Zapata, Ana", "concejo municipal"):
        vocabulario.crear(db, fondo_id=f["fondo"].id, clase="agente", nombre=nombre, subtipo="persona",
                          origen="persona", confianza=None, motor=None, usuario_id=None)
    db.commit()
    datos = cliente.get("/api/instrumentos/indice", headers=archivista, params={"fondo_id": str(f["fondo"].id)}).json()
    _sin_nada_interno(datos)
    assert [g["clase"] for g in datos["grupos"]] == ["agente", "lugar", "forma_documental"]
    agentes = datos["grupos"][0]
    assert [x["letra"] for x in agentes["letras"]] == ["A", "C", "G", "Z"]
    assert [e["nombre"] for x in agentes["letras"] for e in x["entidades"]] == [
        "Ábrego, Luis", "Alcaldía Municipal", "concejo municipal", "Gobernador del Departamento", "Zapata, Ana"]
    assert datos["grupos"][1]["letras"][0]["entidades"][0]["nombre"] == "Boyacá"


# --- Regla de procedencia en todo lo exportado ----------------------------------------------------


def test_ningun_archivo_exportado_lleva_campos_internos(cliente, fondo_descrito, archivista, monkeypatch):
    f = fondo_descrito
    monkeypatch.setattr(motor, "motor_activo", lambda: None)
    xlsx = cliente.post("/api/instrumentos/inventario", headers=archivista, json={"recurso_id": str(f["fondo"].id)}).content
    hoja, filas = _xlsx(xlsx)
    texto = " ".join(str(c) for fila in filas for c in fila if c is not None).lower()
    guia = cliente.post("/api/instrumentos/guia", headers=archivista, json={"fondo_id": str(f["fondo"].id)}).json()
    docx = cliente.post("/api/instrumentos/guia/exportar", headers=archivista,
                        json={"fondo_id": str(f["fondo"].id), "texto": guia["texto"]}).content
    doc = Document(io.BytesIO(docx))
    texto += " " + " ".join(p.text for p in doc.paragraphs).lower()
    texto += " " + " ".join(c.text for t in doc.tables for fila in t.rows for c in fila.cells).lower()
    for marca in (MOTOR_SECRETO.lower(), FRAGMENTO_SECRETO.lower(), "confianza", "propuesto por el motor",
                  "estado de revisión", "0.93", "0.97", "0.41"):
        assert marca not in texto, marca


# --- Permisos -----------------------------------------------------------------------------------


def test_generar_exige_rol_y_consultar_no(cliente, db, fondo_descrito):
    f = fondo_descrito
    crear_usuario(db, "julian@correo.com", "revisor")
    crear_usuario(db, "laura@correo.com", "consulta")
    revisor, consulta = ingresar(cliente, "julian@correo.com"), ingresar(cliente, "laura@correo.com")
    nivel = {"recurso_id": str(f["fondo"].id)}
    for quien in (revisor, consulta, None):
        codigo = 401 if quien is None else 403
        assert cliente.post("/api/instrumentos/inventario", headers=quien, json=nivel).status_code == codigo
        assert cliente.post("/api/instrumentos/inventario/vista-previa", headers=quien, json=nivel).status_code == codigo
        assert cliente.post("/api/instrumentos/guia", headers=quien, json={"fondo_id": str(f["fondo"].id)}).status_code == codigo
        assert cliente.post("/api/instrumentos/guia/exportar", headers=quien,
                            json={"fondo_id": str(f["fondo"].id), "texto": "x"}).status_code == codigo
    for quien in (revisor, consulta):
        assert cliente.get("/api/instrumentos/catalogo", headers=quien, params={"fondo_id": str(f["fondo"].id)}).status_code == 200
        assert cliente.get(f"/api/instrumentos/catalogo/{f['o114'].id}", headers=quien).status_code == 200
        assert cliente.get("/api/instrumentos/indice", headers=quien, params={"fondo_id": str(f["fondo"].id)}).status_code == 200
    assert cliente.get("/api/instrumentos/catalogo", params={"fondo_id": str(f["fondo"].id)}).status_code == 401


# --- Grafo visual --------------------------------------------------------------------------------


def test_grafo_de_un_documento_con_sus_relaciones_ric_y_sin_datos_internos(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    r = cliente.get("/api/instrumentos/grafo", headers=archivista,
                    params={"fondo_id": str(f["fondo"].id), "centro": f"recurso_documental:{f['o114'].id}"})
    assert r.status_code == 200
    g = r.json()
    _sin_nada_interno(g)
    etiquetas = {n["etiqueta"] for n in g["nodos"]}
    assert {"Oficio N.º 114", "Alcaldía Municipal", "Gobernador del Departamento", "Boyacá", "Oficio",
            "Oficio_114_1948.pdf", "1948-03-15", "Correspondencia 1948"} <= etiquetas
    relaciones = {(a["desde"].split(":")[0], a["hacia"].split(":")[0], a["codigo_ric"]) for a in g["aristas"]}
    assert ("recurso_documental", "entidad_vocabulario", "has_creator") in relaciones
    assert ("recurso_documental", "instanciacion", "has_or_had_instantiation") in relaciones
    assert ("fecha", "recurso_documental", "is_creation_date_of") in relaciones
    assert ("recurso_documental", "recurso_documental", "includes_or_included") in relaciones
    assert any(a["uri_rico"] == "rico:hasCreator" and a["etiqueta"] == "producido por" for a in g["aristas"])
    # Toda arista une nodos presentes.
    claves = {n["clave"] for n in g["nodos"]}
    assert all(a["desde"] in claves and a["hacia"] in claves for a in g["aristas"])


def test_grafo_desde_una_entidad_muestra_los_documentos_que_la_comparten(cliente, fondo_descrito, archivista):
    f = fondo_descrito
    g = cliente.get("/api/instrumentos/grafo", headers=archivista,
                    params={"fondo_id": str(f["fondo"].id), "centro": f"entidad_vocabulario:{f['alcaldia'].id}"}).json()
    documentos = {n["etiqueta"] for n in g["nodos"] if n["tipo"] == "recurso_documental"}
    assert documentos == {"Oficio N.º 114", "Oficio N.º 115", "Correspondencia 1949"}  # el borrador no aparece
    assert next(n for n in g["nodos"] if n["etiqueta"] == "Alcaldía Municipal")["documentos"] == 3
    # Por defecto, el fondo; con dos saltos llega a los expedientes.
    raiz = cliente.get("/api/instrumentos/grafo", headers=archivista,
                       params={"fondo_id": str(f["fondo"].id), "profundidad": 2}).json()
    assert {"Correspondencia", "Correspondencia 1948", "Correspondencia 1949"} <= {n["etiqueta"] for n in raiz["nodos"]}
    # El rol consulta también lo ve (es consulta del catálogo); sin sesión, no.
    assert cliente.get("/api/instrumentos/grafo", params={"fondo_id": str(f["fondo"].id)}).status_code == 401
    otro = cliente.get("/api/instrumentos/grafo", headers=archivista,
                       params={"fondo_id": str(f["fondo"].id), "centro": f"recurso_documental:{uuid.uuid4()}"})
    assert otro.status_code == 404
