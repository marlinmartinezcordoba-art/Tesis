"""
Perfil RiC-Col del AGN (Esquema de Metadatos v1.4) aplicado a RICORA:
catálogo de correspondencias verificado contra el OWL oficial, campos que
pide el AGN (estado de conservación, signatura topográfica, datos
personales, accesibilidad, ORCID y ROR) y medición de calidad.
"""

import io
import re
import uuid
from functools import lru_cache
from pathlib import Path

import pytest
from openpyxl import load_workbook
from rdflib import Graph, Namespace
from sqlalchemy import select

from app.models.auditoria import RegistroAuditoria
from app.models.instanciacion import Instanciacion
from app.servicios import autoridad, exportacion_rico, perfil_agn, ric_o
from tests.test_instrumentos import archivista, fondo_descrito  # noqa: F401

RICO = Namespace(ric_o.RICO)
OWL_RICO = Path(__file__).resolve().parent.parent / "app" / "recursos" / "ric-o" / "RiC-O_1-1.rdf"
ENTIDADES_RIC_CM = {f"E{n:02d}" for n in range(1, 19)} | {"E22"}


@lru_cache(maxsize=1)
def _owl():
    g = Graph().parse(OWL_RICO)
    nombres = {str(s)[len(ric_o.RICO):] for s in g.subjects() if str(s).startswith(ric_o.RICO)}
    codigos = set()
    for s in g.subjects():
        if str(s).startswith(ric_o.RICO):
            for o in g.objects(s, None):
                codigos |= set(re.findall(r"RiC-([AR]\d{2,3})\b", str(o)))
    return nombres, codigos


def test_cada_termino_rico_del_perfil_existe_en_el_owl_oficial():
    nombres, _ = _owl()
    faltan = {(f.id, t) for f in perfil_agn.CATALOGO for t in f.ric_o if t not in nombres}
    assert not faltan, faltan


def test_cada_codigo_ric_cm_citado_como_correcto_es_real():
    """Los códigos de la columna «RiC-CM 1.0 (correcto)»: atributos y
    relaciones que el OWL cita, y entidades de la lista oficial."""
    _, codigos = _owl()
    for f in perfil_agn.CATALOGO:
        for c in re.findall(r"RiC-([AER]\d{2,3})\b", f.ric_cm):
            if c.startswith("E"):
                assert c in ENTIDADES_RIC_CM, (f.id, c)
            else:
                assert c in codigos, (f.id, c)


def test_capas_estados_y_niveles_validos_y_ids_unicos():
    ids = [f.id for f in perfil_agn.CATALOGO]
    assert len(ids) == len(set(ids))
    for f in perfil_agn.CATALOGO:
        assert f.capa in perfil_agn.CAPAS and f.estado in perfil_agn.ESTADOS and f.nivel in perfil_agn.NIVELES, f.id
    # Cada fila del FUID del anexo 2 del AGN tiene su correspondencia.
    assert sum(1 for f in perfil_agn.CATALOGO if f.origen == "Anexo 2") == 17
    # Una decisión propia nunca se presenta como del AGN.
    assert all(f.id.startswith("RIC-") for f in perfil_agn.CATALOGO if f.capa == "ricora")


def test_catalogo_y_excel_por_la_api(cliente, archivista):
    r = cliente.get("/api/calidad/correspondencias", headers=archivista)
    assert r.status_code == 200, r.text
    datos = r.json()
    assert datos["fuentes"]["agn"]["version"] == "1.4" and datos["fuentes"]["ric_o"]["version"] == "1.1"
    assert datos["codigos_agn_incorrectos"] > 0 and len(datos["erratas"]) == len(perfil_agn.ERRATAS)
    x = cliente.get("/api/calidad/correspondencias.xlsx", headers=archivista)
    assert x.status_code == 200
    libro = load_workbook(io.BytesIO(x.content))
    assert libro.sheetnames == ["Perfil RiC-Col", "Erratas del esquema AGN", "Fuentes"]
    assert libro["Perfil RiC-Col"].max_row == len(perfil_agn.CATALOGO) + 1
    assert cliente.get("/api/calidad/correspondencias").status_code == 401


# --- ORCID y ROR (AGN, tabla 5) ------------------------------------------------------------------


def test_orcid_y_ror_con_digito_de_control_y_tipo_de_agente(db, fondo_descrito, admin):
    alcaldia = fondo_descrito["alcaldia"]
    i = autoridad.agregar_identificador(db, alcaldia, esquema="ror", valor="https://ror.org/05dxps055",
                                        usuario_id=admin.id)
    assert i.valor == "05dxps055"
    assert autoridad.uri_externa("ror", i.valor) == "https://ror.org/05dxps055"
    with pytest.raises(autoridad.ErrorAutoridad, match="personas"):
        autoridad.agregar_identificador(db, alcaldia, esquema="orcid", valor="0000-0002-1825-0097",
                                        usuario_id=admin.id)
    with pytest.raises(autoridad.ErrorAutoridad, match="control"):
        autoridad.agregar_identificador(db, alcaldia, esquema="ror", valor="05dxps056", usuario_id=admin.id)
    assert autoridad.orcid_normalizado("0000000218250097") == "0000-0002-1825-0097"
    with pytest.raises(autoridad.ErrorAutoridad, match="control"):
        autoridad.orcid_normalizado("0000-0002-1825-0098")


# --- Original físico: estado de conservación y signatura (AGN, tabla 4) ----------------------------


def test_original_fisico_con_estado_y_signatura_se_corrige_con_auditoria_y_sale_en_rico(
        cliente, db, fondo_descrito, archivista):
    o114 = fondo_descrito["o114"]
    r = cliente.post(f"/api/descripcion/registros/{o114.id}/original-fisico", headers=archivista, json={
        "soporte": "papel", "ubicacion": "Archivo central", "estado_conservacion": "regular", "deposito": "2",
        "estante": "14", "entrepano": "3", "caracteristicas_fisicas": "Manchas de humedad en el margen."})
    assert r.status_code == 201, r.text
    [fisico] = [i for i in r.json()["instanciaciones"] if i.get("fisica")]
    assert fisico["estado_conservacion"] == "regular"
    assert fisico["signatura"] == "Depósito 2, estante 14, entrepaño 3"
    malo = cliente.post(f"/api/descripcion/registros/{o114.id}/original-fisico", headers=archivista,
                        json={"soporte": "papel", "estado_conservacion": "pésimo"})
    assert malo.status_code == 422

    r = cliente.patch(f"/api/descripcion/registros/{o114.id}/original-fisico/{fisico['id']}", headers=archivista,
                      json={"estado_conservacion": "restaurado"})
    assert r.status_code == 200, r.text
    db.expire_all()
    assert db.get(Instanciacion, uuid.UUID(fisico["id"])).estado_conservacion == "restaurado"
    a = db.scalar(select(RegistroAuditoria).where(RegistroAuditoria.accion == "original_fisico_actualizado"))
    assert a.valor_anterior == {"estado_conservacion": "regular"} and a.valor_nuevo == {"estado_conservacion": "restaurado"}

    g = exportacion_rico.exportar(db, fondo_descrito["fondo"], incluir_restringidos=True).grafo
    notas = {str(o) for o in g.objects(None, RICO.physicalCharacteristicsNote)}
    assert "Estado de conservación: restaurado; Manchas de humedad en el margen." in notas


# --- Datos personales, accesibilidad y calidad -------------------------------------------------------


def _editar(cliente, cab, recurso_id, **datos):
    trabajo = cliente.post(f"/api/descripcion/registros/{recurso_id}/reabrir", headers=cab)
    assert trabajo.status_code in (200, 201), trabajo.text
    return cliente.patch(f"/api/descripcion/{recurso_id}", headers=cab,
                         json={"trabajo_id": trabajo.json()["trabajo_id"], **datos})


def test_datos_sensibles_en_documento_publico_es_incoherencia_y_la_nota_de_accesibilidad_es_publica(
        cliente, db, fondo_descrito, archivista):
    o114 = fondo_descrito["o114"]
    r = _editar(cliente, archivista, o114.id, proteccion={
        "datos_personales": "sensibles", "nota_accesibilidad": "Transcripción en texto plano para lector de pantalla."})
    assert r.status_code == 200, r.text
    assert r.json()["proteccion"]["datos_personales"] == "sensibles"
    assert db.scalar(select(RegistroAuditoria).where(RegistroAuditoria.accion == "proteccion_datos_declarada"))

    c = cliente.get(f"/api/calidad/descripciones/{o114.id}", headers=archivista).json()
    assert c["incoherencias"] and "Ley 1712" in c["incoherencias"][0]
    criterio = {x["clave"]: x for x in c["criterios"]}
    assert criterio["datos_personales"]["cumple"] and criterio["accesibilidad"]["cumple"]
    assert criterio["identificador"]["cumple"] and criterio["tipo_documental"]["cumple"]
    assert criterio["productor"]["cumple"]

    ficha = cliente.get(f"/api/catalogo/registros/{o114.id}", headers=archivista)
    assert ficha.status_code == 200, ficha.text
    assert ficha.json()["nota_accesibilidad"].startswith("Transcripción")
    assert "proteccion" not in ficha.json() and "datos_personales" not in ficha.json()

    r = _editar(cliente, archivista, o114.id, clasificacion={
        "acceso": "clasificado", "fundamento": "Ley 1712 de 2014, art. 18, literal a"})
    assert r.status_code == 200, r.text
    assert cliente.get(f"/api/calidad/descripciones/{o114.id}", headers=archivista).json()["incoherencias"] == []


def test_productor_heredado_del_nivel_superior_y_kpi_del_fondo(cliente, db, fondo_descrito, archivista):
    """El oficio 115 no tiene productor propio pero el expediente 1949 sí; la
    serie lo hereda solo si un nivel superior lo declara (aquí no)."""
    serie = fondo_descrito["serie"]
    c = perfil_agn.calidad(db, serie)
    assert {x["clave"]: x["cumple"] for x in c["criterios"]}["productor"] is False
    k = cliente.get(f"/api/calidad/fondos/{fondo_descrito['fondo'].id}", headers=archivista)
    assert k.status_code == 200, k.text
    datos = k.json()
    assert datos["descripciones"] >= 5
    por_clave = {x["clave"]: x for x in datos["criterios"]}
    assert 0 <= por_clave["titulo"]["porcentaje"] <= 100 and por_clave["titulo"]["porcentaje"] == 100
    assert sum(datos["nivel_alcanzado"].values()) == datos["descripciones"]
    assert cliente.get(f"/api/calidad/fondos/{uuid.uuid4()}", headers=archivista).status_code == 404


def test_literal_de_conservacion_no_aparece_sin_dato(db, fondo_descrito):
    g = exportacion_rico.exportar(db, fondo_descrito["fondo"], incluir_restringidos=True).grafo
    assert not any("Estado de conservación" in str(o) for o in g.objects(None, RICO.physicalCharacteristicsNote))
