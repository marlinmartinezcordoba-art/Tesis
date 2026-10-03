"""
Evaluación ciega del motor frente a archivistas (objetivo 3 de la tesis).
"""

import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.auditoria import RegistroAuditoria
from app.models.evaluacion import Anotacion
from app.servicios import evaluacion as ev_s
from app.servicios import motor
from tests.conftest import crear_usuario, ingresar
from tests.test_descripcion import MotorDePrueba, aceptar_todo, documento, fondo, iniciar  # noqa: F401

TEXTO = ("Oficio 210 de 1948. La Alcaldía Municipal de Tunja remite al Concejo Municipal el informe de permisos "
         "para las fiestas.")
PROPUESTA = {"titulo": "Oficio de la Alcaldía al Concejo", "alcance_contenido": "Informe de permisos.", "entidades": [
    {"tipo": "agente", "subtipo": "entidad_corporativa", "rol": "productor", "valor": "Alcaldía Municipal de Tunja",
     "fragmento": "La Alcaldía Municipal de Tunja", "documento": 1, "confianza": 0.9},
    {"tipo": "agente", "subtipo": "entidad_corporativa", "rol": "destinatario", "valor": "Concejo Municipal",
     "fragmento": "al Concejo Municipal", "documento": 1, "confianza": 0.9},
    {"tipo": "lugar", "valor": "Tunja", "fragmento": "Tunja", "documento": 1, "confianza": 0.8},
    {"tipo": "fecha", "valor": "1948", "subtipo_fecha": "simple", "edtf": "1948", "fragmento": "de 1948",
     "documento": 1, "confianza": 0.9},
    {"tipo": "forma_documental", "valor": "Oficio", "fragmento": "Oficio", "documento": 1, "confianza": 0.9},
]}
# Lo que describe una archivista a ciegas: el productor sin tilde (igual al
# normalizar), el destinatario más largo (solo coincide con umbral flexible),
# otro lugar (falso positivo y falso negativo), la fecha y la forma iguales.
CIEGA_A = {"titulo": "Oficio 210 de 1948", "entidades": [
    {"tipo": "agente", "rol": "productor", "valor": "Alcaldia Municipal de Tunja"},
    {"tipo": "agente", "rol": "destinatario", "valor": "Concejo Municipal de Tunja"},
    {"tipo": "lugar", "valor": "Boyacá"},
    {"tipo": "fecha", "valor": "1948", "edtf": "1948"},
    {"tipo": "forma_documental", "valor": "oficio"}]}
CIEGA_B = {"titulo": "Oficio sobre permisos", "entidades": [
    {"tipo": "agente", "rol": "productor", "valor": "Alcaldía Municipal de Tunja"},
    {"tipo": "fecha", "valor": "1948", "edtf": "1948"},
    {"tipo": "forma_documental", "valor": "Oficio"}]}


# --- Métricas (a mano) ---------------------------------------------------------------------------


def test_emparejamiento_uno_a_uno_por_tipo_y_rol():
    s = [{"tipo": "agente", "rol": "productor", "valor": "Alcaldía"}, {"tipo": "agente", "rol": "productor", "valor": "Alcaldía"}]
    r = [{"tipo": "agente", "rol": "productor", "valor": "alcaldia"}]
    assert ev_s.emparejar(s, r, 1.0) == 1  # una referencia no cuenta dos veces
    assert ev_s.emparejar(r, s, 1.0) == 1  # ni una entidad del sistema
    otro_rol = [{"tipo": "agente", "rol": "destinatario", "valor": "Alcaldía"}]
    assert ev_s.emparejar(s, otro_rol, 1.0) == 0  # el productor no se empareja con el destinatario
    assert ev_s.emparejar([{"tipo": "fecha", "valor": "marzo 1948", "edtf": "1948-03"}],
                          [{"tipo": "fecha", "valor": "1948-03", "edtf": "1948-03"}], 1.0) == 1  # la fecha, por su EDTF


def test_precision_exhaustividad_y_f1_calculadas_a_mano():
    sistema = ev_s.entidades_de_propuesta(PROPUESTA | {"entidades": [e | {} for e in PROPUESTA["entidades"]]})
    referencia = ev_s.entidades_de_anotacion(CIEGA_A)
    estricto = ev_s.prf(ev_s.conteo(sistema, referencia, 1.0))
    assert (estricto["vp"], estricto["fp"], estricto["fn"]) == (3, 2, 2)
    assert estricto["precision"] == 0.6 and estricto["exhaustividad"] == 0.6 and estricto["f1"] == 0.6
    flexible = ev_s.prf(ev_s.conteo(sistema, referencia, 0.75))  # «Concejo Municipal» ≈ «Concejo Municipal de Tunja»
    assert (flexible["vp"], flexible["fp"], flexible["fn"]) == (4, 1, 1) and flexible["f1"] == 0.8


def test_kappa_ponderado_cuadratico():
    assert ev_s.kappa_ponderado([(1, 1), (2, 2), (3, 3), (4, 4), (5, 5)]) == 1.0
    assert ev_s.kappa_ponderado([(1, 1), (5, 5), (1, 5), (5, 1)]) == 0.0  # tanto acuerdo como el azar
    # A mano: desacuerdo observado 4 · ¼ · 1/16 = 0,0625; esperado (1/16) · (1/16) · Σ(i−j)² = (1/256) · 80
    # = 0,3125 (categorías 1, 2, 4 y 5, cada una con ¼); κ = 1 − 0,0625 / 0,3125 = 0,8.
    assert ev_s.kappa_ponderado([(1, 2), (2, 1), (4, 5), (5, 4)]) == 0.8
    assert ev_s.kappa_ponderado([]) is None


# --- Flujo y ceguera -------------------------------------------------------------------------------


@pytest.fixture()
def motor_eval(monkeypatch):
    m = MotorDePrueba(PROPUESTA)
    monkeypatch.setattr(motor, "motor_activo", lambda: m)
    return m


def evaluadora(cliente, db, correo, nombre):
    crear_usuario(db, correo, "archivista", nombre=nombre)
    return ingresar(cliente, correo)


def preparar(cliente, db, fondo, adm, umbral=0.75):
    d = documento(db, fondo, "Oficio_210_1948.pdf", TEXTO)
    ev = cliente.post("/api/evaluacion", headers=adm, json={"fondo_id": str(fondo.id), "nombre": "Prueba piloto",
                                                             "protocolo": "Dos archivistas, a ciegas y asistidas.",
                                                             "umbral_similitud": umbral}).json()
    assert cliente.post(f"/api/evaluacion/{ev['id']}/documentos", headers=adm,
                        json={"instanciacion_ids": [str(d.id)]}).status_code == 200
    # No se inicia sin las propuestas del motor.
    assert cliente.post(f"/api/evaluacion/{ev['id']}/estado", headers=adm, json={"estado": "en_curso"}).status_code == 409
    r = cliente.post(f"/api/evaluacion/{ev['id']}/propuestas", headers=adm).json()
    assert r["generadas"] == 1 and "entidades" not in str(r)  # se genera sin mostrarla
    assert cliente.post(f"/api/evaluacion/{ev['id']}/estado", headers=adm, json={"estado": "en_curso"}).status_code == 200
    return ev["id"], d


def anotar(cliente, cab, ev_id, inst_id, condicion, datos=None, enviar=True):
    r = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{inst_id}/anotar", headers=cab, json={"condicion": condicion})
    if r.status_code != 200:
        return r
    a = r.json()
    if datos is not None:
        r = cliente.put(f"/api/evaluacion/anotaciones/{a['id']}", headers=cab, json=datos | {"enviar": enviar})
    return r


def test_ciclo_completo_con_metricas_acuerdo_tiempos_y_rubrica(cliente, db, fondo, cabeceras_admin, motor_eval):
    ev_id, d = preparar(cliente, db, fondo, cabeceras_admin)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    luis = evaluadora(cliente, db, "luis@correo.com", "Luis Pardo")
    # A ciegas: el documento sin la propuesta.
    r = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "ciega"})
    assert r.status_code == 200 and r.json()["datos"] == {"entidades": []} and "Alcaldía" in r.json()["documento"]["texto"]
    # Con la descripción ciega abierta no se puede ver la propuesta.
    assert cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/propuesta", headers=ana).status_code == 409
    assert cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana,
                        json={"condicion": "asistida"}).status_code == 409
    assert anotar(cliente, ana, ev_id, d.id, "ciega", CIEGA_A).status_code == 200
    assert anotar(cliente, luis, ev_id, d.id, "ciega", CIEGA_B).status_code == 200
    # Asistida: empieza con la propuesta puesta.
    asistida = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "asistida"}).json()
    assert len(asistida["datos"]["entidades"]) == 5
    cliente.put(f"/api/evaluacion/anotaciones/{asistida['id']}", headers=ana, json=asistida["datos"] | {"enviar": True})
    # Rúbrica: verla primero; los dos califican.
    assert cliente.put(f"/api/evaluacion/{ev_id}/documentos/{d.id}/calificacion", headers=luis,
                       json={"exactitud": 4, "completitud": 3, "pertinencia": 4}).status_code == 409
    for cab, notas in ((ana, {"exactitud": 4, "completitud": 3, "pertinencia": 5}),
                       (luis, {"exactitud": 4, "completitud": 4, "pertinencia": 5})):
        cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/propuesta", headers=cab)
        assert cliente.put(f"/api/evaluacion/{ev_id}/documentos/{d.id}/calificacion", headers=cab, json=notas).status_code == 200
    # Para medir tiempos, la ciega de Ana tardó 30 minutos.
    a = db.scalar(select(Anotacion).where(Anotacion.condicion == "ciega", Anotacion.datos["titulo"].astext == "Oficio 210 de 1948"))
    a.iniciada_en = a.enviada_en - timedelta(minutes=30)
    db.commit()

    # La archivista no ve los resultados.
    assert cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=ana).status_code == 403
    res = cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=cabeceras_admin).json()
    # Estricto, sumando las dos referencias: Ana (3 VP, 2 FP, 2 FN) y Luis (3 VP, 2 FP, 0 FN).
    assert (res["estricto"]["micro"]["vp"], res["estricto"]["micro"]["fp"], res["estricto"]["micro"]["fn"]) == (6, 4, 2)
    assert res["estricto"]["micro"]["precision"] == 0.6 and res["estricto"]["micro"]["exhaustividad"] == 0.75
    assert res["flexible"]["micro"]["vp"] == 7
    assert res["flexible"]["por_tipo"]["fecha"]["f1"] == 1.0 and res["flexible"]["por_tipo"]["lugar"]["vp"] == 0
    # Acuerdo entre Ana y Luis: 3 en común de 5 y 3.
    assert res["acuerdo_entre_archivistas"]["vp"] == 3 and res["acuerdo_entre_archivistas"]["f1"] == 0.75
    assert res["tiempos"]["ciega"]["n"] == 2 and res["tiempos"]["asistida"]["n"] == 1
    assert res["tiempos"]["ciega"]["media_minutos"] >= 15
    assert res["rubrica"]["exactitud"]["kappa_ponderado"] is None or res["rubrica"]["exactitud"]["pares"] == 1
    assert res["rubrica"]["pertinencia"]["media"] == 5 and res["rubrica"]["completitud"]["media"] == 3.5
    assert res["anotaciones"] == {"ciegas": 2, "asistidas": 1, "evaluadores": 2}
    assert res["motores"] == [{"motor": "motor-de-prueba", "version_prompt": motor.VERSION_PROMPT}]
    hoja = cliente.get(f"/api/evaluacion/{ev_id}/resultados/hoja-de-calculo", headers=cabeceras_admin)
    assert load_workbook(io.BytesIO(hoja.content)).sheetnames == ["Entidades", "Acuerdo, tiempos y rúbrica", "Documentos"]
    # Todo queda en la auditoría.
    acciones = {e.accion for e in db.scalars(select(RegistroAuditoria).where(RegistroAuditoria.modulo == "evaluacion"))}
    assert {"evaluacion_creada", "documentos_agregados", "propuestas_generadas", "evaluacion_iniciada",
            "anotacion_iniciada", "anotacion_enviada", "propuesta_vista", "calificacion_registrada",
            "resultados_consultados"} <= acciones


def test_quien_vio_la_propuesta_ya_no_describe_a_ciegas(cliente, db, fondo, cabeceras_admin, motor_eval, admin):
    ev_id, d = preparar(cliente, db, fondo, cabeceras_admin)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    # Asistida primero: después, la ciega queda cerrada para ella.
    assert anotar(cliente, ana, ev_id, d.id, "asistida").status_code == 200
    r = anotar(cliente, ana, ev_id, d.id, "ciega")
    assert r.status_code == 409 and "ya vio la propuesta" in r.json()["detail"]
    # Quien lo abrió en Descripción (donde el motor le mostró su propuesta), tampoco.
    luis = evaluadora(cliente, db, "luis@correo.com", "Luis Pardo")
    iniciar(cliente, luis, [d.id])
    r = anotar(cliente, luis, ev_id, d.id, "ciega")
    assert r.status_code == 409 and "Descripción" in r.json()["detail"]
    # La administradora, después de mirar los resultados, tampoco.
    cliente.get(f"/api/evaluacion/{ev_id}/resultados", headers=cabeceras_admin)
    assert anotar(cliente, cabeceras_admin, ev_id, d.id, "ciega").status_code == 409
    tareas = cliente.get(f"/api/evaluacion/{ev_id}/tareas", headers=ana).json()["tareas"]
    assert tareas[0]["puede_ciega"] is False


def test_solo_documentos_sin_describir_y_reglas_de_estado(cliente, db, fondo, cabeceras_admin, motor_eval):
    descrito = documento(db, fondo, "ya.pdf", TEXTO)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    espacio = iniciar(cliente, ana, [descrito.id])
    assert cliente.post("/api/descripcion/publicar", headers=ana, json=aceptar_todo(espacio)).status_code == 201
    ev = cliente.post("/api/evaluacion", headers=cabeceras_admin, json={"fondo_id": str(fondo.id), "nombre": "Otra"}).json()
    r = cliente.post(f"/api/evaluacion/{ev['id']}/documentos", headers=cabeceras_admin,
                     json={"instanciacion_ids": [str(descrito.id)]})
    assert r.status_code == 422 and "contaminaría" in r.json()["detail"]
    # Ni la archivista crea evaluaciones ni el revisor anota.
    assert cliente.post("/api/evaluacion", headers=ana, json={"fondo_id": str(fondo.id), "nombre": "X" * 5}).status_code == 403
    crear_usuario(db, "revisor@correo.com", "revisor")
    revisor = ingresar(cliente, "revisor@correo.com")
    assert cliente.get("/api/evaluacion", headers=revisor, params={"fondo_id": str(fondo.id)}).status_code == 403
    # La archivista solo ve las evaluaciones en curso.
    assert cliente.get("/api/evaluacion", headers=ana, params={"fondo_id": str(fondo.id)}).json() == []


def test_una_anotacion_ajena_no_se_toca_y_la_anulada_no_se_borra(cliente, db, fondo, cabeceras_admin, motor_eval):
    ev_id, d = preparar(cliente, db, fondo, cabeceras_admin)
    ana = evaluadora(cliente, db, "ana@correo.com", "Ana Ruiz")
    luis = evaluadora(cliente, db, "luis@correo.com", "Luis Pardo")
    a = cliente.post(f"/api/evaluacion/{ev_id}/documentos/{d.id}/anotar", headers=ana, json={"condicion": "ciega"}).json()
    assert cliente.put(f"/api/evaluacion/anotaciones/{a['id']}", headers=luis, json=CIEGA_B).status_code == 403
    sin_rol = CIEGA_B | {"entidades": [{"tipo": "agente", "valor": "Alcaldía"}]}
    assert cliente.put(f"/api/evaluacion/anotaciones/{a['id']}", headers=ana, json=sin_rol).status_code == 422
    assert cliente.post(f"/api/evaluacion/anotaciones/{a['id']}/anular", headers=ana).status_code == 403
    assert cliente.post(f"/api/evaluacion/anotaciones/{a['id']}/anular", headers=cabeceras_admin).json()["estado"] == "anulada"
    assert db.get(Anotacion, uuid.UUID(a["id"])) is not None
